"""Глава «Электронная модель системы теплоснабжения» схемы теплоснабжения (фаза 5): черновик в Word по данным города.

Состав главы — п. 55 Требований к схемам теплоснабжения (утв. постановлением Правительства РФ от 22.02.2012 № 154,
подпункты а–к; п. 56 — при актуализации глава отражает изменения гидравлических режимов). По каждому подпункту —
как требование выполняет ТГИД, состояние данных города и цифры из базы: состав и связность сети, заполненность
паспорта, расчётные единицы, расчёты, переключения, балансы нагрузки по источникам и фрагментам, потери, надёжность,
сценарии. Чего в базе нет — так и пишется («нет данных»): глава — заготовка для разработчика схемы, а не отчёт
о выполнении.

Сеть — выбранные фрагменты; не выбраны — все, кроме сценариев и фрагментов-копий: в Астане магистрали лежат
копиями по сезонам, и все фрагменты разом считают трубу несколько раз. Фрагмент — копия, если не меньше COPY_SHARE его
участков совпадает с участками более нового фрагмента (failure_forecast.canonical_lines); такие фрагменты глава
называет отдельно.
"""
from __future__ import annotations

import io
from collections import defaultdict
from datetime import datetime
from typing import Any, Optional

import asyncpg

from database.net_srid import line_inner_sql, lines_have_inner, net_srid

BASIS = ("п. 55 Требований к схемам теплоснабжения, утверждённых постановлением Правительства Российской "
         "Федерации от 22.02.2012 № 154")
OK, PARTIAL, NO_DATA = "выполняется", "частично", "нет данных"
TOP_ROWS = 30                      # строк в таблицах балансов — остальное одной строкой «прочие»
COPY_SHARE = 0.5                   # доля участков-копий, с которой фрагмент считается копией более нового

# подпункты п. 55 — краткие формулировки (полный текст — в постановлении)
ITEMS: dict[str, str] = {
    "а": "Графическое представление объектов с привязкой к топооснове и описанием связности",
    "б": "Паспортизация объектов системы теплоснабжения",
    "в": "Паспортизация и описание расчётных единиц территориального деления",
    "г": "Гидравлический расчёт сетей любой закольцованности, совместная работа источников",
    "д": "Моделирование переключений, в том числе нагрузок между источниками",
    "е": "Балансы тепловой энергии по источникам и по территориальному признаку",
    "ж": "Потери тепловой энергии через изоляцию и с утечками",
    "з": "Показатели надёжности теплоснабжения",
    "и": "Групповые изменения характеристик объектов для перспективных вариантов",
    "к": "Сравнительные пьезометрические графики для сценариев развития",
}

PASSPORT_COLUMNS = (
    ("diametercondit", "условный диаметр"),
    ("tubingtypeid", "способ прокладки"),
    ("isolmaterialid", "материал изоляции"),
    ("isolationtypeid", "тип изоляции"),
    ("nettype", "вид сети"),
    ("firstpicdatehp", "дата прокладки"),
    ("lasttransdate", "дата перекладки"),
    ("repairdatecapital", "дата капремонта"),
)
DIAMETER_CLASSES = (100, 200, 300, 500, 700)


class ChapterError(Exception):
    def __init__(self, status: int, message: str):
        super().__init__(message)
        self.status = status
        self.message = message


def section(key: str, status: str, text: list[str], tables: Optional[list[dict[str, Any]]] = None) -> dict[str, Any]:
    return {"key": key, "title": ITEMS[key], "status": status, "text": text, "tables": tables or []}


def fmt(v: Any, digits: int = 1) -> str:
    if v is None:
        return "—"
    if isinstance(v, float):
        return f"{v:,.{digits}f}".replace(",", " ").replace(".", ",")
    if isinstance(v, int):
        return f"{v:,}".replace(",", " ")
    return str(v)


def diameter_class(d: Optional[float]) -> str:
    if d is None or d <= 0:
        return "нет данных"
    prev = 0
    for b in DIAMETER_CLASSES:
        if d <= b:
            return f"{prev + 1 if prev else 'до'}{'–' if prev else ' '}{b}"
        prev = b
    return f"свыше {DIAMETER_CLASSES[-1]}"


def graph_stats(lines: dict[int, tuple[int, int]], sources: set[int], closed: set[int]) -> dict[str, Any]:
    """Связность: узлы, компоненты, независимые кольца (цикломатическое число E − V + C), компоненты с несколькими
    источниками, узлы без связи с источником. Нормально закрытые участки не проводят."""
    parent: dict[int, int] = {}

    def find(k: int) -> int:
        parent.setdefault(k, k)
        while parent[k] != k:
            parent[k] = parent[parent[k]]
            k = parent[k]
        return k

    edges = 0
    for lid, (a, b) in lines.items():
        find(a)
        find(b)
        if lid in closed:
            continue
        edges += 1
        ra, rb = find(a), find(b)
        if ra != rb:
            parent[ra] = rb
    nodes = list(parent)
    comps: dict[int, list[int]] = defaultdict(list)
    for n in nodes:
        comps[find(n)].append(n)
    src_in = defaultdict(int)
    for s in sources:
        if s in parent:
            src_in[find(s)] += 1
    fed = {r for r in comps if src_in[r]}
    return {"nodes": len(nodes), "edges": edges, "components": len(comps),
            "rings": edges - len(nodes) + len(comps), "fed_components": len(fed),
            "multi_source_components": sum(1 for r in fed if src_in[r] > 1),
            "unfed_nodes": sum(len(comps[r]) for r in comps if r not in fed),
            "component_of": {n: find(n) for n in nodes}, "sources_in": dict(src_in)}


def top_rows(rows: list[list[Any]], key_index: int, label: str = "прочие") -> list[list[Any]]:
    """Первые TOP_ROWS по убыванию столбца key_index, остальные — одной строкой (числовые столбцы складываются)."""
    rows = sorted(rows, key=lambda r: -(r[key_index] or 0))
    if len(rows) <= TOP_ROWS:
        return rows
    rest = rows[TOP_ROWS:]
    tail: list[Any] = [f"{label} ({len(rest)})"]
    for i in range(1, len(rows[0])):
        values = [r[i] for r in rest]
        tail.append(sum(v for v in values if isinstance(v, (int, float))) if all(
            isinstance(v, (int, float)) or v is None for v in values) else "")
    return rows[:TOP_ROWS] + [tail]


# --- данные --------------------------------------------------------------------------------------

def copy_fragments(lines: list[dict[str, Any]]) -> dict[int, float]:
    """Фрагменты, у которых не меньше COPY_SHARE участков — копии участков более нового фрагмента: {id: доля}."""
    import pandas as pd

    from database.failure_forecast import canonical_lines

    if not lines:
        return {}
    df = pd.DataFrame(lines, columns=["id", "fileid", "ends"])
    copies = canonical_lines(df)
    total = df.groupby("fileid").size()
    copied = df[df.id.isin(copies)].groupby("fileid").size()
    share = (copied / total).dropna()
    return {int(f): round(float(v), 3) for f, v in share.items() if v >= COPY_SHARE}


async def _fileids(conn, fileids: Optional[list[int]]) -> tuple[list[int], dict[int, float]]:
    """Фрагменты модели и исключённые копии {id: доля участков-копий}."""
    scen = set()
    if await conn.fetchval("SELECT to_regclass('public.tgid_scenario') IS NOT NULL"):
        scen = {r[0] for r in await conn.fetch("SELECT fileid FROM tgid_scenario")}
    all_ids = [r[0] for r in await conn.fetch("SELECT id FROM fragments WHERE coalesce(removed, 0) = 0 ORDER BY id")]
    if fileids:
        unknown = set(fileids) - set(all_ids)
        if unknown:
            raise ChapterError(422, f"Нет фрагментов: {', '.join(map(str, sorted(unknown)))}")
        return sorted(set(fileids)), {}
    from database.failure_forecast import ENDS_SQL

    candidates = [f for f in all_ids if f not in scen]
    rows = await conn.fetch(f"""SELECT l.id, l.fileid, {ENDS_SQL} AS ends FROM linesobj l
                                 WHERE l.shape IS NOT NULL AND coalesce(l.removed, 0) = 0
                                   AND l.fileid = ANY($1::int[])""", candidates)
    copies = copy_fragments([dict(r) for r in rows])
    return [f for f in candidates if f not in copies], copies


async def _network(conn, fileids: list[int]) -> dict[str, Any]:
    rows = await conn.fetch(
        f"""SELECT l.id, l.nodeid1, l.nodeid2, l.fileid, h.diametercondit AS dn, h.pipesectlength AS len_m,
                   {', '.join(f'(h.{c} IS NOT NULL) AS f_{c}' for c, _ in PASSPORT_COLUMNS)},
                   EXISTS (SELECT 1 FROM dampers d WHERE d.lineid = l.id) AS has_valve,
                   EXISTS (SELECT 1 FROM dampers d WHERE d.lineid = l.id AND d.damperarmaturestateid = 2) AS closed
              FROM linesobj l
              LEFT JOIN LATERAL (SELECT * FROM heatpipesections hp WHERE hp.lineid = l.id
                                  ORDER BY hp.pipesectlength DESC NULLS LAST, hp.id DESC LIMIT 1) h ON true
             WHERE coalesce(l.removed, 0) = 0 AND {line_inner_sql('l')} IS NULL AND l.fileid = ANY($1::int[])""",
        fileids)
    return {"lines": [dict(r) for r in rows]}


async def _sources(conn, fileids: list[int]) -> list[dict[str, Any]]:
    """Питающие точки: источники и узлы заданного напора (вход квартальной сети из магистрали) — как в изоляции."""
    rows = await conn.fetch(
        """SELECT hs.nodeid, n.fileid, 'source' AS kind,
                  coalesce(nullif(hs.name, ''), nullif(hs.sourcename, ''), nullif(n.externalnodename, ''),
                           'Источник №' || hs.id::text) AS name
             FROM heatsources hs JOIN nodes n ON n.id = hs.nodeid
            WHERE coalesce(n.removed, 0) = 0 AND n.fileid = ANY($1::int[])
           UNION ALL
           SELECT sp.nodeid, n.fileid, 'set_pressure',
                  'ввод ' || coalesce(nullif(n.externalnodename, ''), '№ ' || sp.nodeid::text)
             FROM setpressnodes sp JOIN nodes n ON n.id = sp.nodeid
            WHERE coalesce(n.removed, 0) = 0 AND n.fileid = ANY($1::int[])""", fileids)
    return [dict(r) for r in rows]


async def _consumers(conn, fileids: list[int]) -> list[dict[str, Any]]:
    from database.outage_simulation import _Q_CONSUMERS

    nodes = {r["id"]: r["fileid"] for r in await conn.fetch(
        "SELECT id, fileid FROM nodes WHERE coalesce(removed, 0) = 0 AND fileid = ANY($1::int[])", fileids)}
    out = []
    for r in await conn.fetch(_Q_CONSUMERS, list(nodes)):
        load = sum(float(r[k] or 0) for k in ("heating_load", "ventilation_load", "hot_water_load", "technology_load"))
        out.append({"node": r["nodeid"], "fileid": nodes.get(r["nodeid"]), "load": load, "type": r["ctype"]})
    return out


async def _optional(conn, sql: str, *args: Any) -> Any:
    """Запрос к таблице, которой в базе может не быть (модуль ещё не запускали) — None вместо ошибки."""
    try:
        return await conn.fetch(sql, *args)
    except asyncpg.UndefinedTableError:
        return None


async def collect(conn, fileids: Optional[list[int]] = None) -> dict[str, Any]:
    """Данные главы по подпунктам а–к."""
    ids, copies = await _fileids(conn, fileids)
    if not ids:
        raise ChapterError(422, "Нет фрагментов сети")
    names = {r["id"]: r["name"] for r in await conn.fetch("SELECT id, name FROM fragments WHERE id = ANY($1::int[])",
                                                         ids + list(copies))}
    net = await _network(conn, ids)
    lines = net["lines"]
    sources = await _sources(conn, ids)
    consumers = await _consumers(conn, ids)
    topo = {x["id"]: (x["nodeid1"], x["nodeid2"]) for x in lines if x["nodeid1"] and x["nodeid2"]}
    closed = {x["id"] for x in lines if x["closed"]}
    g = graph_stats(topo, {s["nodeid"] for s in sources}, closed)
    km = sum(float(x["len_m"] or 0) for x in lines) / 1000
    srid = net_srid()
    sections = [
        _sec_a(lines, topo, g, km, srid, len(ids)),
        _sec_b(lines, km, sources, consumers),
        _sec_v(ids, names, lines, consumers, await conn.fetchval(
            "SELECT count(*) FROM administrativnyy_rayon") if await conn.fetchval(
            "SELECT to_regclass('public.administrativnyy_rayon') IS NOT NULL") else None),
        await _sec_g(conn, ids, names, g, sources),
        await _sec_d(conn, ids, lines, closed),
        _sec_e(g, sources, consumers, names),
        await _sec_zh(conn),
        await _sec_z(conn),
        await _sec_i(conn),
        await _sec_k(conn, ids),
    ]
    return {"generated_at": datetime.now().astimezone().isoformat(timespec="minutes"), "basis": BASIS,
            "fragments": [{"id": i, "name": names.get(i)} for i in ids], "km": round(km, 1),
            "copies": [{"id": f, "name": names.get(f), "share": v} for f, v in sorted(copies.items())],
            "lines": len(lines), "sections": sections}


def _sec_a(lines, topo, g, km, srid, n_frag) -> dict[str, Any]:
    no_topo = len(lines) - len(topo)
    text = [
        f"Сеть представлена в ГИС ТГИД участками и узлами с координатами (система координат базы — SRID {srid}); "
        f"участок связан с начальным и конечным узлом, связность хранится в базе, а не выводится из рисунка.",
        f"Состав модели: фрагментов — {fmt(n_frag)}, участков основной сети — {fmt(len(lines))} (суммарная длина по "
        f"паспорту — {fmt(km)} км), узлов — {fmt(g['nodes'])}.",
        f"Связность: связных частей сети — {fmt(g['components'])}, из них с источником или узлом заданного напора — "
        f"{fmt(g['fed_components'])}; узлов вне связи с питанием — {fmt(g['unfed_nodes'])}.",
    ]
    if no_topo:
        text.append(f"Участков без привязки к узлам (линии ГИС без топологии) — {fmt(no_topo)}: в расчёт не входят.")
    status = OK if no_topo == 0 and g["unfed_nodes"] == 0 else PARTIAL
    return section("а", status, text)


def _sec_b(lines, km, sources, consumers) -> dict[str, Any]:
    n = len(lines) or 1
    fill = [[label, fmt(100 * sum(1 for x in lines if x[f"f_{c}"]) / n, 0) + " %"] for c, label in PASSPORT_COLUMNS]
    by_dn: dict[str, list[float]] = defaultdict(lambda: [0, 0.0])
    for x in lines:
        k = diameter_class(float(x["dn"]) if x["dn"] is not None else None)
        by_dn[k][0] += 1
        by_dn[k][1] += float(x["len_m"] or 0) / 1000
    order = [diameter_class(float(b)) for b in DIAMETER_CLASSES] + [f"свыше {DIAMETER_CLASSES[-1]}", "нет данных"]
    dn_rows = [[k, fmt(int(by_dn[k][0])), fmt(by_dn[k][1])] for k in order if k in by_dn]
    real = sum(1 for c in consumers if c["type"] == "real")
    worst = min(fill, key=lambda r: float(r[1].split()[0].replace(",", ".")))
    text = [
        "Паспорт участка, узла, источника, потребителя и арматуры ведётся в карточках объектов ТГИД (поля паспорта "
        "десктопа gid8), с журналами осмотров, дефектов, ремонтов, шурфов и опрессовок.",
        f"Источников — {fmt(sum(1 for x in sources if x['kind'] == 'source'))}, узлов заданного напора — "
        f"{fmt(sum(1 for x in sources if x['kind'] != 'source'))}, потребителей — {fmt(len(consumers))} (реальных — "
        f"{fmt(real)}, обобщённых — {fmt(len(consumers) - real)}). Заполненность паспорта участков — в таблице; хуже всего заполнено поле "
        f"«{worst[0]}» ({worst[1]}).",
    ]
    return section("б", OK, text, [
        {"title": "Заполненность паспорта участков", "header": ["Поле", "Заполнено"], "rows": fill},
        {"title": "Участки по условному диаметру", "header": ["Ду, мм", "Участков", "Длина, км"], "rows": dn_rows},
    ])


def _sec_v(ids, names, lines, consumers, districts) -> dict[str, Any]:
    per_line: dict[int, list[float]] = defaultdict(lambda: [0, 0.0])
    for x in lines:
        per_line[x["fileid"]][0] += 1
        per_line[x["fileid"]][1] += float(x["len_m"] or 0) / 1000
    per_cons: dict[int, list[float]] = defaultdict(lambda: [0, 0.0])
    for c in consumers:
        per_cons[c["fileid"]][0] += 1
        per_cons[c["fileid"]][1] += c["load"]
    rows = top_rows([[names.get(f) or f"№ {f}", int(per_line[f][0]), round(per_line[f][1], 1), int(per_cons[f][0]),
                      round(per_cons[f][1], 2)] for f in ids], 2)
    text = [
        "Расчётная единица ТГИД — фрагмент сети (тепломагистраль, квартальная сеть): по фрагменту выполняются расчёт, "
        "паспорт и отчёты; состав фрагментов — в таблице.",
    ]
    if districts:
        text.append(f"Административные районы есть в справочнике базы ({fmt(districts)}), но без границ, и объекты сети "
                    "на них не ссылаются: деление по районам для балансов нужно загрузить (слой границ районов).")
    else:
        text.append("Административного деления в базе нет: для балансов по районам нужен слой границ районов.")
    return section("в", PARTIAL, text, [{
        "title": "Расчётные единицы — фрагменты сети",
        "header": ["Фрагмент", "Участков", "Длина, км", "Потребителей", "Нагрузка, Гкал/ч"],
        "rows": [[r[0], fmt(r[1]), fmt(r[2]), fmt(r[3]), fmt(r[4], 2)] for r in rows]}])


async def _sec_g(conn, ids, names, g, sources) -> dict[str, Any]:
    rows = await conn.fetch(
        """SELECT DISTINCT ON (fileid) fileid, id, name, date1, tn FROM calculation
            WHERE fileid = ANY($1::int[]) ORDER BY fileid, date1 DESC NULLS LAST, id DESC""", ids)
    text = [
        "Гидравлический и тепловой расчёт выполняет расчётный модуль ТГИД (sety — перенос расчёта десктопа gid8): "
        "тупиковые и кольцевые сети, совместная работа нескольких источников на общую сеть, насосные станции, "
        "регуляторы и дроссельные устройства, наладочный и поверочный режимы; результаты — по участкам и "
        "потребителям, протокол расчёта, сверка с эталонами десктопа.",
        f"Закольцованность модели: независимых колец — {fmt(g['rings'])}; связных частей с несколькими точками "
        f"питания (совместная работа) — {fmt(g['multi_source_components'])}.",
        f"Фрагментов с расчётом — {fmt(len(rows))} из {fmt(len(ids))}.",
    ]
    latest = sorted(rows, key=lambda r: (r["date1"] is None, -(r["date1"].timestamp() if r["date1"] else 0)))
    if len(latest) > TOP_ROWS:
        text.append(f"В таблице — {TOP_ROWS} последних расчётов.")
    status = OK if rows and len(rows) == len(ids) else (PARTIAL if rows else NO_DATA)
    return section("г", status, text, [{
        "title": "Последний расчёт по фрагментам", "header": ["Фрагмент", "Расчёт", "Дата", "Тн, °C"],
        "rows": [[names.get(r["fileid"]) or f"№ {r['fileid']}", str(r["id"]),
                  f"{r['date1']:%d.%m.%Y}" if r["date1"] else "—", fmt(r["tn"]) if r["tn"] is not None else "—"]
                 for r in latest[:TOP_ROWS]]}])


async def _sec_d(conn, ids, lines, closed) -> dict[str, Any]:
    from database import outage_consequence

    valves = sum(1 for x in lines if x["has_valve"])
    # задвижки внутренних схем узлов (камеры, ТРП, насосные) — в Астане почти все задвижки там
    chambers = await conn.fetchval(
        """SELECT count(DISTINCT l.internalnodeid) FROM dampers d JOIN linesobj l ON l.id = d.lineid
             JOIN nodes n ON n.id = l.internalnodeid
            WHERE coalesce(l.removed, 0) = 0 AND n.fileid = ANY($1::int[])""", ids) if lines_have_inner() else 0
    try:
        st = await outage_consequence.status(conn)
    except asyncpg.UndefinedTableError:      # API ещё не доводил схему этой базы
        st = {"computed": False}
    text = [
        "Переключения моделируются изменением состояния задвижек и участков (открыта / закрыта / отключён) и "
        "расчётом нового режима; изоляция аварийного участка — минимальный набор задвижек, отключаемые потребители "
        "в зоне и ниже по сети без обходного питания; переключение нагрузок между источниками — через сценарий "
        "(копию фрагмента) с расчётом и сравнением с основой.",
        f"Узлов с задвижками во внутренней схеме (камеры, ТРП, насосные) — {fmt(chambers)}, участков с "
        f"задвижками — {fmt(valves)}, нормально закрытых участков (граница работы источников) — {fmt(len(closed))}.",
    ]
    if st.get("computed"):
        text.append(f"Последствия аварии рассчитаны для всех участков города ({fmt(st['lines'])}); участков, авария "
                    f"на которых отключает потребителей и ниже по сети, — {fmt(st['lines_with_downstream'])}.")
    return section("д", OK, text)


def _sec_e(g, sources, consumers, names) -> dict[str, Any]:
    comp_of = g["component_of"]
    by_comp: dict[int, list[float]] = defaultdict(lambda: [0, 0.0])
    for c in consumers:
        r = comp_of.get(c["node"])
        if r is not None:
            by_comp[r][0] += 1
            by_comp[r][1] += c["load"]
    src_names: dict[int, list[str]] = defaultdict(list)
    for s in sources:
        r = comp_of.get(s["nodeid"])
        if r is not None:
            src_names[r].append(s["name"])
    rows = [["; ".join(sorted(src_names[r])[:4]) + (f" и ещё {len(src_names[r]) - 4}" if len(src_names[r]) > 4 else ""),
             len(src_names[r]), int(v[0]), round(v[1], 2)] for r, v in by_comp.items() if src_names.get(r)]
    unfed = [v for r, v in by_comp.items() if not src_names.get(r)]
    total = sum(c["load"] for c in consumers)
    text = [
        "Баланс — присоединённая расчётная нагрузка потребителей (отопление, вентиляция, ГВС, технология, Гкал/ч, "
        "как в карточке потребителя) по зонам действия источников: зона — связная часть сети с учётом нормально "
        "закрытых задвижек; источники и вводы (узлы заданного напора), работающие на общую сеть, — одной строкой. Фактический отпуск по расчёту и "
        "приборам учёта — в результатах расчёта и разделе телеметрии.",
        f"Всего потребителей — {fmt(len(consumers))}, нагрузка — {fmt(total, 2)} Гкал/ч.",
    ]
    if unfed:
        text.append(f"Потребителей без питания в своей части сети — {fmt(int(sum(v[0] for v in unfed)))}, их нагрузка — "
                    f"{fmt(sum(v[1] for v in unfed), 2)} Гкал/ч: проверьте связность (подпункт «а»).")
    text.append("По территориальному признаку — по фрагментам сети (подпункт «в»).")
    return section("е", PARTIAL, text, [{
        "title": "Присоединённая нагрузка по зонам источников",
        "header": ["Источники и вводы", "Точек питания", "Потребителей", "Нагрузка, Гкал/ч"],
        "rows": [[r[0], fmt(r[1]), fmt(r[2]), fmt(r[3], 2)] for r in top_rows(rows, 3)]}])


async def _sec_zh(conn) -> dict[str, Any]:
    runs = await conn.fetch(
        "SELECT id, date1 FROM calculation WHERE calc_params LIKE '%heat_losses_norm%' ORDER BY id DESC LIMIT 1")
    sensors = await _optional(conn, "SELECT count(*) AS n FROM tgid_sensor WHERE active AND kind = 'flow'")
    text = [
        "Нормативные потери через изоляцию — модуль «Нормативные теплопотери» (удельные потери участков по месяцам, "
        "отчёт по источникам); утечки — баланс расхода по зонам между расходомерами (факт против модели, ночной "
        "расход) и утечки в гидравлическом расчёте.",
    ]
    if runs:
        d = runs[0]["date1"]
        text.append(f"Последний расчёт нормативных потерь — № {runs[0]['id']}" + (f" от {d:%d.%m.%Y}." if d else "."))
    else:
        text.append("Расчёт нормативных потерь в базе не выполнялся.")
    n = sensors[0]["n"] if sensors else 0
    text.append(f"Расходомеров телеметрии для баланса утечек — {fmt(n)}.")
    status = OK if runs and n else (PARTIAL if runs or n else NO_DATA)
    return section("ж", status, text)


async def _sec_z(conn) -> dict[str, Any]:
    import json

    run = await _optional(conn, "SELECT id, created_at, population, expected_events, horizon_years, metrics "
                                "FROM tgid_failure_run ORDER BY id DESC LIMIT 1")
    text = [
        "Надёжность в ТГИД — вероятностный прогноз повреждений участков по журналу дефектов (проверка на годах, "
        "которых модель не видела) и последствия аварии каждого участка: отключаемые потребители, нагрузка и "
        "социальные объекты; по ним — план перекладок. Нормативные показатели надёжности по методическим указаниям "
        "к схемам теплоснабжения ТГИД пока не рассчитывает.",
    ]
    if run:
        r = run[0]
        m = json.loads(r["metrics"]) if isinstance(r["metrics"], str) else r["metrics"]
        text.append(f"Прогноз № {r['id']} от {r['created_at']:%d.%m.%Y}: участков в расчёте — {fmt(r['population'])}, "
                    f"ожидаемых повреждений за {r['horizon_years']} г. — {fmt(float(r['expected_events']))}; ROC-AUC на "
                    f"проверке — {fmt(m.get('auc'), 3)}.")
        return section("з", PARTIAL, text)
    text.append("Прогноз аварийности в базе не рассчитывался (нужен журнал повреждений с участком и датой).")
    return section("з", NO_DATA, text)


async def _sec_i(conn) -> dict[str, Any]:
    scen = await _optional(conn, "SELECT count(*) AS n FROM tgid_scenario WHERE status <> 'archived'")
    n = scen[0]["n"] if scen else 0
    text = [
        "Групповое изменение характеристик — «Групповые установщики»: значение поля для выделенных на карте объектов, "
        "фрагмента или фильтра по условиям (диаметр, материал, состояние, нагрузки), как в десктопе; перспективные "
        "варианты — сценарии (копии фрагментов) с переносом правок в основу.",
        f"Сценариев в работе — {fmt(n)}.",
    ]
    return section("и", OK, text)


async def _sec_k(conn, ids) -> dict[str, Any]:
    n = await conn.fetchval("SELECT count(*) FROM calculation WHERE fileid = ANY($1::int[])", ids)
    text = [
        "Пьезометрический график строится по пути между выбранными узлами по результатам расчёта (напоры в подаче и "
        "обратке, отметки земли, статический напор) с выгрузкой в Excel; сценарий сравнивается с основой по "
        "потребителям и участкам (напоры, располагаемый напор, расходы).",
        f"Расчётов фрагментов в базе — {fmt(n)}.",
    ]
    return section("к", OK if n else NO_DATA, text)


# --- Word ----------------------------------------------------------------------------------------

def to_docx(data: dict[str, Any], city: str = "") -> bytes:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)
    title = doc.add_heading("Глава 3. Электронная модель системы теплоснабжения" + (f" — {city}" if city else ""), 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph(
        f"Черновик по данным ГИС ТГИД на {data['generated_at'][:16].replace('T', ' ')}. Состав — {data['basis']} "
        "(подпункты «а»–«к»). По каждому подпункту — как требование выполняет ТГИД и состояние данных; «нет данных» "
        "и «частично» требуют доработки разработчиком схемы.")
    doc.add_paragraph(f"В модели: фрагментов сети — {len(data['fragments'])}, участков — {fmt(data['lines'])}, "
                      f"суммарная длина — {fmt(data['km'])} км.")
    if data.get("copies"):
        doc.add_paragraph("Не вошли — копии более новых фрагментов (доля совпадающих участков): " + "; ".join(
            f"{c['name'] or c['id']} ({fmt(100 * c['share'], 0)} %)" for c in data["copies"]) + ".")
    doc.add_heading("Соответствие требованиям", 1)
    t = doc.add_table(rows=1, cols=3)
    t.style = "Table Grid"
    for cell, text in zip(t.rows[0].cells, ("Подпункт", "Требование", "Состояние")):
        cell.text = text
        cell.paragraphs[0].runs[0].bold = True
    for s in data["sections"]:
        row = t.add_row().cells
        row[0].text, row[1].text, row[2].text = f"«{s['key']}»", s["title"], s["status"]
    for s in data["sections"]:
        doc.add_heading(f"3.{list(ITEMS).index(s['key']) + 1}. {s['title']} (подпункт «{s['key']}»)", 1)
        doc.add_paragraph(f"Состояние: {s['status']}.").runs[0].italic = True
        for p in s["text"]:
            doc.add_paragraph(p)
        for tab in s["tables"]:
            doc.add_paragraph(tab["title"]).runs[0].bold = True
            tt = doc.add_table(rows=1, cols=len(tab["header"]))
            tt.style = "Table Grid"
            for cell, text in zip(tt.rows[0].cells, tab["header"]):
                cell.text = text
                cell.paragraphs[0].runs[0].bold = True
            for r in tab["rows"]:
                cells = tt.add_row().cells
                for cell, v in zip(cells, r):
                    cell.text = fmt(v) if not isinstance(v, str) else v
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
