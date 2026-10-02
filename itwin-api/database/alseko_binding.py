"""АЛСЕКО: сверка договорных нагрузок со зданиями и привязка (этап 10).

Эталон десктопа (gid8/gid8/alseco, gid6/gidr/postgresql/sql/alseco):

* сверка — отчёты «ненайденные МЖД» (nenaid1), «ненайденные прочие» (nenaid2),
  «здания с нагрузкой без потребителя» (nenaid3); сопоставление ``nagruzki`` ↔ ``zdaniya_2``
  по микрорайону, улице и дому без пробелов/регистра;
* привязка здания к адресу АЛСЕКО — диалог BigDialog (``fun_alseco_nagr``): в здании
  записываются mkr2/street2/house2, суммы нагрузок адреса (ккал/ч → Гкал/ч, /1e6), nagr, txt;
* привязка зданий к потребителю — ``GidWidget::alseco`` (карточка обобщённого потребителя/ТП):
  ``zdaniya_2.potrebitel`` = «код узел» (externalcodes.name + nodes.externalnodename), старая привязка потребителя снимается;
  суммы нагрузок выбранных зданий по схемам — как ``readAlseco``.

Дополнительно к десктопу сверка показывает неоднозначные адреса, устаревшие привязки
и расхождение записанных в здании нагрузок с текущими ``nagruzki``.
Запись — одна транзакция + audit_log (роутер проверяет editor+ и MUTATIONS_ENABLED).
"""

from __future__ import annotations

import io
import math
from typing import Any, Awaitable, Callable, Optional

import asyncpg

KCAL_TO_GCAL = 1e6
LOAD_TOLERANCE = 1e-6
MAX_CONSUMER_BUILDINGS = 500
EXCEL_ROW_LIMIT = 100_000

AuditRow = Callable[..., Awaitable[str]]


class AlsekoError(Exception):
    def __init__(self, status: int, code: str, message: str, **extra: Any):
        super().__init__(message)
        self.status = status
        self.detail = {"code": code, "message": message, **extra}


# Нормализованный адресный ключ (как в десктопе, плюс '' == NULL для микрорайона/улицы).
N_KEY = "coalesce(n.mkr, ''), coalesce(n.street, ''), lower(replace(coalesce(n.house, ''), ' ', ''))"
Z_KEY = "coalesce(z.mkr2, ''), coalesce(z.street2, ''), lower(replace(coalesce(z.house2, ''), ' ', ''))"
N_KEY_AS = ("coalesce(n.mkr, '') AS k_mkr, coalesce(n.street, '') AS k_street,"
            " lower(replace(coalesce(n.house, ''), ' ', '')) AS k_house")
Z_KEY_AS = ("coalesce(z.mkr2, '') AS k_mkr, coalesce(z.street2, '') AS k_street,"
            " lower(replace(coalesce(z.house2, ''), ' ', '')) AS k_house")

ADDRESS_CTE = f"""
    WITH load_addr AS MATERIALIZED (
        SELECT {N_KEY_AS},
               count(*)::int AS load_count,
               sum(coalesce(n.otop, 0)) AS otop, sum(coalesce(n.gvs, 0)) AS gvs,
               sum(coalesce(n.vent, 0)) AS vent, sum(coalesce(n.par, 0)) AS par
          FROM nagruzki n
         WHERE coalesce(n.house, '') <> ''
         GROUP BY 1, 2, 3
    ), bld_addr AS MATERIALIZED (
        SELECT {Z_KEY_AS},
               count(*)::int AS building_count,
               array_agg(z.id ORDER BY z.id) AS building_ids
          FROM zdaniya_2 z
         WHERE coalesce(z.house2, '') <> ''
         GROUP BY 1, 2, 3
    )
"""

_LOAD_COLUMNS = """
    n.id, n.addr AS source_address, n.mkr AS microdistrict, n.street, n.house,
    n.name AS customer_type, n.owner, n.dogovor AS contract_number, n.numb AS registry_number,
    n.rayon AS operation_district, n.adm_rayon AS administrative_district, n.uchastok AS operation_site,
    n.ist AS heat_source, n.tg AS temperature_graph,
    n.otop AS heating_load, n.gvs AS hot_water_load, n.vent AS ventilation_load, n.par AS steam_load
"""

_UNMATCHED = ADDRESS_CTE + f"""
    SELECT {_LOAD_COLUMNS}
      FROM nagruzki n
      LEFT JOIN bld_addr b ON (b.k_mkr, b.k_street, b.k_house) = ({N_KEY})
     WHERE b.k_house IS NULL AND {{group}}
"""

_BUILDING_COLUMNS = """
    z.id, z.id_adr_mas AS geo_microdistrict, z.street_nam AS geo_street, z.number_1 AS geo_house,
    z.mkr2 AS microdistrict, z.street2 AS street, z.house2 AS house,
    z.otop AS heating_load, z.gvs AS hot_water_load, z.vent AS ventilation_load, z.par AS steam_load,
    z.nagr AS total_load, z.potrebitel AS consumer
"""

CONSUMER_LABELS_CTE = """
    WITH consumer_labels AS MATERIALIZED (
        SELECT DISTINCT ec.name || ' ' || nd.externalnodename AS label
          FROM nodes nd
          JOIN externalcodes ec ON ec.id = nd.externalcodeid
         WHERE nd.externalnodename IS NOT NULL
    )
"""

ISSUE_KINDS: dict[str, dict[str, Any]] = {
    "unmatched_apartment": {
        "label": "Ненайденные МЖД (нагрузка без здания)",
        "desktop": "nenaid1.sql",
        "sql": _UNMATCHED.replace("{group}", "n.name = 'МЖД'"),
        "order": "n.rayon NULLS LAST, n.street NULLS LAST, n.house NULLS LAST, n.id",
    },
    "unmatched_other": {
        "label": "Ненайденные прочие объекты",
        "desktop": "nenaid2.sql",
        "sql": _UNMATCHED.replace("{group}", "n.name IS DISTINCT FROM 'МЖД'"),
        "order": "n.rayon NULLS LAST, n.owner NULLS LAST, n.id",
    },
    "unassigned_building": {
        "label": "Здания с нагрузкой без потребителя",
        "desktop": "nenaid3.sql",
        "sql": f"SELECT {_BUILDING_COLUMNS} FROM zdaniya_2 z WHERE z.potrebitel IS NULL AND z.otop IS NOT NULL",
        "order": "z.mkr2 NULLS LAST, z.street2 NULLS LAST, z.house2 NULLS LAST, z.id",
    },
    "ambiguous_address": {
        "label": "Адрес АЛСЕКО привязан к нескольким зданиям",
        "desktop": None,
        "sql": ADDRESS_CTE + """
            SELECT nullif(l.k_mkr, '') AS microdistrict, nullif(l.k_street, '') AS street, l.k_house AS house,
                   l.load_count, b.building_count, b.building_ids,
                   (l.otop + l.gvs + l.vent + l.par) / 1e6 AS total_load_gcal
              FROM load_addr l
              JOIN bld_addr b USING (k_mkr, k_street, k_house)
             WHERE b.building_count > 1
        """,
        "order": "b.building_count DESC, l.k_street, l.k_house",
    },
    "stale_address": {
        "label": "Здание привязано к адресу, которого нет в nagruzki",
        "desktop": None,
        "sql": ADDRESS_CTE + f"""
            SELECT {_BUILDING_COLUMNS}
              FROM zdaniya_2 z
              LEFT JOIN load_addr l ON (l.k_mkr, l.k_street, l.k_house) = ({Z_KEY})
             WHERE coalesce(z.house2, '') <> '' AND l.k_house IS NULL
        """,
        "order": "z.street2 NULLS LAST, z.house2, z.id",
    },
    "load_mismatch": {
        "label": "Нагрузки здания расходятся с nagruzki",
        "desktop": None,
        "sql": ADDRESS_CTE + f"""
            SELECT {_BUILDING_COLUMNS},
                   l.otop / 1e6 AS source_heating_load, l.gvs / 1e6 AS source_hot_water_load,
                   l.vent / 1e6 AS source_ventilation_load, l.par / 1e6 AS source_steam_load
              FROM zdaniya_2 z
              JOIN load_addr l ON (l.k_mkr, l.k_street, l.k_house) = ({Z_KEY})
             WHERE abs(coalesce(z.otop, 0) - l.otop / 1e6) > {LOAD_TOLERANCE}
                OR abs(coalesce(z.gvs, 0) - l.gvs / 1e6) > {LOAD_TOLERANCE}
                OR abs(coalesce(z.vent, 0) - l.vent / 1e6) > {LOAD_TOLERANCE}
                OR abs(coalesce(z.par, 0) - l.par / 1e6) > {LOAD_TOLERANCE}
        """,
        "order": "z.street2 NULLS LAST, z.house2, z.id",
    },
    "unknown_consumer": {
        "label": "Потребитель здания («код узел») не найден среди узлов сети",
        "desktop": None,
        "sql": CONSUMER_LABELS_CTE + f"""
            SELECT {_BUILDING_COLUMNS}
              FROM zdaniya_2 z
              LEFT JOIN consumer_labels c ON c.label = z.potrebitel
             WHERE z.potrebitel IS NOT NULL AND c.label IS NULL
        """,
        "order": "z.potrebitel, z.id",
    },
}


def _kind(kind: str) -> dict[str, Any]:
    spec = ISSUE_KINDS.get(kind)
    if spec is None:
        raise AlsekoError(422, "bad_kind", f"Неизвестный вид несоответствия: {kind}", kinds=list(ISSUE_KINDS))
    return spec


def _clean(value: Any) -> Any:
    if isinstance(value, float) and not math.isfinite(value):
        return None
    return value


def _row(record: asyncpg.Record) -> dict[str, Any]:
    return {k: _clean(v) for k, v in dict(record).items()}


async def reconciliation_summary(conn: asyncpg.Connection) -> dict[str, Any]:
    counts = {}
    for kind, spec in ISSUE_KINDS.items():
        counts[kind] = await conn.fetchval(f"SELECT count(*) FROM ({spec['sql']}) q")
    totals = await conn.fetchrow(
        "SELECT (SELECT count(*) FROM nagruzki) AS loads,"
        " (SELECT count(*) FROM zdaniya_2 WHERE coalesce(house2, '') <> '') AS bound_buildings,"
        " (SELECT count(*) FROM zdaniya_2 WHERE potrebitel IS NOT NULL) AS buildings_with_consumer"
    )
    return {
        "totals": dict(totals),
        "kinds": [
            {"kind": k, "label": s["label"], "desktop_report": s["desktop"], "count": counts[k]}
            for k, s in ISSUE_KINDS.items()
        ],
    }


async def reconciliation_issues(
    conn: asyncpg.Connection, kind: str, *, limit: int = 100, offset: int = 0
) -> dict[str, Any]:
    spec = _kind(kind)
    total = await conn.fetchval(f"SELECT count(*) FROM ({spec['sql']}) q")
    rows = await conn.fetch(f"{spec['sql']} ORDER BY {spec['order']} LIMIT $1 OFFSET $2", limit, offset)
    return {"kind": kind, "label": spec["label"], "total": total, "items": [_row(r) for r in rows]}


# --- адреса АЛСЕКО для выбора ----------------------------------------------------

async def address_candidates(
    conn: asyncpg.Connection, *, q: Optional[str], building_id: Optional[int], limit: int = 50
) -> dict[str, Any]:
    """Адреса nagruzki с суммами (Гкал/ч). Без q, но с building_id — подсказка по геоадресу здания
    (как BigDialog: id_adr_mas, street_nam, number_1)."""
    building = None
    terms: list[str] = []
    if building_id is not None:
        building = await conn.fetchrow(f"SELECT {_BUILDING_COLUMNS} FROM zdaniya_2 z WHERE z.id=$1", building_id)
        if building is None:
            raise AlsekoError(404, "not_found", f"Здание {building_id} не найдено")
    text = (q or "").strip()
    if text:
        terms = [t for t in text.split() if t]
    elif building is not None:
        street = (building["geo_street"] or "").strip()
        house = (building["geo_house"] or "").strip()
        terms = [t for t in (street, house) if t]
    where = []
    values: list[Any] = []
    for term in terms[:5]:
        values.append(term)
        p = f"${len(values)}"
        where.append(
            f"(coalesce(n.mkr, '') || ' ' || coalesce(n.street, '') || ' ' || coalesce(n.house, '')"
            f" || ' ' || coalesce(n.addr, '')) ILIKE '%' || {p} || '%'"
        )
    values.append(limit)
    sql = ADDRESS_CTE + f"""
        SELECT nullif(k_mkr, '') AS microdistrict, nullif(k_street, '') AS street, house, load_count,
               otop / 1e6 AS heating_load, gvs / 1e6 AS hot_water_load, vent / 1e6 AS ventilation_load,
               par / 1e6 AS steam_load, (otop + gvs + vent + par) / 1e6 AS total_load,
               coalesce(b.building_count, 0) AS building_count, b.building_ids
          FROM (
            SELECT {N_KEY_AS},
                   min(n.house) AS house, count(*)::int AS load_count,
                   sum(coalesce(n.otop, 0)) AS otop, sum(coalesce(n.gvs, 0)) AS gvs,
                   sum(coalesce(n.vent, 0)) AS vent, sum(coalesce(n.par, 0)) AS par
              FROM nagruzki n
             WHERE coalesce(n.house, '') <> '' {"AND " + " AND ".join(where) if where else ""}
             GROUP BY 1, 2, 3
          ) a
          LEFT JOIN bld_addr b USING (k_mkr, k_street, k_house)
         ORDER BY k_street, k_house, k_mkr
         LIMIT ${len(values)}
    """
    rows = await conn.fetch(sql, *values)
    return {
        "building": _row(building) if building is not None else None,
        "query": " ".join(terms),
        "items": [_row(r) for r in rows],
    }


# --- привязка здания к адресу (BigDialog) -----------------------------------------

def alseco_text(mkr: str, street: str, house: str, otop: float, vent: float, gvs: float, par: float) -> str:
    """Подпись здания — как getAlsecoTxt (нагрузки в ккал/ч, в тексте Гкал/ч)."""
    parts = [p for p in (mkr, street, house) if p]
    lines = [" ".join(parts)] if parts else []
    for label, value in (("Qот", otop), ("Qгвс", gvs), ("Qвент", vent), ("Qпар", par)):
        if value:
            lines.append(f"{label}={value / KCAL_TO_GCAL:g}")
    total = (otop + gvs + vent + par) / KCAL_TO_GCAL
    if total:
        lines.append(f"Qсум={total:g}")
    return "\r\n".join(lines)


BUILDING_BIND_FIELDS = ("mkr2", "street2", "house2", "otop", "gvs", "vent", "par", "nagr", "txt")


async def _building_state(conn: asyncpg.Connection, building_id: int, *, lock: bool) -> dict[str, Any]:
    row = await conn.fetchrow(
        f"SELECT id, {', '.join(BUILDING_BIND_FIELDS)}, potrebitel FROM zdaniya_2 WHERE id=$1"
        + (" FOR UPDATE" if lock else ""),
        building_id,
    )
    if row is None:
        raise AlsekoError(404, "not_found", f"Здание {building_id} не найдено")
    return _row(row)


async def _address_loads(conn: asyncpg.Connection, mkr: str, street: str, house: str) -> dict[str, Any]:
    row = await conn.fetchrow(
        f"""SELECT count(*)::int AS load_count, min(n.mkr) AS mkr, min(n.street) AS street, min(n.house) AS house,
                   sum(coalesce(n.otop, 0)) AS otop, sum(coalesce(n.gvs, 0)) AS gvs,
                   sum(coalesce(n.vent, 0)) AS vent, sum(coalesce(n.par, 0)) AS par
              FROM nagruzki n
             WHERE ({N_KEY}) = ($1, $2, lower(replace($3, ' ', '')))""",
        mkr or "", street or "", house or "",
    )
    return dict(row)


def _changes(old: dict[str, Any], new: dict[str, Any]) -> list[dict[str, Any]]:
    out = []
    for field, value in new.items():
        before = old.get(field)
        same = before == value or (
            isinstance(before, (int, float)) and isinstance(value, (int, float)) and abs(before - value) <= 1e-12
        )
        if not same:
            out.append({"field": field, "old": before, "new": value})
    return out


async def bind_building_address(
    conn: asyncpg.Connection,
    building_id: int,
    *,
    mkr: Optional[str],
    street: Optional[str],
    house: Optional[str],
    dry_run: bool,
    audit_row: Optional[AuditRow] = None,
) -> dict[str, Any]:
    """Записать в здание адрес АЛСЕКО и суммы нагрузок адреса (clear — house пустой)."""
    old = await _building_state(conn, building_id, lock=not dry_run)
    clearing = not (house or "").strip()
    if clearing:
        new: dict[str, Any] = {f: None for f in BUILDING_BIND_FIELDS}
        source = None
    else:
        source = await _address_loads(conn, mkr or "", street or "", house or "")
        if not source["load_count"]:
            raise AlsekoError(422, "address_not_found", "Адрес не найден в nagruzki (АЛСЕКО)")
        o, g, v, p = (float(source[k] or 0) for k in ("otop", "gvs", "vent", "par"))
        new = {
            "mkr2": source["mkr"], "street2": source["street"], "house2": source["house"],
            "otop": o / KCAL_TO_GCAL, "gvs": g / KCAL_TO_GCAL, "vent": v / KCAL_TO_GCAL, "par": p / KCAL_TO_GCAL,
            "nagr": (o + g + v + p) / KCAL_TO_GCAL,
            "txt": alseco_text(source["mkr"] or "", source["street"] or "", source["house"] or "", o, v, g, p),
        }
    changes = _changes(old, new)
    others = []
    if not clearing:
        others = [
            r["id"] for r in await conn.fetch(
                f"SELECT z.id FROM zdaniya_2 z WHERE ({Z_KEY}) = ($1, $2, lower(replace($3, ' ', ''))) AND z.id <> $4"
                " ORDER BY z.id LIMIT 20",
                new["mkr2"] or "", new["street2"] or "", new["house2"] or "", building_id,
            )
        ]
    result = {
        "building_id": building_id,
        "dry_run": dry_run,
        "action": "clear" if clearing else "bind",
        "source": source,
        "changes": changes,
        "other_buildings_with_address": others,
        "change_group_id": None,
    }
    if dry_run or not changes:
        return result
    fields = list(new)
    sets = ", ".join(f"{f}=${i + 2}" for i, f in enumerate(fields))
    await conn.execute(f"UPDATE zdaniya_2 SET {sets} WHERE id=$1", building_id, *[new[f] for f in fields])
    if audit_row is not None:
        result["change_group_id"] = await audit_row(
            conn, operation="UPDATE", table="zdaniya_2", record_id=building_id,
            old={c["field"]: c["old"] for c in changes}, new={c["field"]: c["new"] for c in changes},
        )
    return result


# --- привязка зданий к потребителю (GidWidget::alseco) ---------------------------------

SCHEME_SUMS_SQL = """
    SELECT count(*)::int AS buildings,
           sum(otop * (otop_cxema = 1)::int) AS heating_dependent_elevator,
           sum(otop * (otop_cxema = 2)::int) AS heating_dependent_direct,
           sum(otop * (otop_cxema = 3)::int) AS heating_independent,
           sum(otop * (otop_cxema IS NULL OR otop_cxema NOT IN (1, 2, 3))::int) AS heating_no_scheme,
           sum(gvs * (gvs_cxema = 1)::int) AS hot_water_open_supply,
           sum(gvs * (gvs_cxema = 2)::int) AS hot_water_open_return,
           sum(gvs * (gvs_cxema = 3)::int) AS hot_water_closed_parallel,
           sum(gvs * (gvs_cxema = 4)::int) AS hot_water_closed_mixed,
           sum(gvs * (gvs_cxema = 5)::int) AS hot_water_closed_serial,
           sum(gvs * (gvs_cxema = 6)::int) AS hot_water_closed_preheat,
           sum(gvs * (gvs_cxema IS NULL OR gvs_cxema NOT BETWEEN 1 AND 6)::int) AS hot_water_no_scheme,
           sum(vent) AS ventilation, sum(par) AS steam, sum(nagr) AS total
      FROM zdaniya_2 WHERE id = ANY($1::int[])
"""


async def consumer_label(conn: asyncpg.Connection, node_id: int) -> dict[str, Any]:
    row = await conn.fetchrow(
        """SELECT nd.id AS node_id, nd.nodetypeid AS node_type_id, ec.name AS code, nd.externalnodename AS node_name,
                  (SELECT gc.id FROM generalizedconsumers gc WHERE gc.nodeid = nd.id LIMIT 1) AS consumer_id
             FROM nodes nd
             LEFT JOIN externalcodes ec ON ec.id = nd.externalcodeid
            WHERE nd.id = $1""",
        node_id,
    )
    if row is None:
        raise AlsekoError(404, "not_found", f"Узел {node_id} не найден")
    if not row["code"] or not row["node_name"]:
        raise AlsekoError(422, "no_label", "У узла потребителя нет кода или имени — привязка невозможна")
    info = dict(row)
    info["label"] = f"{row['code']} {row['node_name']}"
    return info


async def bind_consumer_buildings(
    conn: asyncpg.Connection,
    node_id: int,
    building_ids: list[int],
    *,
    dry_run: bool,
    audit_row: Optional[AuditRow] = None,
) -> dict[str, Any]:
    """Назначить зданиям потребителя; прежние здания этого потребителя отвязываются."""
    ids = sorted({int(i) for i in building_ids})
    if len(ids) > MAX_CONSUMER_BUILDINGS:
        raise AlsekoError(422, "too_many", f"Не больше {MAX_CONSUMER_BUILDINGS} зданий за раз")
    consumer = await consumer_label(conn, node_id)
    if consumer.get("consumer_id") is None:
        # gid6: кнопка «Нагрузки АЛСЕКО» (alsecoNagr) есть только в карточке generalizedConsumers (QA F47)
        raise AlsekoError(
            422, "not_consumer",
            f"Узел {node_id} ({consumer['label']}) не является потребителем: у него нет карточки "
            "обобщённого потребителя. Здания АЛСЕКО привязываются только к узлу-потребителю.",
            node_type_id=consumer.get("node_type_id"),
        )
    label = consumer["label"]
    lock = "" if dry_run else " FOR UPDATE"
    found = await conn.fetch(f"SELECT id, potrebitel FROM zdaniya_2 WHERE id = ANY($1::int[]){lock}", ids)
    missing = sorted(set(ids) - {r["id"] for r in found})
    if missing:
        raise AlsekoError(404, "not_found", "Здания не найдены", missing=missing[:50])
    previous = await conn.fetch(
        f"SELECT id, potrebitel FROM zdaniya_2 WHERE potrebitel = $1 AND NOT (id = ANY($2::int[])){lock}", label, ids
    )
    assign = [dict(r) for r in found if r["potrebitel"] != label]
    taken = [dict(r) for r in assign if r["potrebitel"]]
    sums = _row(await conn.fetchrow(SCHEME_SUMS_SQL, ids)) if ids else {}
    dependent = (sums.get("heating_dependent_elevator") or 0) + (sums.get("heating_dependent_direct") or 0)
    warnings = []
    if dependent > 0 and (sums.get("heating_independent") or 0) > 0:
        warnings.append("Выбранные здания имеют разные схемы отопления (зависимая и независимая)")
    if taken:
        warnings.append(f"{len(taken)} зданий уже привязаны к другому потребителю — привязка будет заменена")
    result = {
        "consumer": consumer,
        "dry_run": dry_run,
        "assign": assign,
        "unassign": [dict(r) for r in previous],
        "loads": sums,
        "warnings": warnings,
        "change_group_id": None,
    }
    if dry_run or (not assign and not previous):
        return result
    group = None
    for r in previous:
        await conn.execute("UPDATE zdaniya_2 SET potrebitel=NULL WHERE id=$1", r["id"])
        if audit_row is not None:
            group = await audit_row(conn, operation="UPDATE", table="zdaniya_2", record_id=r["id"],
                                    old={"potrebitel": r["potrebitel"]}, new={"potrebitel": None}, group=group)
    for r in assign:
        await conn.execute("UPDATE zdaniya_2 SET potrebitel=$2 WHERE id=$1", r["id"], label)
        if audit_row is not None:
            group = await audit_row(conn, operation="UPDATE", table="zdaniya_2", record_id=r["id"],
                                    old={"potrebitel": r["potrebitel"]}, new={"potrebitel": label}, group=group)
    result["change_group_id"] = group
    return result


# --- Excel: отчёт несоответствий -----------------------------------------------------

async def reconciliation_workbook(conn: asyncpg.Connection, kinds: Optional[list[str]] = None) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    selected = kinds or list(ISSUE_KINDS)
    for k in selected:
        _kind(k)
    wb = Workbook()
    summary = wb.active
    summary.title = "Сводка"
    summary.append(["Вид несоответствия", "Количество", "Отчёт десктопа"])
    for cell in summary[1]:
        cell.font = Font(bold=True)
    for index, kind in enumerate(selected, start=1):
        spec = ISSUE_KINDS[kind]
        rows = await conn.fetch(f"{spec['sql']} ORDER BY {spec['order']} LIMIT $1", EXCEL_ROW_LIMIT + 1)
        summary.append([spec["label"], len(rows) if len(rows) <= EXCEL_ROW_LIMIT else f">{EXCEL_ROW_LIMIT}",
                        spec["desktop"] or "—"])
        ws = wb.create_sheet(f"{index}. {kind}"[:31])
        ws.append([spec["label"]])
        ws["A1"].font = Font(bold=True)
        if not rows:
            ws.append(["Несоответствий нет"])
            continue
        columns = list(rows[0].keys())
        ws.append(columns)
        for cell in ws[2]:
            cell.font = Font(bold=True)
        for record in rows[:EXCEL_ROW_LIMIT]:
            ws.append([
                ", ".join(map(str, v)) if isinstance(v, list) else _clean(v) for v in record.values()
            ])
        ws.freeze_panes = "A3"
    summary.column_dimensions["A"].width = 60
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
