"""Плановые и аварийные отключения (фаза 2): изоляция → задвижки и потребители → уведомления → карта.

Отключение — запись tgid_outage (sql/migrations/20261030_outages.sql): что отключаем (участок или узел),
сроки, причина, статус и снимок расчёта изоляции (database/outage_simulation.py) на момент планирования —
задвижки, потребители зоны и потерявшие питание ниже по сети, зона на карте. Снимок пересчитывается
кнопкой, пока отключение не началось. Пересечения: другие открытые отключения с общими потребителями и
перекрывающимся временем. Уведомления — список потребителей (Excel) и текст объявления; рассылкой не
занимается. Публичная лента — только объявленные и идущие, у которых включён public: зона и число потребителей,
без названий (routers/outages.py, PUBLIC_OUTAGES_ENABLED).
"""
from __future__ import annotations

import io
import json
from datetime import datetime, timezone
from typing import Any, Literal, Optional

from pydantic import BaseModel, Field

from database import poi
from database.outage_simulation import simulate_outage_isolation

KINDS: dict[str, str] = {"planned": "Плановое", "emergency": "Аварийное"}
STATUSES: dict[str, str] = {
    "planned": "Запланировано", "announced": "Объявлено", "active": "Идёт", "finished": "Завершено",
    "cancelled": "Отменено",
}
TRANSITIONS: dict[str, tuple[str, ...]] = {
    "planned": ("announced", "active", "cancelled"),
    "announced": ("active", "planned", "cancelled"),
    "active": ("finished",),
    "finished": (),
    "cancelled": ("planned",),
}
OPEN = ("planned", "announced", "active")
RECOMPUTABLE = ("planned", "announced")
PUBLIC_STATUSES = ("announced", "active")
# зона на карте: буфер вокруг отключаемых участков (единицы системы координат сети — метры) и упрощение
AREA_BUFFER = 40
AREA_SIMPLIFY = 10
ANNOUNCE_NAMES = 20
MAX_LIST = 500

Status = Literal["planned", "announced", "active", "finished", "cancelled"]


class OutageError(Exception):
    def __init__(self, status: int, message: str):
        super().__init__(message)
        self.status = status
        self.message = message


class OutageIn(BaseModel):
    title: str = Field(..., min_length=1, max_length=300)
    kind: Literal["planned", "emergency"] = "planned"
    line_id: Optional[int] = Field(None, ge=1)
    node_id: Optional[int] = Field(None, ge=1)
    work_order_id: Optional[int] = Field(None, ge=1)
    planned_start: datetime
    planned_end: datetime
    reason: str = Field("", max_length=2000)
    public: bool = False
    public_note: str = Field("", max_length=1000)


class OutagePatch(BaseModel):
    revision: int = Field(..., ge=1)
    title: Optional[str] = Field(None, min_length=1, max_length=300)
    planned_start: Optional[datetime] = None
    planned_end: Optional[datetime] = None
    reason: Optional[str] = Field(None, max_length=2000)
    public: Optional[bool] = None
    public_note: Optional[str] = Field(None, max_length=1000)


class OutageTransition(BaseModel):
    to: Status
    revision: int = Field(..., ge=1)
    note: str = Field("", max_length=2000)


def _aware(dt: Optional[datetime]) -> Optional[datetime]:
    return dt.replace(tzinfo=timezone.utc) if dt is not None and dt.tzinfo is None else dt


def _iso(v: Any) -> Any:
    return v.isoformat() if isinstance(v, datetime) else v


def _json(v: Any) -> Any:
    return json.loads(v) if isinstance(v, str) else v


def _row(r, *, full: bool = True) -> dict[str, Any]:
    out = {k: _iso(v) for k, v in dict(r).items()}
    for key in ("valves", "consumers", "area"):
        if key in out:
            out[key] = _json(out[key])
    if not full:
        for key in ("valves", "consumers", "area"):
            out.pop(key, None)
    out["status_label"] = STATUSES.get(out["status"], out["status"])
    out["kind_label"] = KINDS.get(out["kind"], out["kind"])
    out["next"] = list(TRANSITIONS.get(out["status"], ()))
    out["recomputable"] = out["status"] in RECOMPUTABLE
    return out


# --- снимок изоляции ------------------------------------------------------------------------

async def _snapshot(conn, line_id: Optional[int], node_id: Optional[int]) -> dict[str, Any]:
    """Расчёт изоляции → задвижки, потребители (зоны и ниже по сети), нагрузка, зона и точка на карте."""
    try:
        iso = await simulate_outage_isolation(conn, line_id=line_id, node_id=node_id)
    except ValueError as exc:
        raise OutageError(422, str(exc)) from exc
    valves = [{
        "id": v["id"], "name": v.get("display_name"), "line_id": v.get("lineid"), "required": bool(v.get("required")),
        "diameter": v.get("nominal_diameter"), "state": v.get("state_name"), "lng": v.get("lng"), "lat": v.get("lat"),
    } for v in iso["valves_to_close"]]
    valves.sort(key=lambda v: (not v["required"], v["id"]))

    def consumer(c: dict[str, Any], downstream: bool) -> dict[str, Any]:
        return {"id": c["id"], "type": c["consumer_type"], "node_id": c["node_id"], "name": c["name"],
                "load": c["total_load"], "lng": c["longitude"], "lat": c["latitude"], "downstream": downstream}

    consumers = [consumer(c, False) for c in iso["affected_consumers"]]
    consumers += [consumer(c, True) for c in iso["downstream_consumers"]]
    gj = iso["geojson"]
    line_ids = [f["properties"]["id"] for f in gj["isolated_pipes"]["features"] + gj["downstream_pipes"]["features"]]
    area = None
    if line_ids:
        raw = await conn.fetchval(
            f"""SELECT ST_AsGeoJSON(ST_Transform(ST_SimplifyPreserveTopology(
                       ST_Buffer(ST_Union(shape), {AREA_BUFFER}), {AREA_SIMPLIFY}), 4326), 6)
                  FROM linesobj WHERE id = ANY($1::int[]) AND shape IS NOT NULL""", line_ids)
        area = json.loads(raw) if raw else None
    if line_id:
        place = await conn.fetchrow(
            "SELECT fileid, ST_X(ST_Transform(ST_PointOnSurface(shape), 4326)) AS lon, "
            "ST_Y(ST_Transform(ST_PointOnSurface(shape), 4326)) AS lat FROM linesobj WHERE id = $1", line_id)
    else:
        place = await conn.fetchrow(
            "SELECT fileid, ST_X(ST_Transform(shape, 4326)) AS lon, ST_Y(ST_Transform(shape, 4326)) AS lat "
            "FROM nodes WHERE id = $1", node_id)
    summary = iso["summary"]
    return {
        "valves": valves, "consumers": consumers, "area": area,
        "consumers_count": len(consumers), "load_gcal_h": float(summary.get("total_with_downstream_load_gcal_h") or 0),
        "fileid": place["fileid"] if place else None, "lon": place["lon"] if place else None,
        "lat": place["lat"] if place else None,
    }


# --- чтение ---------------------------------------------------------------------------------

async def list_outages(conn, *, status: Optional[str] = None, open_only: bool = False,
                       q: Optional[str] = None, limit: int = MAX_LIST) -> dict[str, Any]:
    where, args = [], []
    if status:
        if status not in STATUSES:
            raise OutageError(422, f"Статус — один из: {', '.join(STATUSES)}")
        args.append(status)
        where.append(f"status = ${len(args)}")
    elif open_only:
        args.append(list(OPEN))
        where.append(f"status = ANY(${len(args)}::text[])")
    if q and q.strip():
        args.append(q.strip()[:100])
        where.append(f"(title ILIKE '%' || ${len(args)} || '%' OR number ILIKE '%' || ${len(args)} || '%')")
    args.append(min(max(limit, 1), MAX_LIST))
    rows = await conn.fetch(
        f"""SELECT id, number, title, kind, status, line_id, node_id, work_order_id, fileid, lon, lat,
                   planned_start, planned_end, actual_start, actual_end, reason, public, public_note,
                   consumers_count, load_gcal_h, computed_at, created_by, created_at, updated_at, revision
              FROM tgid_outage {'WHERE ' + ' AND '.join(where) if where else ''}
             ORDER BY (status IN ('finished', 'cancelled')), planned_start DESC, id DESC LIMIT ${len(args)}""", *args)
    counts = {r["status"]: r["n"] for r in await conn.fetch(
        "SELECT status, count(*) AS n FROM tgid_outage GROUP BY status")}
    return {"items": [_row(r, full=False) for r in rows], "counts": {k: counts.get(k, 0) for k in STATUSES}}


def _consumer_keys(consumers: list[dict[str, Any]]) -> set[str]:
    return {f"{c.get('type')}:{c.get('id')}" for c in consumers}


async def conflicts(conn, outage: dict[str, Any]) -> list[dict[str, Any]]:
    """Открытые отключения с перекрывающимся временем и общими потребителями."""
    rows = await conn.fetch(
        """SELECT id, number, title, status, planned_start, planned_end, consumers FROM tgid_outage
            WHERE id <> $1 AND status = ANY($2::text[]) AND planned_start < $4 AND planned_end > $3""",
        outage["id"], list(OPEN), _aware(datetime.fromisoformat(outage["planned_start"])),
        _aware(datetime.fromisoformat(outage["planned_end"])))
    mine = _consumer_keys(outage.get("consumers") or [])
    out = []
    for r in rows:
        shared = mine & _consumer_keys(_json(r["consumers"]) or [])
        if shared:
            out.append({"id": r["id"], "number": r["number"], "title": r["title"], "status": r["status"],
                        "planned_start": _iso(r["planned_start"]), "planned_end": _iso(r["planned_end"]),
                        "shared_consumers": len(shared)})
    return out


async def get_outage(conn, outage_id: int) -> dict[str, Any]:
    row = await conn.fetchrow("SELECT * FROM tgid_outage WHERE id = $1", outage_id)
    if row is None:
        raise OutageError(404, f"Отключение {outage_id} не найдено")
    out = _row(row)
    events = await conn.fetch(
        "SELECT at, by_user, from_status, to_status, note FROM tgid_outage_event WHERE outage_id = $1 ORDER BY at, id",
        outage_id)
    out["events"] = [{**{k: _iso(v) for k, v in dict(e).items()},
                      "to_label": STATUSES.get(e["to_status"]) if e["to_status"] else None} for e in events]
    out["conflicts"] = await conflicts(conn, out)
    # школы, больницы и другие социальные объекты в зоне — первыми в уведомлениях (database/poi.py)
    out["social"] = await poi.in_area(conn, out.get("area"))
    return out


async def map_geojson(conn, *, statuses: tuple[str, ...] = OPEN, public_only: bool = False) -> dict[str, Any]:
    """Зоны отключений полигонами; публичная лента — без названий потребителей и причины для служб."""
    rows = await conn.fetch(
        f"""SELECT id, number, title, kind, status, planned_start, planned_end, consumers_count, load_gcal_h,
                   public_note, area
              FROM tgid_outage WHERE status = ANY($1::text[]) AND area IS NOT NULL
               {'AND public' if public_only else ''}
             ORDER BY planned_start""", list(statuses))
    features = []
    for r in rows:
        props: dict[str, Any] = {
            "id": r["id"], "number": r["number"], "kind": r["kind"], "kind_label": KINDS[r["kind"]],
            "status": r["status"], "status_label": STATUSES[r["status"]],
            "planned_start": _iso(r["planned_start"]), "planned_end": _iso(r["planned_end"]),
            "consumers_count": r["consumers_count"],
        }
        if public_only:
            props["note"] = r["public_note"]
        else:
            props.update(title=r["title"], load_gcal_h=round(r["load_gcal_h"] or 0, 3))
        features.append({"type": "Feature", "id": r["id"], "geometry": _json(r["area"]), "properties": props})
    return {"type": "FeatureCollection", "features": features}


# --- запись ---------------------------------------------------------------------------------

async def _event(conn, outage_id: int, user: str, note: str, from_status: Optional[str] = None,
                 to_status: Optional[str] = None) -> None:
    await conn.execute(
        "INSERT INTO tgid_outage_event (outage_id, by_user, from_status, to_status, note) VALUES ($1, $2, $3, $4, $5)",
        outage_id, user, from_status, to_status, note)


async def _target_from_work_order(conn, work_order_id: int) -> tuple[Optional[int], Optional[int]]:
    row = None
    if await conn.fetchval("SELECT to_regclass('public.tgid_work_order') IS NOT NULL"):
        row = await conn.fetchrow("SELECT line_id, node_id FROM tgid_work_order WHERE id = $1", work_order_id)
    if row is None:
        raise OutageError(422, f"Наряда {work_order_id} нет")
    return row["line_id"], row["node_id"]


async def create(conn, body: OutageIn, user: str) -> dict[str, Any]:
    start, end = _aware(body.planned_start), _aware(body.planned_end)
    if end <= start:
        raise OutageError(422, "Окончание отключения раньше начала")
    line_id, node_id = body.line_id, body.node_id
    if body.work_order_id and not (line_id or node_id):
        line_id, node_id = await _target_from_work_order(conn, body.work_order_id)
    if not (line_id or node_id):
        raise OutageError(422, "Укажите участок или узел (или наряд с местом работ)")
    snap = await _snapshot(conn, line_id, None if line_id else node_id)
    async with conn.transaction():
        outage_id = await conn.fetchval(
            """INSERT INTO tgid_outage (title, kind, line_id, node_id, work_order_id, fileid, lon, lat, planned_start,
                                         planned_end, reason, public, public_note, consumers_count, load_gcal_h,
                                         valves, consumers, area, computed_at, created_by, updated_by)
               VALUES ($1, $2, $3, $4, $5, $6, $7, $8, $9, $10, $11, $12, $13, $14, $15, $16::jsonb, $17::jsonb,
                       $18::jsonb, now(), $19, $19) RETURNING id""",
            body.title.strip(), body.kind, line_id, None if line_id else node_id, body.work_order_id, snap["fileid"],
            snap["lon"], snap["lat"], start, end, body.reason, body.public, body.public_note,
            snap["consumers_count"], snap["load_gcal_h"], json.dumps(snap["valves"], ensure_ascii=False),
            json.dumps(snap["consumers"], ensure_ascii=False),
            json.dumps(snap["area"]) if snap["area"] else None, user)
        await conn.execute(
            "UPDATE tgid_outage SET number = 'О-' || to_char(created_at, 'YYYY') || '-' || lpad(id::text, 5, '0') "
            "WHERE id = $1", outage_id)
        await _event(conn, outage_id, user,
                     f"Запланировано: потребителей {snap['consumers_count']}, задвижек {len(snap['valves'])}",
                     None, "planned")
    return await get_outage(conn, outage_id)


async def recompute(conn, outage_id: int, revision: int, user: str) -> dict[str, Any]:
    current = await get_outage(conn, outage_id)
    if current["status"] not in RECOMPUTABLE:
        raise OutageError(409, f"Отключение {current['number']} {current['status_label'].lower()} — пересчёт закрыт")
    snap = await _snapshot(conn, current["line_id"], None if current["line_id"] else current["node_id"])
    async with conn.transaction():
        ok = await conn.fetchval(
            """UPDATE tgid_outage SET valves = $3::jsonb, consumers = $4::jsonb, area = $5::jsonb, consumers_count = $6,
                      load_gcal_h = $7, computed_at = now(), revision = revision + 1, updated_at = now(), updated_by = $8
                WHERE id = $1 AND revision = $2 RETURNING id""",
            outage_id, revision, json.dumps(snap["valves"], ensure_ascii=False),
            json.dumps(snap["consumers"], ensure_ascii=False), json.dumps(snap["area"]) if snap["area"] else None,
            snap["consumers_count"], snap["load_gcal_h"], user)
        if ok is None:
            raise OutageError(409, "Отключение уже изменили — обновите карточку")
        await _event(conn, outage_id, user, f"Пересчитано: потребителей {current['consumers_count']} → "
                                            f"{snap['consumers_count']}, задвижек {len(current['valves'])} → "
                                            f"{len(snap['valves'])}")
    return await get_outage(conn, outage_id)


_LABELS = {"title": "название", "planned_start": "начало", "planned_end": "окончание", "reason": "причина",
           "public": "публикация", "public_note": "текст для жителей"}


async def update(conn, outage_id: int, patch: OutagePatch, user: str) -> dict[str, Any]:
    current = await get_outage(conn, outage_id)
    if current["status"] in ("finished", "cancelled"):
        raise OutageError(409, f"Отключение {current['number']} {current['status_label'].lower()} — правка закрыта")
    values = patch.model_dump(exclude_unset=True, exclude={"revision"})
    start = _aware(values.get("planned_start")) or _aware(datetime.fromisoformat(current["planned_start"]))
    end = _aware(values.get("planned_end")) or _aware(datetime.fromisoformat(current["planned_end"]))
    if end <= start:
        raise OutageError(422, "Окончание отключения раньше начала")
    if not values:
        return current
    sets = [f"{k} = ${i}" for i, k in enumerate(values, start=3)]
    async with conn.transaction():
        ok = await conn.fetchval(
            f"""UPDATE tgid_outage SET {', '.join(sets)}, revision = revision + 1, updated_at = now(),
                       updated_by = ${len(values) + 3} WHERE id = $1 AND revision = $2 RETURNING id""",
            outage_id, patch.revision, *values.values(), user)
        if ok is None:
            raise OutageError(409, "Отключение уже изменили — обновите карточку")
        await _event(conn, outage_id, user, "Изменено: " + ", ".join(_LABELS[k] for k in values))
    return await get_outage(conn, outage_id)


async def transition(conn, outage_id: int, body: OutageTransition, user: str) -> dict[str, Any]:
    current = await get_outage(conn, outage_id)
    allowed = TRANSITIONS.get(current["status"], ())
    if body.to not in allowed:
        names = ", ".join(STATUSES[s] for s in allowed) or "нет"
        raise OutageError(409, f"Из «{current['status_label']}» нельзя в «{STATUSES[body.to]}» (можно: {names})")
    stamps = {"active": "actual_start = COALESCE(actual_start, $4)", "finished": "actual_end = $4"}
    extra = stamps.get(body.to, "")
    async with conn.transaction():
        ok = await conn.fetchval(
            f"""UPDATE tgid_outage SET status = $3, revision = revision + 1, updated_at = $4, updated_by = $5
                       {', ' + extra if extra else ''}
                 WHERE id = $1 AND revision = $2 RETURNING id""",
            outage_id, body.revision, body.to, datetime.now(timezone.utc), user)
        if ok is None:
            raise OutageError(409, "Отключение уже изменили — обновите карточку")
        await _event(conn, outage_id, user, body.note.strip(), current["status"], body.to)
    return await get_outage(conn, outage_id)


# --- уведомления ----------------------------------------------------------------------------

def _local(iso: str) -> str:
    return f"{datetime.fromisoformat(iso).astimezone():%d.%m.%Y %H:%M}"


def announcement(outage: dict[str, Any]) -> str:
    """Текст объявления для потребителей (SMS, сайт, письмо) — рассылку делает служба сама."""
    consumers = outage.get("consumers") or []
    names = [c["name"] for c in consumers if c.get("name")][:ANNOUNCE_NAMES]
    more = len(consumers) - len(names)
    kind = "Плановое" if outage["kind"] == "planned" else "Аварийное"
    parts = [f"{kind} отключение теплоснабжения {outage['number']}: с {_local(outage['planned_start'])} "
             f"до {_local(outage['planned_end'])}."]
    if outage.get("reason"):
        parts.append(f"Причина: {outage['reason'].strip()}.")
    if consumers:
        listed = "; ".join(names) + (f" и ещё {more}" if more > 0 else "")
        parts.append(f"Отключаются потребители ({len(consumers)}): {listed}.")
    return " ".join(parts)


def notification_workbook(outage: dict[str, Any]) -> bytes:
    """Excel для обзвона и писем: шапка (сроки, причина, текст объявления), потребители, задвижки."""
    from openpyxl import Workbook
    from openpyxl.styles import Font

    bold = Font(bold=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "Уведомление"
    ws.append([f"{outage['kind_label']} отключение {outage['number']}: {outage['title']}"])
    ws["A1"].font = Font(bold=True, size=13)
    ws.append(["Сроки", f"с {_local(outage['planned_start'])} до {_local(outage['planned_end'])}"])
    ws.append(["Причина", outage.get("reason") or "—"])
    ws.append(["Потребителей", outage["consumers_count"], "нагрузка, Гкал/ч", round(outage["load_gcal_h"] or 0, 3)])
    ws.append(["Объявление", announcement(outage)])
    ws.append([])
    ws.append(["№", "Потребитель", "Тип", "Узел", "Нагрузка, Гкал/ч", "Ниже по сети", "Широта", "Долгота"])
    for cell in ws[ws.max_row]:
        cell.font = bold
    for i, c in enumerate(outage.get("consumers") or [], start=1):
        ws.append([i, c.get("name"), "обобщённый" if c.get("type") == "generalized" else "реальный",
                   c.get("node_id"), c.get("load"), "да" if c.get("downstream") else "", c.get("lat"), c.get("lng")])
    ws.column_dimensions["A"].width = 14
    ws.column_dimensions["B"].width = 48
    vs = wb.create_sheet("Задвижки")
    vs.append(["id", "Задвижка", "Участок", "Ду", "Обязательная", "Состояние", "Широта", "Долгота"])
    for cell in vs[1]:
        cell.font = bold
    for v in outage.get("valves") or []:
        vs.append([v.get("id"), v.get("name"), v.get("line_id"), v.get("diameter"), "да" if v.get("required") else "",
                   v.get("state"), v.get("lat"), v.get("lng")])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
