"""Генерация официального Excel-отчета по пьезометрическому графику.

Портировано из legacy C++: gid8/gid8/pjezo/p_excel.cpp с поддержкой openpyxl,
технологической спецификации по участкам (подающий/обратный трубопроводы)
и профильного линейного графика напоров и отметок рельефа.
"""

import io
import math
from typing import Any
import openpyxl
from openpyxl.chart import LineChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side


def _r(value, digits: int):
    return round(float(value), digits) if value is not None else None


def generate_piezometer_excel(path_data: list[dict[str, Any]], segment_details: list[dict[str, Any]]) -> bytes:
    wb = openpyxl.Workbook()

    # -------------------------------------------------------------
    # Лист 1: Технологическая информация (Техн.информация)
    # -------------------------------------------------------------
    ws_tech = wb.active
    ws_tech.title = "Техн.информация"

    # Стили
    font_title = Font(name="Calibri", size=13, bold=True, color="1F497D")
    font_header = Font(name="Calibri", size=9, bold=True, color="FFFFFF")
    font_data = Font(name="Calibri", size=9)
    fill_header = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
    fill_sub = PatternFill(start_color="2962FF", end_color="2962FF", fill_type="solid")
    fill_alt = PatternFill(start_color="F5F7FA", end_color="F5F7FA", fill_type="solid")

    thin_border = Border(
        left=Side(style="thin", color="CCCCCC"),
        right=Side(style="thin", color="CCCCCC"),
        top=Side(style="thin", color="CCCCCC"),
        bottom=Side(style="thin", color="CCCCCC"),
    )
    align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    align_left = Alignment(horizontal="left", vertical="center")
    align_right = Alignment(horizontal="right", vertical="center")

    # Заголовок
    ws_tech.merge_cells("A3:R3")
    title_cell = ws_tech["A3"]
    title_cell.value = "ТЕХНОЛОГИЧЕСКАЯ ИНФОРМАЦИЯ К ПЬЕЗОМЕТРИЧЕСКОМУ ГРАФИКУ"
    title_cell.font = font_title
    title_cell.alignment = align_center

    # Шапка таблицы (строки 6..13)
    def style_merged(range_str, text, fill=fill_header, font=font_header):
        ws_tech.merge_cells(range_str)
        first = ws_tech[range_str.split(":")[0]]
        first.value = text
        for row in ws_tech[range_str]:
            for cell in row:
                cell.fill = fill
                cell.font = font
                cell.alignment = align_center
                cell.border = thin_border

    style_merged("A6:C7", "Начальный узел участка")
    style_merged("D6:F7", "Конечный узел участка")
    style_merged("G6:N7", "Параметры режима")
    style_merged("O6:O13", "Длина участка\nтрубопровода,\nм")
    style_merged("P6:P13", "Внутренний\nдиаметр,\nмм")
    style_merged("Q6:Q13", "Расстояние от\nначала,\nм")
    style_merged("R6:R13", "Объем от\nначала,\nм³")

    style_merged("A8:A13", "Код\nсхемы", fill=fill_sub)
    style_merged("B8:B13", "Наименование\nузла", fill=fill_sub)
    style_merged("C8:C13", "Признак\nтрубы", fill=fill_sub)

    style_merged("D8:D13", "Код\nсхемы", fill=fill_sub)
    style_merged("E8:E13", "Наименование\nузла", fill=fill_sub)
    style_merged("F8:F13", "Признак\nтрубы", fill=fill_sub)

    style_merged("G8:G13", "Расход\nводы,\nт/ч", fill=fill_sub)
    style_merged("H8:I9", "Пьезометрический напор, м", fill=fill_sub)
    style_merged("H10:H13", "в нач.\nузле", fill=fill_sub)
    style_merged("I10:I13", "в кон.\nузле", fill=fill_sub)

    style_merged("J8:M9", "Потери напора на участке, м", fill=fill_sub)
    style_merged("J10:J13", "удель-\nные,\nмм/м", fill=fill_sub)
    style_merged("K10:K13", "линей-\nные", fill=fill_sub)
    style_merged("L10:L13", "мест-\nные", fill=fill_sub)
    style_merged("M10:M13", "общие", fill=fill_sub)

    style_merged("N8:N13", "Скорость\nпотока,\nм/c", fill=fill_sub)

    # Строка 14: номера колонок 1..18
    for c in range(1, 19):
        cell = ws_tech.cell(row=14, column=c, value=c)
        cell.fill = fill_sub
        cell.font = font_header
        cell.alignment = align_center
        cell.border = thin_border

    # Данные участков (начиная со строки 15)
    row_idx = 15
    cum_vol = 0.0

    for seg in segment_details:
        length = float(seg.get("length") or 0.0)
        diam = float(seg.get("diameter") or 0.0)
        r_m = (diam / 2000.0) if diam > 0 else 0.0
        v_seg = math.pi * (r_m ** 2) * length
        cum_vol += v_seg

        dist = round(float(seg.get("distance_to_end") or 0.0), 1)

        def pipe_row(pipe: dict, label: str, h_start, h_end) -> list:
            """Строка участка по данным своей трубы (ut_out подачи или обратки)."""
            return [
                seg.get("node1_id"),
                seg.get("node1_label") or f"Узел {seg.get('node1_id')}",
                label,
                seg.get("node2_id"),
                seg.get("node2_label") or f"Узел {seg.get('node2_id')}",
                label,
                _r(pipe.get("flow"), 2),
                _r(h_start, 2),
                _r(h_end, 2),
                _r(pipe.get("spec_loss"), 2),
                _r(pipe.get("loss_linear"), 3),
                _r(pipe.get("loss_local"), 3),
                _r(pipe.get("loss_total"), 3),
                _r(pipe.get("velocity"), 2),
                length,
                diam,
                dist,
                round(cum_vol, 3),
            ]

        row_pod = pipe_row(seg.get("supply") or {}, "Подающий", seg.get("h_pod_start"), seg.get("h_pod_end"))
        row_obr = pipe_row(seg.get("return") or {}, "Обратный", seg.get("h_obr_start"), seg.get("h_obr_end"))

        for r_data in [row_pod, row_obr]:
            for col_idx, val in enumerate(r_data, start=1):
                c = ws_tech.cell(row=row_idx, column=col_idx, value=val)
                c.font = font_data
                c.border = thin_border
                if r_data is row_obr:
                    c.fill = fill_alt
                if col_idx in (2, 5):
                    c.alignment = align_left
                elif col_idx in (3, 6):
                    c.alignment = align_center
                else:
                    c.alignment = align_right
            row_idx += 1

    # Ширина столбцов
    col_widths = {
        "A": 10, "B": 22, "C": 12, "D": 10, "E": 22, "F": 12,
        "G": 14, "H": 12, "I": 12, "J": 12, "K": 12, "L": 12,
        "M": 12, "N": 12, "O": 14, "P": 14, "Q": 16, "R": 16,
    }
    for col_letter, width in col_widths.items():
        ws_tech.column_dimensions[col_letter].width = width

    # -------------------------------------------------------------
    # Лист 2: Профиль и график (Пьезометрический профиль)
    # -------------------------------------------------------------
    ws_chart = wb.create_sheet(title="Пьезометрический профиль")

    ws_chart.append(["Расстояние (м)", "Геодезическая отметка Z (м)", "Напор H_под (м)", "Напор H_обр (м)", "Узел"])
    for cell in ws_chart[1]:
        cell.fill = fill_header
        cell.font = font_header
        cell.alignment = align_center

    for p in path_data:
        dist = float(p.get("distance") or 0.0)
        z = float(p.get("z") or 0.0)
        h_pod = float(p.get("h_pod")) if p.get("h_pod") is not None else None
        h_obr = float(p.get("h_obr")) if p.get("h_obr") is not None else None
        label = p.get("label") or str(p.get("node_id"))
        ws_chart.append([dist, z, h_pod, h_obr, label])

    # График профиля напоров
    if len(path_data) >= 2:
        chart = LineChart()
        chart.title = "Пьезометрический график сети"
        chart.style = 13
        chart.y_axis.title = "Напор / Отметка (м)"
        chart.x_axis.title = "Расстояние (м)"
        chart.width = 24
        chart.height = 14

        data_ref = Reference(ws_chart, min_col=2, min_row=1, max_col=4, max_row=len(path_data) + 1)
        cats_ref = Reference(ws_chart, min_col=1, min_row=2, max_row=len(path_data) + 1)
        chart.add_data(data_ref, titles_from_data=True)
        chart.set_categories(cats_ref)

        ws_chart.add_chart(chart, "G3")

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output.getvalue()
