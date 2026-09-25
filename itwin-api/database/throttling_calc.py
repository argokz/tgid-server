"""Инженерные расчеты дросселирования: расчет шайб (диафрагм) и элеваторов.

Эталоны десктопа:
- gid8/python/dross/dross.py — бланк расчёта диафрагм теплового ввода (формулы ячеек G20..G43);
- sety/dross/drsh2.py — сопло и горловина элеватора, номер элеватора по диаметру горловины;
- sety/dross/drvary1.py — предупреждения о недостаточном/повышенном напоре на элеваторе.

Расходы считаются в т/ч, напоры — в м вод. ст., нагрузки — в Гкал/ч
(на десктопе нагрузки в ккал/ч: Gо = Qо / ((T1 - T2) * 1000) — то же самое).
"""

from __future__ import annotations

import io
import math
from typing import Any, Optional

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.protection import SheetProtection

# Номер элеватора по диаметру горловины, мм (drsh2.py: <=15 №1 … <=47 №6, <=59 №7).
# drsh2 при dg <= 10 выдаёт «№0»; стандартного элеватора №0 нет — такой ввод получает №1.
ELEVATOR_NECKS: tuple[tuple[int, float], ...] = (
    (1, 15.0), (2, 20.0), (3, 25.0), (4, 30.0), (5, 35.0), (6, 47.0), (7, 59.0),
)

MIN_ORIFICE_MM = 3.0

# Схемы шайб бланка dross.py: (коэффициент формулы, поправка к располагаемому напору, м)
ORIFICE_SCHEMES: dict[str, tuple[float, float]] = {
    "bezelevator": (10.0, -5.0),      # диафрагма безэлеваторного ввода, G26: Нгас = Нрас - 5
    "pump_mix": (10.0, -2.0),         # перед насосами смешения, G29: Нгас = Нрас - 2
    "pre_nozzle": (10.0, 0.0),        # перед соплом элеватора, G32: Нгас = Нрас
    "nozzle": (9.6, 0.0),             # сопло элеватора, G35: Нрас
    "ventilation": (10.0, -5.0),      # на вентиляцию, G38: Нрас - 5
    "heater": (10.0, -5.0),           # перед водоводяным подогревателем, G41: Нрас - 5 (расход отопления)
    "gvs_circulation": (10.0, -5.0),  # на циркуляционную линию ГВС, G43: Нрас - 5 (расход ГВС)
    "gvs": (10.0, -5.0),              # прежнее имя циркуляционной схемы
}


def flow_from_load(q_gcal: float, t_supply: float, t_return: float) -> float:
    """Расход сетевой воды, т/ч, по нагрузке Гкал/ч и перепаду температур."""
    if t_supply <= t_return:
        raise ValueError("Температура подачи должна быть выше температуры обратки.")
    return q_gcal * 1000.0 / (t_supply - t_return)


def gvs_flow_from_max_load(q_gvs_max_gcal: float) -> float:
    """Расход на ГВС, т/ч: Qгвс.ср = Qгвс.max / 2.4; G = Qгвс.ср / 60000 (ккал/ч), dross.py G17/G22."""
    return q_gvs_max_gcal * 1_000_000.0 / 2.4 / 60000.0


def calculate_orifice_diameter(
    *,
    flow_g: float,
    delta_h: float,
    coeff: float = 10.0,
    min_diameter: float = MIN_ORIFICE_MM,
) -> Optional[float]:
    """Диаметр отверстия шайбы (сопла), мм: d = coeff * (G² / Н)^(1/4), не меньше 3 мм.

    None — если гасимого напора нет (Н <= 0): на десктопе формула в этом случае не считается.
    """
    if delta_h <= 0:
        return None
    if flow_g <= 0:
        return min_diameter
    val = round(coeff * math.pow((flow_g * flow_g) / delta_h, 0.25), 1)
    return max(min_diameter, val)


def _resolve_flow(flow_g, q_heating_gcal, q_heating_kcal, t_supply, t_return) -> float:
    if flow_g is not None and flow_g > 0:
        return flow_g
    if q_heating_gcal is not None and q_heating_gcal > 0:
        return flow_from_load(q_heating_gcal, t_supply, t_return)
    if q_heating_kcal is not None and q_heating_kcal > 0:
        return flow_from_load(q_heating_kcal / 1_000_000.0, t_supply, t_return)
    raise ValueError("Укажите расход G или тепловую нагрузку Q.")


def _resolve_head(delta_h, p1, p2) -> float:
    if delta_h is not None and delta_h > 0:
        return delta_h
    if p1 is not None and p2 is not None:
        head = (p1 - p2) * 10.0
        if head <= 0:
            raise ValueError("Располагаемый напор (P1 - P2) должен быть больше нуля.")
        return head
    raise ValueError("Укажите располагаемый напор ΔH или давления P1 и P2.")


def calculate_orifice_plate_full(
    *,
    flow_g: Optional[float] = None,
    delta_h: Optional[float] = None,
    p1: Optional[float] = None,
    p2: Optional[float] = None,
    q_heating_gcal: Optional[float] = None,
    q_heating_kcal: Optional[float] = None,
    t_supply: float = 130.0,
    t_return: float = 70.0,
    scheme: str = "bezelevator",
) -> dict[str, Any]:
    """Расчёт одной шайбы по схеме установки (см. ORIFICE_SCHEMES)."""
    if scheme not in ORIFICE_SCHEMES:
        raise ValueError(f"Неизвестная схема шайбы: {scheme}")
    flow = _resolve_flow(flow_g, q_heating_gcal, q_heating_kcal, t_supply, t_return)
    head = _resolve_head(delta_h, p1, p2)
    coeff, correction = ORIFICE_SCHEMES[scheme]
    h_dissipated = head + correction

    d_orifice = calculate_orifice_diameter(flow_g=flow, delta_h=h_dissipated, coeff=coeff)
    warning = None
    if d_orifice is None:
        warning = (f"Недостаточный располагаемый напор ({head:.1f} м): "
                   f"для этой схемы нужно больше {-correction:g} м, шайба не рассчитывается.")
    elif scheme == "pre_nozzle" and h_dissipated <= 35.0:
        warning = "Диафрагма перед соплом устанавливается при гасимом напоре более 35 м."

    return {
        "diameter_orifice_mm": d_orifice,
        "recommended_standard_diameter": (round(d_orifice * 2.0) / 2.0) if d_orifice is not None else None,
        "flow_g": round(flow, 2),
        "available_head_m": round(head, 2),
        "head_loss_dissipated": round(h_dissipated, 2),
        "scheme": scheme,
        "warning": warning,
    }


def select_elevator_number(neck_mm: float) -> int:
    for number, max_neck in ELEVATOR_NECKS:
        if neck_mm <= max_neck:
            return number
    return 7


def calculate_elevator_parameters(
    *,
    q_heating_gcal: Optional[float] = None,
    flow_g: Optional[float] = None,
    p1: float = 6.0,
    p2: float = 4.0,
    t1: float = 130.0,
    t2: float = 70.0,
    t3: float = 95.0,
    delta_h_system: float = 1.5,
) -> dict[str, Any]:
    """Элеватор: коэффициент смешения, сопло и горловина (drsh2.py), номер по горловине."""
    if not (t1 > t3 > t2):
        raise ValueError("Нужно T1 > T3 > T2 (подача, после смешения, обратка).")
    if delta_h_system <= 0:
        raise ValueError("Потери напора в системе отопления hс должны быть больше нуля.")
    # Коэффициент смешения u = (T1 - T3) / (T3 - T2)
    u = (t1 - t3) / (t3 - t2)
    flow = _resolve_flow(flow_g, q_heating_gcal, None, t1, t2)
    h_avail = _resolve_head(None, p1, p2)

    # Сопло: dc = 9.6 * sqrt(G / sqrt(Нрас)) — полный располагаемый напор, как в drsh2
    dc = max(MIN_ORIFICE_MM, round(9.6 * math.sqrt(flow / math.sqrt(h_avail)), 1))
    # Горловина: dg = 8.5 * sqrt(G * (1 + u) / sqrt(hс)) — Апарцев 7.2
    d_neck = round(8.5 * math.sqrt(flow * (1.0 + u) / math.sqrt(delta_h_system)), 1)
    number = select_elevator_number(d_neck)

    # Требуемый напор перед элеватором ≈ 1.4 * hс * (1 + u)²
    h_required = 1.4 * delta_h_system * (1.0 + u) ** 2
    warnings: list[str] = []
    if d_neck > ELEVATOR_NECKS[-1][1]:
        warnings.append("Горловина больше, чем у элеватора №7: элеваторное смешение "
                        "должно быть заменено на насосное.")
    if h_avail < h_required:
        warnings.append(f"Располагаемый напор {h_avail:.1f} м меньше требуемого "
                        f"{h_required:.1f} м: элеватор не обеспечит смешение.")
    elif h_avail > 2 * h_required:
        warnings.append("Элеватор работает при повышенном напоре: возможны вибрация и шум "
                        "(избыток погасить диафрагмой перед соплом).")

    return {
        "mixing_ratio_u": round(u, 2),
        "nozzle_diameter_mm": dc,
        "mixing_chamber_diameter_mm": d_neck,
        "elevator_number": number,
        "available_head_m": round(h_avail, 2),
        "dissipated_head_m": round(h_avail, 2),
        "required_head_m": round(h_required, 2),
        "flow_g": round(flow, 2),
        "t1": t1,
        "t2": t2,
        "t3": t3,
        "warnings": warnings,
    }


def calculate_throttling_sheet(data: dict[str, Any]) -> dict[str, Any]:
    """Все величины бланка dross.py по исходным данным ввода."""
    t1 = float(data.get("t1") or 130.0)
    t2 = float(data.get("t2") or 70.0)
    q_ot = float(data.get("q_heating_gcal") or 0.0)
    q_v = float(data.get("q_vent_gcal") or 0.0)
    q_gvs_max = float(data.get("q_gvs_gcal") or 0.0)
    if q_ot <= 0 and q_v <= 0 and q_gvs_max <= 0:
        raise ValueError("Укажите хотя бы одну тепловую нагрузку (отопление, вентиляция или ГВС).")
    g_ot = flow_from_load(q_ot, t1, t2)
    g_v = flow_from_load(q_v, t1, t2)
    g_gvs = gvs_flow_from_max_load(q_gvs_max)
    h_ras = _resolve_head(None, data.get("p1"), data.get("p2"))

    def orifice(flow: float, scheme: str) -> Optional[float]:
        coeff, correction = ORIFICE_SCHEMES[scheme]
        return calculate_orifice_diameter(flow_g=flow, delta_h=h_ras + correction, coeff=coeff)

    return {
        "g_ot": g_ot, "g_v": g_v, "g_gvs": g_gvs, "q_gvs_avg": q_gvs_max / 2.4, "h_ras": h_ras,
        "d_bezelevator": orifice(g_ot, "bezelevator"),
        "d_pump_mix": orifice(g_ot, "pump_mix"),
        "d_pre_nozzle": orifice(g_ot, "pre_nozzle"),
        "d_nozzle": orifice(g_ot, "nozzle"),
        "d_ventilation": orifice(g_v, "ventilation"),
        "d_heater": orifice(g_ot, "heater"),
        "d_gvs_circulation": orifice(g_gvs, "gvs_circulation"),
    }


def generate_throttling_excel(data: dict[str, Any]) -> bytes:
    """Бланк расчёта дроссельных устройств теплового ввода (структура gid8/python/dross/dross.py)."""
    calc = calculate_throttling_sheet(data)

    wb = Workbook()
    ws = wb.active
    ws.title = "Расчет диафрагм"

    bold_font = Font(name="Arial", size=10, bold=True)
    regular_font = Font(name="Arial", size=10, bold=False)
    title_font = Font(name="Arial", size=12, bold=True, underline="single")
    fill_input = PatternFill("solid", fgColor="FFF9B1")  # исходные данные
    fill_calc = PatternFill("solid", fgColor="E8F5E9")   # расчётные величины

    for col, width in {"A": 4, "B": 50, "C": 14, "D": 10, "E": 12, "F": 6, "G": 14,
                       "H": 8, "I": 8, "J": 8, "K": 8, "L": 10}.items():
        ws.column_dimensions[col].width = width

    ws.merge_cells("B2:L2")
    ws["B2"] = "РАСЧЕТ ДИАМЕТРОВ ОТВЕРСТИЙ ДРОССЕЛЬНЫХ УСТРОЙСТВ"
    ws["B2"].font = title_font
    ws["B2"].alignment = Alignment(horizontal="center", vertical="center")
    ws.merge_cells("B3:L3")
    ws["B3"] = "ДЛЯ ТЕПЛОВОГО ВВОДА ПОТРЕБИТЕЛЯ"
    ws["B3"].font = title_font
    ws["B3"].alignment = Alignment(horizontal="center", vertical="center")

    def put(cell: str, value: Any, *, font=regular_font, fill=None) -> None:
        ws[cell] = value
        ws[cell].font = font
        if fill is not None:
            ws[cell].fill = fill

    put("B5", "Район:", font=bold_font)
    put("C5", data.get("district") or "", fill=fill_input)
    put("G5", "Участок / ТК:", font=bold_font)
    put("H5", data.get("site_name") or "", fill=fill_input)
    put("B7", "Давление P1 (подача):")
    put("D7", data.get("p1"), fill=fill_input)
    put("E7", "атм")
    put("G7", "Давление P2 (обратка):")
    put("I7", data.get("p2"), fill=fill_input)
    put("J7", "атм")
    put("B9", "Потребитель:", font=bold_font)
    put("C9", data.get("consumer_name") or "", fill=fill_input)
    put("B11", "Адрес:")
    put("C11", data.get("address") or "", fill=fill_input)

    put("B13", "Тепловые нагрузки:", font=bold_font)
    put("B14", "а) на отопление (Qо):")
    put("G14", data.get("q_heating_gcal") or 0.0, fill=fill_input)
    put("L14", "Гкал/ч")
    put("B15", "б) на вентиляцию (Qв):")
    put("G15", data.get("q_vent_gcal") or 0.0, fill=fill_input)
    put("L15", "Гкал/ч")
    put("B16", "в) на ГВС: максимальная (Qгвс.max)")
    put("G16", data.get("q_gvs_gcal") or 0.0, fill=fill_input)
    put("L16", "Гкал/ч")
    put("B17", "                   среднечасовая (Qгвс.ср = Qгвс.max / 2.4)")
    put("G17", round(calc["q_gvs_avg"], 4), fill=fill_calc)
    put("L17", "Гкал/ч")
    put("B18", "Температурный график T1 / T2:", font=bold_font)
    put("G18", f"{data.get('t1', 130)} / {data.get('t2', 70)}", fill=fill_input)
    put("L18", "°С")

    put("B19", "Расходы воды:", font=bold_font)
    rows: list[tuple[str, str, Any, str]] = [
        ("B20", "а) на отопление (Gо):", round(calc["g_ot"], 3), "т/ч"),
        ("B21", "б) на вентиляцию (Gв):", round(calc["g_v"], 3), "т/ч"),
        ("B22", "в) на ГВС (Gгвс):", round(calc["g_gvs"], 3), "т/ч"),
        ("B23", "Располагаемый напор (Нрас):", round(calc["h_ras"], 2), "м"),
    ]
    diaphragms = [
        (24, "Диафрагма безэлеваторного ввода:", calc["h_ras"] - 5, calc["d_bezelevator"],
         "напор после шайбы 5 м", False),
        (27, "Диафрагма перед насосами смешения:", calc["h_ras"] - 2, calc["d_pump_mix"], "", False),
        (30, "Дроссельная диафрагма перед соплом элеватора:", calc["h_ras"], calc["d_pre_nozzle"],
         "устанавливается при гасимом напоре более 35 м", False),
        (33, "Элеватор (сопло):", calc["h_ras"], calc["d_nozzle"], "", True),
        (36, "Дроссельная диафрагма на вентиляцию:", calc["h_ras"] - 5, calc["d_ventilation"], "", False),
        (39, "Дроссельная диафрагма перед водоводяным подогревателем:", calc["h_ras"] - 5,
         calc["d_heater"], "", False),
        (42, "Дроссельная диафрагма на циркуляционную линию г.в.с.:", calc["h_ras"] - 5,
         calc["d_gvs_circulation"], "всегда одна шайба", False),
    ]
    for cell, label, value, unit in rows:
        put(cell, label)
        put("G" + cell[1:], value, fill=fill_calc)
        put("L" + cell[1:], unit)
    for row, title, head, diameter, note, is_nozzle in diaphragms:
        put(f"B{row}", title, font=bold_font)
        put(f"B{row + 1}", "Располагаемый напор (Нрас):" if is_nozzle else "Гасимый напор (Нгас):")
        put(f"G{row + 1}", round(head, 2), fill=fill_calc)
        put(f"L{row + 1}", "м")
        put(f"B{row + 2}", "Диаметр отверстия (Дс):" if is_nozzle else "Диаметр отверстия (Дш):")
        put(f"G{row + 2}", diameter if diameter is not None else "не рассчитывается",
            font=bold_font, fill=fill_calc)
        put(f"L{row + 2}", "мм" if diameter is not None else "")
        if note:
            put(f"N{row + 1}", note)

    ws.merge_cells("B46:L47")
    ws["B46"] = (
        "Примечание: диаметры отверстий дроссельных устройств могут быть скорректированы в "
        "зависимости от параметров теплоносителя на тепловом вводе, при изменении тепловой нагрузки "
        "на отопление и вентиляцию, либо при обосновании необходимости корректировки дроссельных "
        "устройств в акте обследования."
    )
    ws["B46"].font = Font(name="Arial", size=9, italic=True)
    ws["B46"].alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)

    # Подписи — из запроса; по умолчанию пустые строки для заполнения от руки
    signers = data.get("signers") or []
    for i, row in enumerate((50, 52)):
        signer = signers[i] if i < len(signers) else {}
        position = signer.get("position") or "Должность"
        name = signer.get("name") or ""
        put(f"B{row}", f"{position} ______________________ / {name:<20} /")
    if data.get("organization"):
        put("B54", data["organization"])

    # Защита листа с разрешением менять ширину колонок (как в dross.py)
    ws.protection = SheetProtection(sheet=True, formatColumns=False)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
