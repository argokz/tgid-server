"""Электросеть: сверка геометрической привязки и привязка по правилам десктопа (этап 10).

Эталон — gid6 ``GeoFile.cpp``:

* ``CGeoFile::createObj`` для ``liniya_elektroperedach``: после создания ЛЭП ищет источник
  (``istochnik_elektrosnabzheniya``) сначала у первой, затем у последней точки линии, так же
  приёмник (``priemnik_elektrosnabzheniya``); найденные id пишутся в
  ``naimenovanie_istochnika`` / ``naimenovanie_priemnika``.
* ``CGeoFile::createObjElPoint`` для муфты, опоры ЭС (точечные, ``isElPoint``), гильзы и
  кабельного канала ЭС (``isElSq``): ищет ЛЭП в радиусе ``D5*100`` от первой точки объекта
  (для площадных классов — ещё ×10), муфту/опору проецирует на ЛЭП (``getProject``), в объект
  пишет ``naimenovanie_lep`` = id ЛЭП.

Десктоп берёт первый найденный объект в радиусе, веб — ближайший (при равенстве — меньший id).
Отчёт «ЛЭП» (OnElectroRemont) в gid6 выключен, отдельной сверки в десктопе нет — виды
несоответствий ниже выведены из этих правил. Допуск задаётся в метрах (SRID 9998 — метры),
по умолчанию 8 м (``D5``). ``naimenovanie_lep`` у источников/приёмников десктоп не пишет
(ветка ``isElEnd`` недостижима), поэтому здесь не сверяется.
"""

from __future__ import annotations

import io
from typing import Any, Awaitable, Callable, Optional

import asyncpg

from database.sql_ident import UnknownIdentifierError, quote_ident, resolve_table

DEFAULT_TOLERANCE_M = 8.0
AREA_TOLERANCE_FACTOR = 10.0
DISTANCE_EPS_M = 0.01
EXCEL_ROW_LIMIT = 100_000

AuditRow = Callable[..., Awaitable[str]]

LINE_TABLE = "liniya_elektroperedach"
SOURCE_TABLE = "istochnik_elektrosnabzheniya"
RECEIVER_TABLE = "priemnik_elektrosnabzheniya"
# object_type -> (таблица, колонка номера, проецировать на ЛЭП как десктоп getProject)
POINT_TABLES: dict[str, tuple[str, str, bool]] = {
    "channel": ("kabelnyy_kanal_es", "nomer_kanala_es", False),
    "coupling": ("mufta", "nomer_mufty_es", True),
    "support": ("opora_es", "nomer_opory_es", True),
    "sleeve": ("gilza_es", "nomer_gilzy_es", False),
}
ALLOWED_TABLES = frozenset({LINE_TABLE, SOURCE_TABLE, RECEIVER_TABLE, *(t for t, _, _ in POINT_TABLES.values())})
OBJECT_TYPES = ("line", *POINT_TABLES)

LINE_KINDS: dict[str, str] = {
    "line_no_geometry": "ЛЭП без геометрии",
    "line_no_source": "ЛЭП без источника",
    "line_source_missing": "Источник ЛЭП не найден",
    "line_source_far": "Концы ЛЭП не у своего источника",
    "line_no_receiver": "ЛЭП без приёмника",
    "line_receiver_missing": "Приёмник ЛЭП не найден",
    "line_receiver_far": "Концы ЛЭП не у своего приёмника",
}
POINT_KINDS: dict[str, str] = {
    "point_no_geometry": "Объект без геометрии",
    "point_no_line": "Объект без ЛЭП",
    "point_line_missing": "ЛЭП объекта не найдена",
    "point_off_line": "Объект не на своей ЛЭП",
    "point_closer_other": "Объект ближе к другой ЛЭП",
}
ISSUE_KINDS: dict[str, str] = {**LINE_KINDS, **POINT_KINDS}
TYPE_TITLES = {
    "line": "ЛЭП", "channel": "Кабельный канал", "coupling": "Муфта", "support": "Опора", "sleeve": "Гильза",
}


class ElectricalError(Exception):
    def __init__(self, status: int, code: str, message: str, **extra: Any):
        super().__init__(message)
        self.status = status
        self.detail = {"code": code, "message": message, **extra}


def _check_tolerance(tolerance: float) -> float:
    value = float(tolerance)
    if not 0 < value <= 1000:
        raise ElectricalError(422, "bad_tolerance", "Допуск должен быть в диапазоне (0; 1000] м")
    return value


async def _tables(conn: asyncpg.Connection) -> Optional[dict[str, str]]:
    """{логическое имя: имя в кавычках} через sql_ident; None — в БД нет таблиц электросети."""
    out: dict[str, str] = {}
    try:
        for name in ALLOWED_TABLES:
            out[name] = quote_ident(await resolve_table(conn, name, ALLOWED_TABLES))
    except UnknownIdentifierError:
        return None
    return out


def _lonlat(expr: str) -> str:
    return (f"CASE WHEN {expr} IS NULL THEN NULL ELSE ST_X(ST_Transform({expr}, 4326)) END AS longitude, "
            f"CASE WHEN {expr} IS NULL THEN NULL ELSE ST_Y(ST_Transform({expr}, 4326)) END AS latitude")


def _line_sql(t: dict[str, str]) -> str:
    def candidate(table: str, alias: str) -> str:
        # десктоп: сначала первая точка ЛЭП, потом последняя; внутри конца — ближайший
        return f"""
        LEFT JOIN LATERAL (
            SELECT c.id, least(ST_Distance(c.shape, line.p1), ST_Distance(c.shape, line.p2)) AS distance
              FROM {table} c
             WHERE c.shape IS NOT NULL AND line.p1 IS NOT NULL
               AND (ST_DWithin(c.shape, line.p1, $1) OR ST_DWithin(c.shape, line.p2, $1))
             ORDER BY CASE WHEN ST_DWithin(c.shape, line.p1, $1) THEN 0 ELSE 1 END,
                      CASE WHEN ST_DWithin(c.shape, line.p1, $1) THEN ST_Distance(c.shape, line.p1)
                           ELSE ST_Distance(c.shape, line.p2) END, c.id
             LIMIT 1
        ) {alias} ON true"""

    return f"""
        WITH line AS (
            SELECT l.id, coalesce(nullif(l.naimenovanie_lep, ''), nullif(l.mestopolozhenie, ''))::text AS name,
                   l.naimenovanie_istochnika AS source_id, l.naimenovanie_priemnika AS receiver_id,
                   CASE WHEN l.shape IS NULL OR ST_IsEmpty(l.shape) THEN NULL
                        ELSE ST_StartPoint(ST_GeometryN(l.shape, 1)) END AS p1,
                   CASE WHEN l.shape IS NULL OR ST_IsEmpty(l.shape) THEN NULL
                        ELSE ST_EndPoint(ST_GeometryN(l.shape, ST_NumGeometries(l.shape))) END AS p2
              FROM {t[LINE_TABLE]} l
        )
        SELECT 'line'::text AS object_type, line.id, line.name, line.p1 IS NOT NULL AS has_geometry,
               line.source_id, s.id IS NOT NULL AS source_exists, s.naimenovanie_istochnika_es::text AS source_name,
               least(ST_Distance(line.p1, s.shape), ST_Distance(line.p2, s.shape)) AS source_distance,
               line.receiver_id, r.id IS NOT NULL AS receiver_exists,
               r.naimenovanie_priemnika_es::text AS receiver_name,
               least(ST_Distance(line.p1, r.shape), ST_Distance(line.p2, r.shape)) AS receiver_distance,
               cs.id AS candidate_source_id, cs.distance AS candidate_source_distance,
               cr.id AS candidate_receiver_id, cr.distance AS candidate_receiver_distance,
               {_lonlat("line.p1")}
          FROM line
          LEFT JOIN {t[SOURCE_TABLE]} s ON s.id = line.source_id
          LEFT JOIN {t[RECEIVER_TABLE]} r ON r.id = line.receiver_id
          {candidate(t[SOURCE_TABLE], "cs")}
          {candidate(t[RECEIVER_TABLE], "cr")}
         ORDER BY line.id
    """


def _point_sql(t: dict[str, str]) -> str:
    parts = []
    for object_type, (table, number_col, snap) in POINT_TABLES.items():
        parts.append(
            f"SELECT '{object_type}'::text AS object_type, o.id,"
            f" coalesce(nullif(o.naimenovanie_lep2, ''), nullif(o.{quote_ident(number_col)}, ''))::text AS name,"
            f" o.naimenovanie_lep AS line_id, o.shape, {'true' if snap else 'false'} AS snap_type"
            f" FROM {t[table]} o"
        )
    union = "\n UNION ALL ".join(parts)
    return f"""
        WITH obj AS ({union})
        SELECT obj.object_type, obj.id, obj.name, obj.shape IS NOT NULL AS has_geometry,
               GeometryType(obj.shape) AS geometry_type, obj.snap_type,
               obj.line_id, l.id IS NOT NULL AS line_exists, l.naimenovanie_lep::text AS line_name,
               ST_Distance(obj.shape, l.shape) AS line_distance,
               n.id AS nearest_line_id, n.distance AS nearest_line_distance,
               CASE WHEN GeometryType(obj.shape) IN ('POLYGON', 'MULTIPOLYGON') THEN $1 * {AREA_TOLERANCE_FACTOR}
                    ELSE $1 END AS search_tolerance,
               {_lonlat("ST_PointOnSurface(obj.shape)")}
          FROM obj
          LEFT JOIN {t[LINE_TABLE]} l ON l.id = obj.line_id
          LEFT JOIN LATERAL (
              SELECT x.id, ST_Distance(x.shape, obj.shape) AS distance
                FROM {t[LINE_TABLE]} x
               WHERE x.shape IS NOT NULL AND obj.shape IS NOT NULL
               ORDER BY ST_Distance(x.shape, obj.shape), x.id
               LIMIT 1
          ) n ON true
         ORDER BY obj.object_type, obj.id
    """


def _num(value: Any) -> Optional[float]:
    return None if value is None else round(float(value), 3)


def classify_line(row: dict[str, Any], tolerance: float) -> list[str]:
    if not row["has_geometry"]:
        return ["line_no_geometry"]
    issues = []
    for role in ("source", "receiver"):
        if row[f"{role}_id"] is None:
            issues.append(f"line_no_{role}")
        elif not row[f"{role}_exists"]:
            issues.append(f"line_{role}_missing")
        elif row[f"{role}_distance"] is None or row[f"{role}_distance"] > tolerance:
            issues.append(f"line_{role}_far")
    return issues


def classify_point(row: dict[str, Any]) -> list[str]:
    if not row["has_geometry"]:
        return ["point_no_geometry"]
    if row["line_id"] is None:
        return ["point_no_line"]
    if not row["line_exists"]:
        return ["point_line_missing"]
    issues = []
    distance = row["line_distance"]
    if distance is None or distance > row["search_tolerance"]:
        issues.append("point_off_line")
    nearest = row["nearest_line_distance"]
    if (row["nearest_line_id"] not in (None, row["line_id"]) and nearest is not None and distance is not None
            and nearest < distance - DISTANCE_EPS_M):
        issues.append("point_closer_other")
    return issues


def _line_item(row: dict[str, Any], issues: list[str], tolerance: float) -> dict[str, Any]:
    return {
        "object_type": "line", "id": row["id"], "name": row["name"], "issues": issues,
        "source_id": row["source_id"], "source_name": row["source_name"],
        "source_distance": _num(row["source_distance"]),
        "candidate_source_id": row["candidate_source_id"],
        "candidate_source_distance": _num(row["candidate_source_distance"]),
        "receiver_id": row["receiver_id"], "receiver_name": row["receiver_name"],
        "receiver_distance": _num(row["receiver_distance"]),
        "candidate_receiver_id": row["candidate_receiver_id"],
        "candidate_receiver_distance": _num(row["candidate_receiver_distance"]),
        "tolerance": tolerance, "longitude": row["longitude"], "latitude": row["latitude"],
    }


def _point_candidate(row: dict[str, Any]) -> tuple[Optional[int], Optional[float]]:
    nearest = row["nearest_line_distance"]
    if row["nearest_line_id"] is None or nearest is None or nearest > row["search_tolerance"]:
        return None, None
    return row["nearest_line_id"], nearest


def _point_item(row: dict[str, Any], issues: list[str]) -> dict[str, Any]:
    candidate, distance = _point_candidate(row)
    return {
        "object_type": row["object_type"], "id": row["id"], "name": row["name"], "issues": issues,
        "geometry_type": row["geometry_type"], "line_id": row["line_id"], "line_name": row["line_name"],
        "line_distance": _num(row["line_distance"]),
        "nearest_line_id": row["nearest_line_id"], "nearest_line_distance": _num(row["nearest_line_distance"]),
        "candidate_line_id": candidate, "candidate_line_distance": _num(distance),
        "tolerance": _num(row["search_tolerance"]), "longitude": row["longitude"], "latitude": row["latitude"],
    }


async def _fetch(conn: asyncpg.Connection, tolerance: float) -> tuple[list[dict], list[dict]]:
    t = await _tables(conn)
    if t is None:
        return [], []
    lines = [dict(r) for r in await conn.fetch(_line_sql(t), tolerance)]
    points = [dict(r) for r in await conn.fetch(_point_sql(t), tolerance)]
    return lines, points


async def reconcile(conn: asyncpg.Connection, tolerance: float = DEFAULT_TOLERANCE_M) -> dict[str, Any]:
    """Полная сверка: сводка по видам и все объекты с несоответствиями."""
    tolerance = _check_tolerance(tolerance)
    t = await _tables(conn)
    lines, points = await _fetch(conn, tolerance) if t is not None else ([], [])
    items: list[dict[str, Any]] = []
    for row in lines:
        issues = classify_line(row, tolerance)
        if issues:
            items.append(_line_item(row, issues, tolerance))
    for row in points:
        issues = classify_point(row)
        if issues:
            items.append(_point_item(row, issues))
    counts = {kind: 0 for kind in ISSUE_KINDS}
    for item in items:
        for kind in item["issues"]:
            counts[kind] += 1
    checked = {"line": len(lines)}
    for object_type in POINT_TABLES:
        checked[object_type] = sum(1 for p in points if p["object_type"] == object_type)
    return {
        "tolerance": tolerance,
        "tables_present": t is not None,
        "checked": checked,
        "kinds": [{"kind": k, "label": ISSUE_KINDS[k], "count": counts[k]} for k in ISSUE_KINDS],
        "items": items,
    }


def filter_items(result: dict[str, Any], *, kind: Optional[str], object_type: Optional[str]) -> list[dict]:
    if kind is not None and kind not in ISSUE_KINDS:
        raise ElectricalError(422, "unknown_kind", f"Неизвестный вид несоответствия: {kind}")
    if object_type is not None and object_type not in OBJECT_TYPES:
        raise ElectricalError(422, "unknown_object_type", f"Неизвестный вид объекта: {object_type}")
    return [
        item for item in result["items"]
        if (kind is None or kind in item["issues"]) and (object_type is None or item["object_type"] == object_type)
    ]


# --- привязка ----------------------------------------------------------------------------

def plan_line(row: dict[str, Any], tolerance: float, overwrite: bool) -> list[dict[str, Any]]:
    """Изменения полей ЛЭП: пустой (или, с overwrite, неверный) источник/приёмник → найденный у конца."""
    if not row["has_geometry"]:
        return []
    issues = set(classify_line(row, tolerance))
    changes = []
    for role, field in (("source", "naimenovanie_istochnika"), ("receiver", "naimenovanie_priemnika")):
        candidate = row[f"candidate_{role}_id"]
        if candidate is None or candidate == row[f"{role}_id"]:
            continue
        replace = row[f"{role}_id"] is None or (
            overwrite and ({f"line_{role}_missing", f"line_{role}_far"} & issues)
        )
        if replace:
            changes.append({"field": field, "old": row[f"{role}_id"], "new": candidate,
                            "distance": _num(row[f"candidate_{role}_distance"])})
    return changes


def plan_point(row: dict[str, Any], overwrite: bool, snap: bool) -> tuple[list[dict[str, Any]], Optional[int]]:
    """Изменение naimenovanie_lep (и признак проекции муфты/опоры на ЛЭП); вторым — итоговая ЛЭП."""
    if not row["has_geometry"]:
        return [], None
    issues = set(classify_point(row))
    candidate, distance = _point_candidate(row)
    changes = []
    target = row["line_id"] if row["line_exists"] and "point_off_line" not in issues else None
    if candidate is not None and candidate != row["line_id"]:
        replace = row["line_id"] is None or (
            overwrite and ({"point_line_missing", "point_off_line", "point_closer_other"} & issues)
        )
        if replace:
            changes.append({"field": "naimenovanie_lep", "old": row["line_id"], "new": candidate,
                            "distance": _num(distance)})
            target = candidate
    if snap and row["snap_type"] and row["geometry_type"] == "POINT" and target is not None:
        own = distance if target == candidate else row["line_distance"]
        if own is not None and own > 0.001:
            changes.append({"field": "shape", "old": None, "new": f"проекция на ЛЭП {target}",
                            "distance": _num(own), "snap_line_id": target})
    return changes, target


def _selected(targets: Optional[list[tuple[str, int]]], object_type: str, object_id: int) -> bool:
    return targets is None or (object_type, object_id) in targets


async def plan_binding(
    conn: asyncpg.Connection,
    *,
    tolerance: float = DEFAULT_TOLERANCE_M,
    overwrite: bool = False,
    snap_points: bool = False,
    targets: Optional[list[tuple[str, int]]] = None,
) -> dict[str, Any]:
    tolerance = _check_tolerance(tolerance)
    for object_type, _ in targets or []:
        if object_type not in OBJECT_TYPES:
            raise ElectricalError(422, "unknown_object_type", f"Неизвестный вид объекта: {object_type}")
    target_set = set(targets) if targets is not None else None
    lines, points = await _fetch(conn, tolerance)
    records: list[dict[str, Any]] = []
    unresolved: list[dict[str, Any]] = []
    for row in lines:
        if not _selected(target_set, "line", row["id"]):
            continue
        changes = plan_line(row, tolerance, overwrite)
        if changes:
            records.append({"object_type": "line", "table": LINE_TABLE, "id": row["id"], "name": row["name"],
                            "changes": changes, "longitude": row["longitude"], "latitude": row["latitude"]})
        elif classify_line(row, tolerance):
            unresolved.append({"object_type": "line", "id": row["id"], "issues": classify_line(row, tolerance)})
    for row in points:
        if not _selected(target_set, row["object_type"], row["id"]):
            continue
        changes, _ = plan_point(row, overwrite, snap_points)
        if changes:
            records.append({"object_type": row["object_type"], "table": POINT_TABLES[row["object_type"]][0],
                            "id": row["id"], "name": row["name"], "changes": changes,
                            "longitude": row["longitude"], "latitude": row["latitude"]})
        elif classify_point(row):
            unresolved.append({"object_type": row["object_type"], "id": row["id"], "issues": classify_point(row)})
    if target_set is not None:
        known = {("line", r["id"]) for r in lines} | {(p["object_type"], p["id"]) for p in points}
        missing = sorted(target_set - known)
        if missing:
            raise ElectricalError(404, "object_not_found", "Объекты электросети не найдены",
                                  missing=[{"object_type": o, "id": i} for o, i in missing])
    return {
        "tolerance": tolerance, "overwrite": overwrite, "snap_points": snap_points,
        "records": records, "unresolved": unresolved,
        "counts": {"records": len(records), "fields": sum(len(r["changes"]) for r in records),
                   "unresolved": len(unresolved)},
    }


async def bind(
    conn: asyncpg.Connection,
    *,
    tolerance: float = DEFAULT_TOLERANCE_M,
    overwrite: bool = False,
    snap_points: bool = False,
    targets: Optional[list[tuple[str, int]]] = None,
    dry_run: bool = True,
    audit_row: Optional[AuditRow] = None,
) -> dict[str, Any]:
    """dry_run — только план «было → станет»; иначе запись (вызывать внутри одной транзакции)."""
    plan = await plan_binding(conn, tolerance=tolerance, overwrite=overwrite, snap_points=snap_points,
                              targets=targets)
    result = {**plan, "dry_run": dry_run, "applied": 0, "change_group_id": None}
    if dry_run or not plan["records"]:
        return result
    t = await _tables(conn)
    assert t is not None
    group: Optional[str] = None
    for record in plan["records"]:
        table = t[record["table"]]
        old: dict[str, Any] = {}
        new: dict[str, Any] = {}
        for change in record["changes"]:
            if change["field"] == "shape":
                continue
            field = quote_ident(change["field"])
            status = await conn.execute(
                f"UPDATE {table} SET {field}=$2 WHERE id=$1 AND {field} IS NOT DISTINCT FROM $3",
                record["id"], change["new"], change["old"],
            )
            if not status.endswith(" 1"):
                raise ElectricalError(409, "concurrent_change", "Объект изменён другим пользователем, повторите",
                                      object_type=record["object_type"], id=record["id"], field=change["field"])
            old[change["field"]] = change["old"]
            new[change["field"]] = change["new"]
        snap = next((c for c in record["changes"] if c["field"] == "shape"), None)
        if snap is not None:
            geometry = await conn.fetchrow(
                f"""UPDATE {table} o SET shape = ST_ClosestPoint(l.shape, o.shape)
                      FROM {t[LINE_TABLE]} l
                     WHERE o.id=$1 AND l.id=$2
                 RETURNING ST_AsText(o.shape) AS new_wkt,
                           (SELECT ST_AsText(shape) FROM {table} WHERE id=$1) AS old_wkt""",
                record["id"], snap["snap_line_id"],
            )
            if geometry is None:
                raise ElectricalError(409, "concurrent_change", "Объект или ЛЭП удалены, повторите",
                                      object_type=record["object_type"], id=record["id"])
            old["shape"] = geometry["old_wkt"]
            new["shape"] = geometry["new_wkt"]
        if audit_row is not None:
            group = await audit_row(conn, operation="UPDATE", table=record["table"], record_id=record["id"],
                                    old=old, new=new, group=group)
        result["applied"] += 1
    result["change_group_id"] = group
    return result


# --- Excel ---------------------------------------------------------------------------------

LINE_COLUMNS = [
    ("id", "ID ЛЭП"), ("name", "Наименование"), ("issues", "Несоответствия"),
    ("source_id", "Источник"), ("source_name", "Наименование источника"), ("source_distance", "До источника, м"),
    ("candidate_source_id", "Источник у конца"), ("candidate_source_distance", "До него, м"),
    ("receiver_id", "Приёмник"), ("receiver_name", "Наименование приёмника"),
    ("receiver_distance", "До приёмника, м"),
    ("candidate_receiver_id", "Приёмник у конца"), ("candidate_receiver_distance", "До него, м"),
    ("longitude", "Долгота"), ("latitude", "Широта"),
]
POINT_COLUMNS = [
    ("object_type", "Вид"), ("id", "ID"), ("name", "Наименование / номер"), ("issues", "Несоответствия"),
    ("geometry_type", "Геометрия"), ("line_id", "ЛЭП"), ("line_name", "Наименование ЛЭП"),
    ("line_distance", "До своей ЛЭП, м"), ("nearest_line_id", "Ближайшая ЛЭП"),
    ("nearest_line_distance", "До ближайшей, м"), ("candidate_line_id", "ЛЭП для привязки"),
    ("tolerance", "Допуск, м"), ("longitude", "Долгота"), ("latitude", "Широта"),
]


def _cell(key: str, value: Any) -> Any:
    if key == "issues":
        return "; ".join(ISSUE_KINDS.get(v, v) for v in value)
    if key == "object_type":
        return TYPE_TITLES.get(value, value)
    return value


async def reconciliation_workbook(conn: asyncpg.Connection, tolerance: float = DEFAULT_TOLERANCE_M) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    result = await reconcile(conn, tolerance)
    plan = await plan_binding(conn, tolerance=result["tolerance"])
    wb = Workbook()
    summary = wb.active
    summary.title = "Сводка"
    summary.append([f"Сверка привязки электросети, допуск {result['tolerance']:g} м"])
    summary["A1"].font = Font(bold=True)
    summary.append(["Вид несоответствия", "Количество"])
    for kind in result["kinds"]:
        summary.append([kind["label"], kind["count"]])
    summary.append([])
    summary.append(["Проверено объектов"])
    for object_type, count in result["checked"].items():
        summary.append([TYPE_TITLES[object_type], count])
    summary.append([])
    summary.append(["Можно привязать автоматически (пустые поля)", plan["counts"]["records"]])
    summary.column_dimensions["A"].width = 48

    def sheet(title: str, columns: list[tuple[str, str]], rows: list[dict[str, Any]]) -> None:
        ws = wb.create_sheet(title)
        ws.append([label for _, label in columns])
        for c in ws[1]:
            c.font = Font(bold=True)
        if not rows:
            ws.append(["Несоответствий нет"])
        for row in rows[:EXCEL_ROW_LIMIT]:
            ws.append([_cell(key, row.get(key)) for key, _ in columns])
        ws.freeze_panes = "A2"

    sheet("ЛЭП", LINE_COLUMNS, [i for i in result["items"] if i["object_type"] == "line"])
    sheet("Точечные объекты", POINT_COLUMNS, [i for i in result["items"] if i["object_type"] != "line"])
    ws = wb.create_sheet("Предлагаемая привязка")
    ws.append(["Вид", "ID", "Наименование", "Поле", "Было", "Станет", "Расстояние, м"])
    for c in ws[1]:
        c.font = Font(bold=True)
    for record in plan["records"]:
        for change in record["changes"]:
            ws.append([TYPE_TITLES[record["object_type"]], record["id"], record["name"], change["field"],
                       change["old"], change["new"], change["distance"]])
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
