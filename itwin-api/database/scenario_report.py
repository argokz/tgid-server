"""Excel-отчёт «сценарий против основы» по результату database/scenario_compare.py.

Листы: «Сводка» (расчёты, предупреждения, пороги, итоги по потребителям и участкам), «Потребители» и
«Участки» — существенные изменения режима: значение в основе, в сценарии и разница.
"""
from __future__ import annotations

import io
from datetime import datetime
from typing import Any, Optional

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from database.scenario_compare import CHANGE_KINDS

BOLD = Font(bold=True)
# изменений в отчёте — больше, чем в окне (там 300)
REPORT_ITEMS_LIMIT = 20000


def _delta(pair: Optional[list]) -> Optional[float]:
    if not pair or pair[0] is None or pair[1] is None:
        return None
    return round(pair[1] - pair[0], 3)


def _pair(item: dict[str, Any], key: str) -> list:
    pair = item.get(key) or [None, None]
    return [pair[0], pair[1], _delta(pair)]


def _calc_label(calc: dict[str, Any]) -> str:
    mode = "аварийный" if calc.get("emergency") else "плановый"
    name = f" «{calc['name']}»" if calc.get("name") else ""
    date = ""
    if calc.get("date"):
        try:
            date = f", {datetime.fromisoformat(calc['date']):%d.%m.%Y %H:%M}"
        except ValueError:
            date = f", {calc['date']}"
    return f"№{calc['id']}{name}, фрагмент {calc['fileid']}, Тн {calc['tn']:g} °C, {mode}{date}"


def _sheet(wb: Workbook, title: str, header: list[str], rows: list[list[Any]]) -> None:
    ws = wb.create_sheet(title)
    ws.append(header)
    for cell in ws[1]:
        cell.font = BOLD
    for row in rows:
        ws.append(row)
    ws.freeze_panes = "A2"
    for i, name in enumerate(header, start=1):
        ws.column_dimensions[get_column_letter(i)].width = max(12, min(40, len(name) + 4))


def workbook(result: dict[str, Any]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Сводка"
    ws.append([f"Сценарий «{result['name']}» против основы"])
    ws["A1"].font = Font(bold=True, size=13)
    ws.append(["Расчёт основы", _calc_label(result["base_calculation"])])
    ws.append(["Расчёт сценария", _calc_label(result["scenario_calculation"])])
    for warning in result.get("warnings", []):
        ws.append(["Внимание", warning])
    th = result["thresholds"]
    ws.append(["Пороги", f"недотоп — обеспеченность ниже {th['underheat_pct']:g} %; заметный напор — "
                         f"{th['head_delta_m']:g} м; удельные потери — {th['spec_loss_pa_m']:g} Па/м; "
                         f"скорость — {th['velocity_ms']:g} м/с"])
    ws.append([])
    ws.append(["Показатель", "Основа", "Сценарий", "Разница"])
    for cell in ws[ws.max_row]:
        cell.font = BOLD
    c, ln = result["consumers"], result["lines"]
    for label, a, b in (
        ("Потребителей рассчитано", c["base"], c["scenario"]),
        ("С недотопом", c["underheated_base"], c["underheated_scenario"]),
        ("Нагрузка, Гкал/ч", c["load_base"], c["load_scenario"]),
        ("Расход потребителей, т/ч", c["flow_base"], c["flow_scenario"]),
        ("Участков рассчитано", ln["base"], ln["scenario"]),
        ("Участков с потерями выше нормы", ln["overloaded_base"], ln["overloaded_scenario"]),
    ):
        ws.append([label, a, b, round(b - a, 3) if a is not None and b is not None else None])
    ws.append([])
    ws.append(["Изменения", "Потребители", "Участки"])
    for cell in ws[ws.max_row]:
        cell.font = BOLD
    for kind, label in CHANGE_KINDS.items():
        cons, lines = c.get(kind), ln.get(kind)
        if cons or lines:
            ws.append([label, cons or 0, lines or 0])
    if result.get("truncated"):
        ws.append([])
        ws.append([f"В листах — первые {len(result['items'])} изменений из {result['total']}"])
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 70

    consumers, lines_rows = [], []
    for item in result["items"]:
        ids = [item.get("base_id"), item.get("copy_id")]
        coords = [item.get("lat"), item.get("lng")]
        label = CHANGE_KINDS.get(item["kind"], item["kind"])
        if item["object"] == "consumer":
            consumers.append([label, item.get("name"), *ids, *_pair(item, "head"), *_pair(item, "supply"),
                              *_pair(item, "flow"), *coords])
        else:
            lines_rows.append([label, *ids, *_pair(item, "flow"), *_pair(item, "velocity"), *_pair(item, "loss"),
                               *coords])
    _sheet(wb, "Потребители", [
        "Изменение", "Узел", "Узел основы", "Узел сценария",
        "Напор основы, м", "Напор сценария, м", "Δ напора, м",
        "Обеспеченность основы, %", "Обеспеченность сценария, %", "Δ, п.п.",
        "Расход основы, т/ч", "Расход сценария, т/ч", "Δ расхода, т/ч", "Широта", "Долгота"], consumers)
    _sheet(wb, "Участки", [
        "Изменение", "Участок основы", "Участок сценария",
        "Расход основы, т/ч", "Расход сценария, т/ч", "Δ расхода, т/ч",
        "Скорость основы, м/с", "Скорость сценария, м/с", "Δ скорости, м/с",
        "Потери основы, Па/м", "Потери сценария, Па/м", "Δ потерь, Па/м", "Широта", "Долгота"], lines_rows)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
