"""План перекладок теплосети (фаза 5, шаг 2): что переложить в пределах бюджета, чтобы снять больше риска.

Польза перекладки участка = вероятность его повреждения за горизонт (прогноз аварийности, tgid_failure_risk) ×
последствия повреждения — по выбранному критерию (OBJECTIVES):
- consumers — сколько потребителей отключит авария на участке (зона задвижек + ниже по сети без обходного питания,
  database/outage_consequence.py, шаг 3); load — их нагрузка, Гкал/ч;
- wear — множитель оценки износа (wear_out.consequence из паспорта, 1…2,4) × множитель расхода (wear.flow_factor);
- failures — без последствий: просто ожидаемое число повреждений.
Новая труба за горизонт прогноза считается неповреждаемой — польза равна снятому ожидаемому ущербу.

Стоимость — длина × расценка за метр по классу диаметра (tgid_relay_price, заполняет пользователь); расценок нет —
бюджет в километрах. Единица перекладки — трасса (build_runs): участки одного фрагмента между развилками, подача и
обратка — один коридор; выбор по отдельным участкам давал тысячи обрезков по 2 м. Отбор — жадный по пользе на
единицу стоимости до исчерпания бюджета.

Копии сети друг на друге (в Астане — магистрали по отопительным сезонам, 11 фрагментов) — одна физическая труба:
участки разных фрагментов с совпадающими концами (сетка 1 м) и длиной склеиваются (merge_copies), иначе план
перекладывает одну трубу дважды.

Для сравнения — план того же бюджета «самые старые первыми»: сколько ожидаемых повреждений снял бы он. Ретроспектива
на реальных повреждениях — в проверке прогноза (metrics.capture_by_km).
"""
from __future__ import annotations

import io
import json
from collections import defaultdict
from typing import Any, Optional

from database import outage_consequence
from database.wear import flow_factor

MAX_SELECTED = 5000
JOB_OVERHEAD_M = 25.0      # постоянные затраты объекта (вскрытие, отключение, благоустройство) — как 25 м трубы
DEFAULT_CONSEQUENCE = 1.0
OBJECTIVES = {
    "consumers": "отключаемые потребители",
    "load": "отключаемая нагрузка",
    "wear": "последствия из оценки износа",
    "failures": "число повреждений",
}
OUTAGE_OBJECTIVES = ("consumers", "load")


class PlanError(Exception):
    def __init__(self, status: int, message: str):
        super().__init__(message)
        self.status = status
        self.message = message


# --- чистые функции ------------------------------------------------------------------------------

def price_for(dn_mm: Optional[float], prices: list[dict[str, Any]]) -> Optional[float]:
    """Расценка за метр: первый класс с dn ≤ dn_max (классы по возрастанию); диаметр больше всех — последний класс;
    диаметра нет — средний класс."""
    if not prices:
        return None
    if dn_mm is None:
        return prices[len(prices) // 2]["price_per_m"]
    for p in prices:
        if dn_mm <= p["dn_max_mm"]:
            return p["price_per_m"]
    return prices[-1]["price_per_m"]


def select(items: list[dict[str, Any]], budget: float, overhead: float = 0.0) -> list[dict[str, Any]]:
    """Жадный отбор по пользе на единицу стоимости, пока хватает бюджета (дорогие пропускаются, дешёвые — дальше).
    overhead — постоянные затраты объекта: входят в порядок отбора, бюджет тратится на сам объект. Без них
    модель выбирает обрезки по 2 м — вероятность повреждения почти не растёт с длиной, «риск на метр» у отводов
    огромный."""
    order = sorted((x for x in items if x["cost"] > 0 and x["benefit"] > 0),
                   key=lambda x: (-x["benefit"] / (x["cost"] + overhead), x["lineid"]))
    out, spent = [], 0.0
    for x in order:
        if spent + x["cost"] <= budget:
            out.append(x)
            spent += x["cost"]
            if len(out) >= MAX_SELECTED:
                break
    return out


def build_runs(items: list[dict[str, Any]]) -> list[dict[str, Any]]:
    """Трассы — единицы перекладки: участки одного фрагмента между развилками. Подача и обратка между одними узлами —
    один коридор; коридоры сцепляются через узел, у которого ровно два коридора одного диаметра (смена диаметра —
    граница трассы). Участок без узлов — трасса сам по себе. Польза и стоимость трассы — суммы её участков."""
    corridors: dict[Any, list[dict[str, Any]]] = defaultdict(list)
    for x in items:
        if x["node1"] is None or x["node2"] is None:
            corridors[("line", x["lineid"])].append(x)
        else:
            a, b = sorted((x["node1"], x["node2"]))
            corridors[(x["fileid"], a, b)].append(x)
    keys = list(corridors)
    dn = {k: max((x["dn"] or 0) for x in corridors[k]) for k in keys}
    at_node: dict[Any, list[Any]] = defaultdict(list)
    for k in keys:
        if k[0] != "line":
            at_node[(k[0], k[1])].append(k)
            at_node[(k[0], k[2])].append(k)
    parent = {k: k for k in keys}

    def find(k: Any) -> Any:
        while parent[k] != k:
            parent[k] = parent[parent[k]]
            k = parent[k]
        return k

    for ks in at_node.values():
        if len(ks) == 2 and dn[ks[0]] == dn[ks[1]]:
            parent[find(ks[0])] = find(ks[1])
    groups: dict[Any, list[dict[str, Any]]] = defaultdict(list)
    for k in keys:
        groups[find(k)].extend(corridors[k])
    runs = []
    for lines in groups.values():
        runs.append({"lines": lines, "lineid": min(x["lineid"] for x in lines), "fileid": lines[0]["fileid"],
                     "benefit": sum(x["benefit"] for x in lines), "cost": sum(x["cost"] for x in lines),
                     "prob": sum(x["prob"] for x in lines), "len_m": sum(x["len_m"] for x in lines)})
    return runs


def oldest_first(items: list[dict[str, Any]], budget: float) -> list[dict[str, Any]]:
    """План сравнения: самые старые участки первыми (возраст неизвестен — в конце) в том же бюджете."""
    order = sorted(items, key=lambda x: (x["age"] is None, -(x["age"] or 0), x["lineid"]))
    out, spent = [], 0.0
    for x in order:
        if x["cost"] > 0 and spent + x["cost"] <= budget:
            out.append(x)
            spent += x["cost"]
    return out


# --- расценки ------------------------------------------------------------------------------------

async def list_prices(conn) -> list[dict[str, Any]]:
    rows = await conn.fetch("SELECT * FROM tgid_relay_price ORDER BY dn_max_mm")
    return [{**dict(r), "updated_at": r["updated_at"].isoformat()} for r in rows]


async def set_prices(conn, rows: list[dict[str, Any]], user: str) -> list[dict[str, Any]]:
    """Расценки целиком: классы диаметра должны различаться."""
    if len({r["dn_max_mm"] for r in rows}) != len(rows):
        raise PlanError(422, "Классы диаметра повторяются")
    async with conn.transaction():
        await conn.execute("DELETE FROM tgid_relay_price")
        await conn.executemany(
            """INSERT INTO tgid_relay_price (dn_max_mm, price_per_m, currency, source, updated_by)
               VALUES ($1, $2, $3, $4, $5)""",
            [(r["dn_max_mm"], r["price_per_m"], r.get("currency") or "KZT", r.get("source"), user) for r in rows])
    return await list_prices(conn)


# --- план ----------------------------------------------------------------------------------------

def severity(objective: str, x: dict[str, Any]) -> float:
    """Последствия одного повреждения участка по критерию плана. Последствия отключения неизвестны (участок вне
    расчётной схемы) — 0: такой участок по этим критериям не выбирается."""
    if objective == "consumers":
        return float(x["consumers"] or 0)
    if objective == "load":
        return float(x["load"] or 0.0)
    if objective == "wear":
        return (x["consequence"] or DEFAULT_CONSEQUENCE) * flow_factor(x["flow_t_h"])
    return 1.0


def merge_copies(items: list[dict[str, Any]]) -> tuple[list[dict[str, Any]], int]:
    """Склеить копии одной трубы из разных фрагментов (ключ ends — концы и длина). Остаются участки фрагмента с
    наибольшим суммарным риском (трассы не рвутся между копиями); вероятность — наибольшая по копиям (попарно по
    убыванию: подача и обратка фрагмента — две трубы), последствия — наибольшие. Возвращает участки и число
    отброшенных копий."""
    weight: dict[Any, float] = defaultdict(float)
    for x in items:
        weight[x["fileid"]] += x["prob"]
    groups: dict[Any, list[dict[str, Any]]] = defaultdict(list)
    out = []
    for x in items:
        if x.get("ends") is None:
            out.append(x)
        else:
            groups[x["ends"]].append(x)
    dropped = 0
    for lines in groups.values():
        by_file: dict[Any, list[dict[str, Any]]] = defaultdict(list)
        for x in lines:
            by_file[x["fileid"]].append(x)
        if len(by_file) == 1:
            out.extend(lines)
            continue
        main = min(by_file, key=lambda f: (-weight[f], f if f is not None else 0))
        kept = sorted((dict(x) for x in by_file[main]), key=lambda x: -x["prob"])
        for f, others in by_file.items():
            if f == main:
                continue
            dropped += len(others)
            for k, o in zip(kept, sorted(others, key=lambda x: -x["prob"])):
                k["prob"] = max(k["prob"], o["prob"])
        for key in ("consumers", "load", "social"):
            values = [x[key] for x in lines if x[key] is not None]
            for k in kept:
                k[key] = max(values) if values else None
        out.extend(kept)
    return out, dropped


async def _items(conn, run_id: int, fileid: Optional[int]) -> list[dict[str, Any]]:
    rows = await conn.fetch(
        """SELECT r.lineid, r.prob, r.age, r.len_m, l.fileid, f.name AS fragment, l.nodeid1, l.nodeid2,
                  h.diametercondit AS dn, w.consequence, w.flow_t_h, w.category,
                  round(ST_X(ST_StartPoint(ST_GeometryN(l.shape, 1)))) AS sx,
                  round(ST_Y(ST_StartPoint(ST_GeometryN(l.shape, 1)))) AS sy,
                  round(ST_X(ST_EndPoint(ST_GeometryN(l.shape, ST_NumGeometries(l.shape))))) AS ex,
                  round(ST_Y(ST_EndPoint(ST_GeometryN(l.shape, ST_NumGeometries(l.shape))))) AS ey,
                  c.zone_consumers + c.down_consumers AS consumers, c.zone_load + c.down_load AS load,
                  c.zone_social + c.down_social AS social
             FROM tgid_failure_risk r JOIN linesobj l ON l.id = r.lineid
             LEFT JOIN tgid_line_consequence c ON c.lineid = r.lineid
             LEFT JOIN fragments f ON f.id = l.fileid
             LEFT JOIN LATERAL (SELECT diametercondit FROM heatpipesections WHERE lineid = l.id ORDER BY id LIMIT 1) h
                    ON true
             LEFT JOIN wear_out w ON w.lineid = r.lineid
                   AND w.calculationid = (SELECT max(calculationid) FROM wear_out)
            WHERE r.run_id = $1 AND ($2::int IS NULL OR l.fileid = $2)""", run_id, fileid)
    out = []
    for r in rows:
        ends = None
        if None not in (r["sx"], r["sy"], r["ex"], r["ey"]):
            ends = (*sorted([(r["sx"], r["sy"]), (r["ex"], r["ey"])]), round(r["len_m"] or 0))
        out.append({"lineid": r["lineid"], "prob": r["prob"], "age": r["age"], "len_m": r["len_m"] or 0.0,
                    "fileid": r["fileid"], "fragment": r["fragment"], "node1": r["nodeid1"], "node2": r["nodeid2"],
                    "dn": r["dn"], "wear_category": r["category"], "consequence": r["consequence"],
                    "flow_t_h": r["flow_t_h"], "consumers": r["consumers"], "load": r["load"], "social": r["social"],
                    "ends": ends})
    return out


def expected(lines: list[dict[str, Any]], totals: dict[str, float]) -> dict[str, Any]:
    """Снятый ожидаемый ущерб: повреждения, отключения потребителей, нагрузки, социальных объектов — и доли."""
    out: dict[str, Any] = {"lines": len(lines)}
    for key, field in (("failures", None), ("consumers", "consumers"), ("load", "load"), ("social", "social")):
        v = sum(x["prob"] * (1 if field is None else (x[field] or 0)) for x in lines)
        out[f"expected_{key}"] = round(v, 2)
        out[f"expected_{key}_share"] = round(v / totals[key], 3) if totals[key] else None
    return out


async def plan(conn, *, budget: float, unit: str = "km", run_id: Optional[int] = None,
               fileid: Optional[int] = None, job_overhead_m: float = JOB_OVERHEAD_M,
               objective: str = "consumers", with_lines: bool = False) -> dict[str, Any]:
    """План перекладок в пределах бюджета: объекты, участки, снятый ожидаемый ущерб и сравнение со «старыми».
    with_lines — ещё и все выбранные участки (выгрузка в Excel)."""
    if unit not in ("km", "money"):
        raise PlanError(422, "Единица бюджета — km или money")
    if objective not in OBJECTIVES:
        raise PlanError(422, f"Критерий плана — один из: {', '.join(OBJECTIVES)}")
    if budget <= 0:
        raise PlanError(422, "Бюджет — больше нуля")
    if run_id is None:
        run_id = await conn.fetchval("SELECT max(id) FROM tgid_failure_run")
        if run_id is None:
            raise PlanError(422, "Нет прогноза аварийности — сначала рассчитайте его")
    run = await conn.fetchrow("SELECT horizon_years, cutoffs, metrics FROM tgid_failure_run WHERE id = $1", run_id)
    if run is None:
        raise PlanError(404, "Расчёт прогноза не найден")
    prices = await list_prices(conn)
    if unit == "money" and not prices:
        raise PlanError(422, "Расценки перекладки не заполнены — план в деньгах не составить, задайте бюджет в км")
    if objective in OUTAGE_OBJECTIVES:
        await outage_consequence.ensure(conn)
    items, copies = merge_copies(await _items(conn, run_id, fileid))
    for x in items:
        x["severity"] = severity(objective, x)
        x["benefit"] = x["prob"] * x["severity"]
        if unit == "km":
            x["cost"] = x["len_m"] / 1000
        else:
            x["cost"] = x["len_m"] * price_for(x["dn"], prices)
    runs = build_runs(items)
    # постоянные затраты объекта — в единицах бюджета: км или метры × средняя расценка
    if unit == "km":
        overhead = job_overhead_m / 1000
    else:
        overhead = job_overhead_m * sum(p["price_per_m"] for p in prices) / len(prices)
    objects = sorted((r["lines"] for r in select(runs, budget, overhead)), key=lambda g: -sum(x["benefit"] for x in g))
    chosen = [x for g in objects for x in g]
    total_benefit = sum(x["benefit"] for x in items) or 1.0
    totals = {"failures": sum(x["prob"] for x in items),
              "consumers": sum(x["prob"] * (x["consumers"] or 0) for x in items),
              "load": sum(x["prob"] * (x["load"] or 0) for x in items),
              "social": sum(x["prob"] * (x["social"] or 0) for x in items)}
    old = oldest_first(items, budget)
    geo = await conn.fetch(
        "SELECT id, ST_AsGeoJSON(ST_Transform(shape, 4326), 7) AS g FROM linesobj WHERE id = ANY($1::int[])",
        [x["lineid"] for x in chosen])
    geom = {r["id"]: r["g"] for r in geo}
    rank = {x["lineid"]: i for i, g in enumerate(objects) for x in g}
    features = [{"type": "Feature", "geometry": json.loads(geom[x["lineid"]]),
                 "properties": {"lineid": x["lineid"], "object": rank[x["lineid"]] + 1, "prob": x["prob"]}}
                for x in chosen if geom.get(x["lineid"])]

    def obj(i: int, g: list[dict[str, Any]]) -> dict[str, Any]:
        dns = [x["dn"] for x in g if x["dn"]]
        return {"n": i + 1, "fragment": g[0]["fragment"], "fileid": g[0]["fileid"], "lines": len(g),
                "lineids": [x["lineid"] for x in g][:50], "length_m": round(sum(x["len_m"] for x in g), 1),
                "cost": round(sum(x["cost"] for x in g), 3), "expected_failures": round(sum(x["prob"] for x in g), 3),
                "benefit": round(sum(x["benefit"] for x in g), 3), "max_prob": round(max(x["prob"] for x in g), 4),
                "dn_min": min(dns) if dns else None, "dn_max": max(dns) if dns else None,
                # худшее отключение при повреждении одного участка трассы
                "consumers_max": max((x["consumers"] or 0) for x in g),
                "load_max": round(max((x["load"] or 0) for x in g), 3),
                "social_max": max((x["social"] or 0) for x in g)}

    metrics = json.loads(run["metrics"])
    out = {
        "run_id": run_id, "horizon_years": run["horizon_years"], "unit": unit, "budget": budget,
        "objective": objective, "objective_label": OBJECTIVES[objective],
        "job_overhead_m": job_overhead_m,
        "currency": prices[0]["currency"] if unit == "money" and prices else None,
        "spent": round(sum(x["cost"] for x in chosen), 3),
        "length_km": round(sum(x["len_m"] for x in chosen) / 1000, 3),
        **expected(chosen, totals),
        "benefit_share": round(sum(x["benefit"] for x in chosen) / total_benefit, 3),
        "oldest_first": expected(old, totals),
        "consequence_values": len({x["consequence"] for x in items if x["consequence"] is not None}),
        "flow_known": sum(1 for x in items if x["flow_t_h"] is not None),
        "outage_known": sum(1 for x in items if x["consumers"] is not None),
        "outage": await outage_consequence.status(conn),
        "runs": len(runs), "population": len(items), "copies_merged": copies, "backtest_capture_by_km": metrics.get("capture_by_km"),
        "objects": [obj(i, g) for i, g in enumerate(objects)],
        "geojson": {"type": "FeatureCollection", "features": features},
    }
    if with_lines:
        out["chosen_lines"] = [{**{k: x[k] for k in ("lineid", "fragment", "len_m", "dn", "prob", "age", "consumers",
                                                     "load", "social", "cost")},
                                "object": rank[x["lineid"]] + 1} for x in chosen]
    return out



# --- выгрузка ------------------------------------------------------------------------------------

def _pct(v: Optional[float]) -> str:
    return "—" if v is None else f"{v * 100:.1f} %".replace(".", ",")


def plan_workbook(p: dict[str, Any]) -> bytes:
    """Excel для инвестпрограммы: итоги и параметры, объекты перекладки, участки объектов."""
    from openpyxl import Workbook
    from openpyxl.styles import Font

    bold = Font(bold=True)
    money = p["unit"] == "money"
    cost_label = f"Стоимость, {p['currency']}" if money else "Длина в бюджете, км"
    wb = Workbook()
    ws = wb.active
    ws.title = "Итоги"
    ws.append([f"План перекладок: бюджет {p['budget']} {p['currency'] if money else 'км'}"])
    ws["A1"].font = Font(bold=True, size=13)
    ws.append(["Критерий отбора", p["objective_label"]])
    ws.append(["Горизонт прогноза, лет", p["horizon_years"]])
    ws.append(["Постоянные затраты объекта, м трубы", p["job_overhead_m"]])
    ws.append(["Расчёт прогноза", p["run_id"]])
    ws.append([])
    ws.append(["", "По плану", "«Самые старые первыми»"])
    for cell in ws[ws.max_row]:
        cell.font = bold
    old = p["oldest_first"]
    ws.append(["Участков", p["lines"], old["lines"]])
    for key, label in (("failures", "Ожидаемых повреждений снято"),
                       ("consumers", "Ожидаемых отключений потребителей снято"),
                       ("load", "Ожидаемой отключаемой нагрузки снято, Гкал/ч"),
                       ("social", "Ожидаемых отключений социальных объектов снято")):
        ws.append([label, p[f"expected_{key}"], old[f"expected_{key}"]])
        ws.append(["  доля от всей сети в расчёте", _pct(p[f"expected_{key}_share"]),
                   _pct(old[f"expected_{key}_share"])])
    ws.append([])
    ws.append(["Оценка по топологии: отключение — зона задвижек и потребители ниже по сети без обходного питания; "
               "вероятность — прогноз аварийности за горизонт."])
    ws.column_dimensions["A"].width = 52
    ws.column_dimensions["B"].width = 18
    ws.column_dimensions["C"].width = 24

    objs = wb.create_sheet("Объекты")
    objs.append(["№", "Фрагмент", "Участков", "Длина, м", "Ø от, мм", "Ø до, мм", "Ожид. повреждений",
                "Отключит потребителей (худший участок)", "Нагрузка, Гкал/ч", "Соцобъектов", cost_label])
    for cell in objs[1]:
        cell.font = bold
    for o in p["objects"]:
        objs.append([o["n"], o["fragment"], o["lines"], o["length_m"], o["dn_min"], o["dn_max"], o["expected_failures"],
                    o["consumers_max"], o["load_max"], o["social_max"], o["cost"]])
    objs.column_dimensions["B"].width = 36

    ls = wb.create_sheet("Участки")
    ls.append(["Объект", "Участок (id)", "Фрагмент", "Длина, м", "Ø, мм", "Возраст, лет", "Вероятность повреждения",
               "Отключит потребителей", "Нагрузка, Гкал/ч", "Соцобъектов", cost_label])
    for cell in ls[1]:
        cell.font = bold
    for x in p.get("chosen_lines") or []:
        ls.append([x["object"], x["lineid"], x["fragment"], round(x["len_m"], 1), x["dn"],
                   None if x["age"] is None else round(x["age"], 1), round(x["prob"], 4), x["consumers"],
                   None if x["load"] is None else round(x["load"], 3), x["social"], round(x["cost"], 3)])
    ls.column_dimensions["C"].width = 36
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
