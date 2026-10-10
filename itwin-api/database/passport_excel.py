"""Excel-паспорт участка (перенос gid8 passport_module/p.py::passport) — общий код для синхронного
эндпоинта POST /api/db/object/{table}/{obj_id} и фоновой задачи Celery (file_jobs, kind=passport).

Синхронный код (psycopg2 + openpyxl + ThreadPoolExecutor форм f1…f15): из async-кода вызывать
через поток (run_in_threadpool / asyncio.to_thread).
"""

from __future__ import annotations

import io
import os
import sys
from typing import Iterable, Optional

import psycopg2

PASSPORT_TABLES = ("linesobj", "nodes", "uchastok_ms", "uchastok_rs")
XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


class PassportError(Exception):
    """Ошибка входных данных паспорта с HTTP-кодом для эндпоинта / сообщением для задачи."""

    def __init__(self, status_code: int, detail: str):
        super().__init__(detail)
        self.status_code = status_code
        self.detail = detail


def _passport_modules():
    # passport_module — перенесённый из gid8 код с плоскими импортами
    # (import config, import connect …). Чтобы они разрешались, каталог
    # модуля должен быть на sys.path.
    module_dir = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "passport_module")
    if module_dir not in sys.path:
        sys.path.insert(0, module_dir)
    from passport_module.p import async_do_passport, read_ms_rs
    import connect as passport_connect
    import db2 as passport_db2
    import ms1 as passport_ms1
    import rs1 as passport_rs1
    import sort_graph as passport_sort_graph
    import sql_pass as passport_sql_pass

    return (async_do_passport, read_ms_rs, passport_connect, passport_db2, passport_ms1, passport_rs1,
            passport_sort_graph, passport_sql_pass)


def resolve_passport_site(cur, table: str, obj_id: int) -> tuple[str, int]:
    """Участок (ms/rs, id) объекта: труба → узел 1, узел → belong*site или heatpipesections."""
    line_id = None
    if table == "linesobj":
        cur.execute("SELECT nodeid1 FROM linesobj WHERE id = %s", (obj_id,))
        res = cur.fetchone()
        if not res:
            raise PassportError(404, "Труба не найдена")
        node_id = res[0]
        line_id = obj_id
    elif table == "nodes":
        node_id = obj_id
    elif table == "uchastok_ms":
        return "ms", obj_id
    elif table == "uchastok_rs":
        return "rs", obj_id
    else:
        raise PassportError(400, "Неизвестная таблица")

    from database.passport_site import resolve_site_from_node_row, resolve_site_via_heatpipesections

    cur.execute("SELECT belongmagistralsite, belongdistsite FROM nodes WHERE id = %s", (node_id,))
    res = cur.fetchone()
    if not res:
        raise PassportError(404, "Узел не найден")
    resolved = resolve_site_from_node_row(res[0], res[1])
    if resolved is None:
        resolved = resolve_site_via_heatpipesections(cur, node_id, line_id=line_id)
    if resolved is None:
        raise PassportError(
            404,
            f"Узел {node_id} не привязан ни к магистральному, ни к распределительному "
            "участку — паспорт формируется только по участку. "
            "Запустите scripts/sql/backfill_belong_site.sql на копии БД "
            "или откройте паспорт по uchastok_ms / uchastok_rs.",
        )
    return resolved


def fragments_arg(fragments: Optional[Iterable[int]]) -> str:
    """Фрагменты для форм паспорта: «1,2,3» — как ``-fragments`` десктопа (подключённые фрагменты);
    пусто — все фрагменты базы."""
    return ",".join(str(int(f)) for f in sorted(set(fragments or [])))


def build_passport_xlsx(table: str, obj_id: int, fragments: Optional[Iterable[int]] = None) -> tuple[bytes, str]:
    """Полный паспорт участка объекта: (содержимое .xlsx, имя файла). PassportError — ошибки данных.

    fragments — фрагменты карты: десктоп строит паспорт по подключённым фрагментам
    (``GidWidget::Passport`` → ``-fragments m_par``). Без них в базе с фрагментами-вариантами
    (копии магистралей в Астане) трубы участка попадают в паспорт по несколько раз.
    """
    import openpyxl

    (async_do_passport, read_ms_rs, passport_connect, passport_db2, passport_ms1, passport_rs1,
     passport_sort_graph, passport_sql_pass) = _passport_modules()

    # Формы паспорта (f1…f15) сами открывают соединение через
    # connect.connect(**c), поэтому им передаются ПАРАМЕТРЫ подключения,
    # а не готовый объект соединения (иначе TypeError на **c).
    passport_conn_params = {
        "rdbms": "postgreSQL",
        "server": os.getenv("DB_HOST"),
        "user": os.getenv("DB_USER"),
        "password": os.getenv("DB_PASSWORD"),
        "db": os.getenv("DB_NAME"),
        "port": os.getenv("DB_PORT"),
    }

    # Отдельное соединение для определения участка объекта
    conn = psycopg2.connect(
        host=os.getenv("DB_HOST"),
        database=os.getenv("DB_NAME"),
        user=os.getenv("DB_USER"),
        password=os.getenv("DB_PASSWORD"),
        port=os.getenv("DB_PORT"),
    )
    try:
        ms_rs, site_id = resolve_passport_site(conn.cursor(), table, obj_id)
    finally:
        conn.close()

    # Полный сценарий desktop-паспорта (passport_module/p.py::passport):
    # титульный лист участка → состав участка → граф участка (marked lines) →
    # 15 форм. Без make_graph формам приходил mark_line=0 и SQL падал.
    fragments = fragments_arg(fragments)
    passport_conn = passport_connect.connect(**passport_conn_params)
    try:
        wb = openpyxl.Workbook()
        title_sheet = wb.active
        title_sheet.title = "Паспорт"

        head_rows = passport_db2.read_q(passport_conn, passport_sql_pass.passport(ms_rs, site_id))
        if ms_rs == "ms":
            passport_ms1.write_ms(title_sheet, head_rows)
        else:
            passport_rs1.write_rs(title_sheet, head_rows)

        site_rows = read_ms_rs(passport_conn, ms_rs, site_id)
        graph = passport_sort_graph.make_graph(passport_conn, fragments, ms_rs, site_id)
    finally:
        passport_conn.close()

    if not graph or not graph[0]:
        raise PassportError(
            404,
            f"Для участка {ms_rs}/{site_id} нет ни одного трубопровода. "
            "Привяжите трубы к участку (инструмент «Участки ПТС», "
            "heatpipesections.magistralSite / distSite) — без неё паспорт сформировать нельзя.",
        )

    mark_line, mark_node, mark_pts = graph

    # Формы f1…f15 открывают собственные соединения из параметров (ThreadPoolExecutor)
    async_do_passport(
        c=passport_conn_params,
        wb=wb,
        ms_rs=ms_rs,
        id=site_id,
        fragments=fragments,
        mark_line=mark_line,
        mark_pts=mark_pts,
        mark_node=mark_node,
        vals=site_rows,
    )

    output = io.BytesIO()
    wb.save(output)
    return output.getvalue(), f"Passport_{ms_rs}_{site_id}.xlsx"


# --- паспорта по начальнику участка --------------------------------------------------------
#
# Десктоп строит паспорт одного участка МС/РС (док ПТС → «Паспорт»: GidWidget::Passport,
# passport_ps/p.py -type ms|rs -id N); в дереве дока участки сгруппированы по начальникам.
# Паспорта по начальнику — те же паспорта всех его участков (формат десктопа без изменений),
# одним архивом с перечнем: что вошло и что пропущено.

KIND_TITLE = {"ms": "МС", "rs": "РС"}
SITE_NAME_COLUMN = {"ms": "opisanie_uchastka_ms", "rs": "naimenovanie_uchastka_rs"}
_BAD_FILENAME_CHARS = str.maketrans({c: "_" for c in '\\/:*?"<>|\r\n\t'})


def safe_filename(name: str, limit: int = 120) -> str:
    cleaned = " ".join(str(name).translate(_BAD_FILENAME_CHARS).split()).strip(" .")
    return (cleaned[:limit].rstrip(" .") or "file")


def chief_sites(cur, nach_id: int, kinds: Iterable[str], fragments: Optional[Iterable[int]]):
    """ФИО начальника и его участки (порядок дока ПТС) с числом труб в выбранных фрагментах —
    отбор как в sort_graph.make_graph (трубы вне внутренних схем, фрагмент по узлу 1)."""
    cur.execute("SELECT fio FROM nachalniki_uchastkov WHERE id = %s", (nach_id,))
    row = cur.fetchone()
    if not row:
        raise PassportError(404, f"Начальник участка {nach_id} не найден")
    fio = (row[0] or "").strip() or f"Начальник {nach_id}"
    frags = sorted(set(int(f) for f in fragments or []))
    sites = []
    for kind in ("ms", "rs"):
        if kind not in kinds:
            continue
        table, name_col = f"uchastok_{kind}", SITE_NAME_COLUMN[kind]
        pipe_col = "magistralsite" if kind == "ms" else "distsite"
        cur.execute(
            f"""
            WITH s AS (
                SELECT s.id, s.{name_col} AS name
                  FROM {table} s
                  JOIN uchastki_ekspluatatsii ue ON ue.id = s.nomer_uchastka
                 WHERE ue.nachalnik_uchastka = %s
            ), p AS (
                SELECT h.{pipe_col} AS site_id, count(*) AS pipes
                  FROM heatpipesections h
                  JOIN linesobj l ON l.id = h.lineid AND l.removed = 0
                  JOIN nodes n1 ON n1.id = l.nodeid1 AND n1.removed = 0 AND n1.internalnodeid IS NULL
                  JOIN nodes n2 ON n2.id = l.nodeid2 AND n2.removed = 0
                 WHERE h.{pipe_col} IN (SELECT id FROM s)
                   AND (cardinality(%s::int[]) = 0 OR n1.fileid = ANY(%s::int[]))
                 GROUP BY 1
            )
            SELECT s.id, s.name, COALESCE(p.pipes, 0) FROM s LEFT JOIN p ON p.site_id = s.id
             ORDER BY s.name, s.id
            """,
            (nach_id, frags, frags),
        )
        sites += [(kind, sid, (name or "").strip(), int(pipes)) for sid, name, pipes in cur.fetchall()]
    return fio, sites


def _index_workbook(fio: str, fragments: str, rows: list[tuple]) -> bytes:
    import openpyxl
    from openpyxl.styles import Font

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Перечень участков"
    ws.append([f"Паспорта участков начальника участка {fio}"])
    ws["A1"].font = Font(bold=True, size=12)
    ws.append([f"Фрагменты: {fragments or 'все'}"])
    ws.append([])
    ws.append(["№", "Вид", "Участок", "Наименование", "Труб", "Файл паспорта / причина"])
    for cell in ws[4]:
        cell.font = Font(bold=True)
    for i, row in enumerate(rows, 1):
        ws.append([i, *row])
    for col, width in zip("ABCDEF", (5, 6, 9, 60, 7, 70)):
        ws.column_dimensions[col].width = width
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def build_chief_passports_zip(nach_id: int, kinds: Iterable[str] = ("ms", "rs"),
                              fragments: Optional[Iterable[int]] = None,
                              progress=None) -> tuple[bytes, str]:
    """Паспорта всех участков начальника (МС и/или РС) — ZIP: по .xlsx на участок + перечень.

    Участок без труб в выбранных фрагментах пропускается (десктоп: «Нет участков»); ошибка одного
    паспорта не останавливает остальные — попадает в перечень. progress(message) — ход работы.
    """
    import zipfile

    kinds = tuple(k for k in ("ms", "rs") if k in set(kinds))
    conn = psycopg2.connect(
        host=os.getenv("DB_HOST"), database=os.getenv("DB_NAME"), user=os.getenv("DB_USER"),
        password=os.getenv("DB_PASSWORD"), port=os.getenv("DB_PORT"),
    )
    try:
        fio, sites = chief_sites(conn.cursor(), nach_id, kinds, fragments)
    finally:
        conn.close()
    if not sites:
        raise PassportError(404, f"У начальника участка {fio} нет участков "
                                 f"{' и '.join(KIND_TITLE[k] for k in kinds)}")
    if not any(pipes for *_, pipes in sites):
        raise PassportError(404, f"У участков начальника {fio} нет труб в выбранных фрагментах — "
                                 "подключите фрагменты или привяжите трубы к участкам (инструмент «Участки ПТС»)")

    buf = io.BytesIO()
    index_rows, built, used = [], 0, set()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zf:
        for n, (kind, site_id, name, pipes) in enumerate(sites, 1):
            title = f"{KIND_TITLE[kind]} {site_id}"
            if not pipes:
                index_rows.append((KIND_TITLE[kind], site_id, name, 0, "пропущен: нет труб в выбранных фрагментах"))
                continue
            if progress:
                progress(f"Участок {n} из {len(sites)}: {title}")
            try:
                content, _ = build_passport_xlsx(f"uchastok_{kind}", site_id, fragments)
            except PassportError as e:
                index_rows.append((KIND_TITLE[kind], site_id, name, pipes, f"не сформирован: {e.detail}"))
                continue
            except Exception as e:  # noqa: BLE001 — один участок не должен ронять весь архив
                index_rows.append((KIND_TITLE[kind], site_id, name, pipes, f"ошибка: {e}"))
                continue
            fname = safe_filename(f"{title} — {name}" if name else title) + ".xlsx"
            if fname in used:
                fname = safe_filename(f"{title} ({n})") + ".xlsx"
            used.add(fname)
            zf.writestr(fname, content)
            index_rows.append((KIND_TITLE[kind], site_id, name, pipes, fname))
            built += 1
        zf.writestr("Перечень участков.xlsx", _index_workbook(fio, fragments_arg(fragments), index_rows))
    if not built:
        raise PassportError(500, f"Ни один паспорт участков начальника {fio} не сформирован: "
                                 + "; ".join(f"{k} {i}: {r}" for k, i, _, _, r in index_rows if r.startswith(("не", "ош")))[:900])
    return buf.getvalue(), safe_filename(f"Паспорта участков — {fio}") + ".zip"
