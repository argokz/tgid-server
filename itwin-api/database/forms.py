"""Формы обходов (фаза 2, п. 8): конструктор форм осмотра, проверка ответов, импорт XLSForm, выгрузка.

Форма — список полей (подмножество XLSForm: https://xlsform.org): текст, целое, дробное, выбор одного, выбор
нескольких, дата, да/нет (acknowledge), пояснение (note), фото (image). У поля — обязательность, подсказка,
варианты, границы чисел и условие показа («показать, если поле = значение»). Правка формы — новая версия
(tgid_form_version); заполнение ссылается на версию и проверяется по ней. Заполняет бригада в полевом приложении
— раз на наряд (scope order) или на каждом узле маршрута обхода (scope point).

Ответ поля «фото» — список client_id фото наряда (tgid_work_order_photo): фото уходят из очереди телефона как
обычные фото наряда.
"""
from __future__ import annotations

import io
import json
import re
from datetime import date, datetime
from typing import Any, Literal, Optional
from uuid import UUID

from pydantic import BaseModel, Field, field_validator, model_validator

FieldType = Literal["text", "integer", "decimal", "select_one", "select_multiple", "date", "acknowledge", "note",
                    "image"]
TYPE_LABELS: dict[str, str] = {
    "text": "Текст", "integer": "Целое число", "decimal": "Число", "select_one": "Выбор одного",
    "select_multiple": "Выбор нескольких", "date": "Дата", "acknowledge": "Да / отметка", "note": "Пояснение",
    "image": "Фото",
}
MAX_FIELDS = 100
MAX_CHOICES = 200
MAX_TEXT = 2000
MAX_PHOTOS = 5
_NAME = re.compile(r"^[a-z][a-z0-9_]{0,39}$")


class FormError(Exception):
    def __init__(self, status: int, message: str):
        super().__init__(message)
        self.status = status
        self.message = message


# --- описание формы ----------------------------------------------------------------------------

class Choice(BaseModel):
    value: str = Field(min_length=1, max_length=100)
    label: str = Field(min_length=1, max_length=200)


class Relevant(BaseModel):
    """Показать поле, если ответ поля field равен / не равен value (у выбора нескольких — содержит value).
    Пока на поле field не ответили, поле скрыто при любом условии."""
    field: str
    op: Literal["=", "!=", "selected"] = "="
    value: str = Field(max_length=100)


class FormField(BaseModel):
    name: str
    type: FieldType
    label: str = Field(min_length=1, max_length=300)
    hint: str = Field("", max_length=500)
    required: bool = False
    choices: Optional[list[Choice]] = Field(None, max_length=MAX_CHOICES)
    min: Optional[float] = None
    max: Optional[float] = None
    relevant: Optional[Relevant] = None

    @field_validator("name")
    @classmethod
    def _name(cls, v: str) -> str:
        if not _NAME.match(v):
            raise ValueError(f"имя поля «{v}» — латиница, цифры и _, с буквы, до 40 символов")
        return v

    @model_validator(mode="after")
    def _shape(self) -> "FormField":
        if self.type in ("select_one", "select_multiple"):
            if not self.choices:
                raise ValueError(f"у поля «{self.label}» нет вариантов")
            values = [c.value for c in self.choices]
            if len(set(values)) != len(values):
                raise ValueError(f"у поля «{self.label}» повторяются значения вариантов")
        else:
            self.choices = None
        if self.type not in ("integer", "decimal"):
            self.min = self.max = None
        elif self.min is not None and self.max is not None and self.min > self.max:
            raise ValueError(f"у поля «{self.label}» минимум больше максимума")
        if self.type == "note":
            self.required = False
        return self


class FormSchema(BaseModel):
    fields: list[FormField] = Field(min_length=1, max_length=MAX_FIELDS)

    @model_validator(mode="after")
    def _refs(self) -> "FormSchema":
        seen: dict[str, FormField] = {}
        for f in self.fields:
            if f.name in seen:
                raise ValueError(f"имя поля «{f.name}» повторяется")
            if f.relevant:
                ref = seen.get(f.relevant.field)
                if ref is None:
                    raise ValueError(f"условие поля «{f.label}» ссылается на поле «{f.relevant.field}», которого нет выше")
                if ref.type in ("note", "image"):
                    raise ValueError(f"условие поля «{f.label}»: по полю «{ref.label}» условие не задаётся")
            seen[f.name] = f
        if all(f.type == "note" for f in self.fields):
            raise ValueError("в форме нет ни одного поля для ответа")
        return self


class FormIn(BaseModel):
    name: str = Field(min_length=1, max_length=200)
    description: str = Field("", max_length=2000)
    scope: Literal["order", "point"] = "order"
    schema_: FormSchema = Field(alias="schema")

    model_config = {"populate_by_name": True}

    @field_validator("name")
    @classmethod
    def _strip(cls, v: str) -> str:
        v = v.strip()
        if not v:
            raise ValueError("название формы пустое")
        return v


class FormPatch(BaseModel):
    active: bool


# --- проверка ответов --------------------------------------------------------------------------

def _empty(v: Any) -> bool:
    return v is None or v == "" or v == [] or v is False


def is_relevant(field: FormField, answers: dict[str, Any]) -> bool:
    r = field.relevant
    if r is None:
        return True
    got = answers.get(r.field)
    if _empty(got):
        return False  # на поле-условие ещё не ответили — зависимое поле скрыто (при любом условии)
    if r.op == "selected":
        return isinstance(got, list) and r.value in got
    text = "" if got is None else ("1" if got is True else "0" if got is False else str(got))
    return (text == r.value) if r.op == "=" else (text != r.value)


def _number(f: FormField, v: Any, errors: dict[str, str]) -> Optional[float]:
    if isinstance(v, bool) or not isinstance(v, (int, float, str)):
        errors[f.name] = "нужно число"
        return None
    try:
        num = float(str(v).replace(",", ".")) if isinstance(v, str) else float(v)
    except ValueError:
        errors[f.name] = "нужно число"
        return None
    if f.type == "integer" and num != int(num):
        errors[f.name] = "нужно целое число"
        return None
    if f.min is not None and num < f.min:
        errors[f.name] = f"не меньше {f.min:g}"
    elif f.max is not None and num > f.max:
        errors[f.name] = f"не больше {f.max:g}"
    return int(num) if f.type == "integer" else num


def validate_answers(schema: FormSchema, answers: dict[str, Any]) -> tuple[dict[str, Any], dict[str, str]]:
    """Ответы по схеме: (очищенные ответы, ошибки по именам полей). Поля, скрытые условием, отбрасываются."""
    clean: dict[str, Any] = {}
    errors: dict[str, str] = {}
    for f in schema.fields:
        if f.type == "note" or not is_relevant(f, clean):
            continue
        v = answers.get(f.name)
        if _empty(v):
            if f.required:
                errors[f.name] = "отметьте" if f.type == "acknowledge" else "обязательное поле"
            continue
        if f.type == "text":
            if not isinstance(v, str):
                errors[f.name] = "нужен текст"
                continue
            clean[f.name] = v.strip()[:MAX_TEXT]
        elif f.type in ("integer", "decimal"):
            num = _number(f, v, errors)
            if num is not None and f.name not in errors:
                clean[f.name] = num
        elif f.type == "select_one":
            values = {c.value for c in f.choices or ()}
            if str(v) not in values:
                errors[f.name] = "нет такого варианта"
                continue
            clean[f.name] = str(v)
        elif f.type == "select_multiple":
            values = [c.value for c in f.choices or ()]
            if not isinstance(v, list) or any(str(x) not in values for x in v):
                errors[f.name] = "нет такого варианта"
                continue
            picked = {str(x) for x in v}
            clean[f.name] = [x for x in values if x in picked]
        elif f.type == "date":
            try:
                clean[f.name] = date.fromisoformat(str(v)[:10]).isoformat()
            except ValueError:
                errors[f.name] = "дата — ГГГГ-ММ-ДД"
        elif f.type == "acknowledge":
            if v is not True:
                errors[f.name] = "нужно да или пусто"
                continue
            clean[f.name] = True
        elif f.type == "image":
            try:
                ids = [str(UUID(str(x))) for x in v] if isinstance(v, list) else None
            except ValueError:
                ids = None
            if ids is None or len(ids) > MAX_PHOTOS:
                errors[f.name] = f"фото — до {MAX_PHOTOS}"
                continue
            clean[f.name] = ids
    return clean, errors


# --- XLSForm -----------------------------------------------------------------------------------

_XLS_TYPES = {
    "text": "text", "string": "text", "integer": "integer", "int": "integer", "decimal": "decimal",
    "date": "date", "note": "note", "acknowledge": "acknowledge", "image": "image", "photo": "image",
}
_REL_EQ = re.compile(r"^\$\{(\w+)\}\s*(=|!=)\s*'([^']*)'$")
_REL_SEL = re.compile(r"^selected\(\s*\$\{(\w+)\}\s*,\s*'([^']*)'\s*\)$")
_CON_PART = re.compile(r"^\.\s*(>=|<=|>|<)\s*(-?\d+(?:\.\d+)?)$")


def _cell(v: Any) -> str:
    return "" if v is None else str(v).strip()


def _label_column(header: list[str]) -> Optional[str]:
    """label или label::Русский (ru) — русская колонка, если есть, иначе первая label*."""
    labels = [h for h in header if h == "label" or h.startswith("label::")]
    ru = [h for h in labels if "(ru)" in h.lower() or "русск" in h.lower()]
    return (ru or labels or [None])[0]


def _rows(ws) -> list[dict[str, str]]:
    it = ws.iter_rows(values_only=True)
    header = [_cell(h).lower() for h in next(it, ())]
    return [{header[i]: _cell(v) for i, v in enumerate(row) if i < len(header) and header[i]} for row in it]


def _slug(text: str, used: set[str]) -> str:
    base = re.sub(r"[^a-z0-9_]+", "_", text.lower()).strip("_") or "field"
    if not base[0].isalpha():
        base = "f_" + base
    name, n = base[:40], 2
    while name in used:
        name = f"{base[:36]}_{n}"
        n += 1
    return name


def parse_xlsform(data: bytes) -> tuple[dict[str, Any], list[str]]:
    """XLSForm (.xlsx: листы survey и choices) → описание формы и список того, что не перенесено."""
    from openpyxl import load_workbook

    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001 — не xlsx или битый файл
        raise FormError(422, "Файл не похож на XLSForm (.xlsx)") from exc
    sheets = {ws.title.lower(): ws for ws in wb.worksheets}
    if "survey" not in sheets:
        raise FormError(422, "В файле нет листа survey")
    survey = _rows(sheets["survey"])
    choice_rows = _rows(sheets["choices"]) if "choices" in sheets else []
    header = list(survey[0].keys()) if survey else []
    label_col = _label_column(header) or "label"
    choice_label = _label_column(list(choice_rows[0].keys())) if choice_rows else None
    lists: dict[str, list[dict[str, str]]] = {}
    for r in choice_rows:
        if r.get("list_name") and r.get("name"):
            lists.setdefault(r["list_name"], []).append(
                {"value": r["name"][:100], "label": (r.get(choice_label or "label") or r["name"])[:200]})
    title = ""
    if "settings" in sheets:
        settings = _rows(sheets["settings"])
        title = (settings[0].get("form_title") or "") if settings else ""

    fields: list[dict[str, Any]] = []
    warnings: list[str] = []
    used: set[str] = set()
    for i, r in enumerate(survey, start=2):
        raw_type = r.get("type", "")
        if not raw_type:
            continue
        parts = raw_type.split()
        kind = parts[0].lower()
        label = r.get(label_col) or r.get("label") or r.get("name") or ""
        if kind in ("begin_group", "end_group", "begin", "end", "begin_repeat", "end_repeat"):
            if kind.startswith("begin"):
                warnings.append(f"строка {i}: группа «{label or r.get('name', '')}» — поля перенесены без группы")
            continue
        if kind in ("start", "end", "today", "deviceid", "username", "calculate", "hidden"):
            continue  # служебные поля XLSForm: время, устройство — ТГИД пишет их сам
        ftype = _XLS_TYPES.get(kind)
        if kind in ("select_one", "select_multiple"):
            ftype = kind
        if ftype is None:
            warnings.append(f"строка {i}: тип «{raw_type}» не поддерживается — поле «{label}» пропущено")
            continue
        name = r.get("name", "")
        name = name if _NAME.match(name or "") and name not in used else _slug(name or label, used)
        used.add(name)
        field: dict[str, Any] = {"name": name, "type": ftype, "label": (label or name)[:300],
                                 "hint": r.get("hint", "")[:500],
                                 "required": r.get("required", "").lower() in ("yes", "true", "true()", "1", "да")}
        if ftype in ("select_one", "select_multiple"):
            list_name = parts[1] if len(parts) > 1 else ""
            choices = lists.get(list_name)
            if not choices:
                warnings.append(f"строка {i}: нет вариантов «{list_name}» на листе choices — поле «{label}» пропущено")
                continue
            field["choices"] = choices[:MAX_CHOICES]
        relevant = r.get("relevant", "")
        if relevant:
            m, s = _REL_EQ.match(relevant), _REL_SEL.match(relevant)
            if m:
                field["relevant"] = {"field": m.group(1), "op": m.group(2), "value": m.group(3)}
            elif s:
                field["relevant"] = {"field": s.group(1), "op": "selected", "value": s.group(2)}
            else:
                warnings.append(f"строка {i}: условие «{relevant}» сложнее «${{поле}} = 'значение'» — поле «{label}» показывается всегда")
        constraint = r.get("constraint", "")
        if constraint and ftype in ("integer", "decimal"):
            ok = True
            for part in re.split(r"\s+and\s+", constraint):
                c = _CON_PART.match(part.strip())
                if not c:
                    ok = False
                    break
                key = "min" if c.group(1).startswith(">") else "max"
                field[key] = float(c.group(2))
            if not ok:
                field.pop("min", None)
                field.pop("max", None)
                warnings.append(f"строка {i}: ограничение «{constraint}» не перенесено (только «. >= N and . <= M»)")
        elif constraint:
            warnings.append(f"строка {i}: ограничение «{constraint}» не перенесено")
        fields.append(field)
    # условие на поле, которое пропущено или ниже, — убрать
    names: set[str] = set()
    for f in fields:
        rel = f.get("relevant")
        if rel and rel["field"] not in names:
            warnings.append(f"условие поля «{f['label']}» ссылается на «{rel['field']}», которого нет выше — убрано")
            f.pop("relevant")
        names.add(f["name"])
    if len(fields) > MAX_FIELDS:
        warnings.append(f"полей больше {MAX_FIELDS} — перенесены первые {MAX_FIELDS}")
        fields = fields[:MAX_FIELDS]
    if not fields:
        raise FormError(422, "В форме не найдено ни одного поддерживаемого поля")
    return {"name": title, "schema": {"fields": fields}}, warnings


# --- хранение ----------------------------------------------------------------------------------

def _iso(v: Any) -> Any:
    return v.isoformat() if isinstance(v, (datetime, date)) else v


def _json(v: Any) -> Any:
    return json.loads(v) if isinstance(v, str) else v


_FORM_COLUMNS = ("f.id, f.name, f.description, f.scope, f.current_version, f.active, f.created_by, f.created_at, "
                 "f.updated_by, f.updated_at")


def _form(row, schema: Any = None) -> dict[str, Any]:
    out = {k: _iso(row[k]) for k in row.keys() if k != "schema"}
    if schema is not None or "schema" in row.keys():
        out["schema"] = _json(schema if schema is not None else row["schema"])
    return out


async def list_forms(conn, *, active_only: bool = False) -> list[dict[str, Any]]:
    """Формы с полями текущей версии и числом заполнений."""
    where = "WHERE f.active" if active_only else ""
    rows = await conn.fetch(
        f"""SELECT {_FORM_COLUMNS}, v.schema,
                   (SELECT count(*) FROM tgid_form_submission s WHERE s.form_id = f.id) AS submissions
              FROM tgid_form f JOIN tgid_form_version v ON v.form_id = f.id AND v.version = f.current_version
              {where} ORDER BY f.name, f.id""")
    return [_form(r) for r in rows]


async def get_form(conn, form_id: int, version: Optional[int] = None) -> dict[str, Any]:
    row = await conn.fetchrow(
        f"""SELECT {_FORM_COLUMNS}, v.version, v.schema FROM tgid_form f
              JOIN tgid_form_version v ON v.form_id = f.id AND v.version = coalesce($2, f.current_version)
             WHERE f.id = $1""", form_id, version)
    if row is None:
        raise FormError(404, f"Форма {form_id}" + (f" версии {version}" if version else "") + " не найдена")
    return _form(row)


async def create_form(conn, body: FormIn, user: str) -> dict[str, Any]:
    async with conn.transaction():
        form_id = await conn.fetchval(
            "INSERT INTO tgid_form (name, description, scope, created_by, updated_by) VALUES ($1, $2, $3, $4, $4)"
            " RETURNING id", body.name, body.description.strip(), body.scope, user)
        await conn.execute(
            "INSERT INTO tgid_form_version (form_id, version, schema, created_by) VALUES ($1, 1, $2::jsonb, $3)",
            form_id, body.schema_.model_dump_json(exclude_none=True), user)
    return await get_form(conn, form_id)


async def update_form(conn, form_id: int, body: FormIn, user: str) -> dict[str, Any]:
    """Название, описание, где заполняется; поля изменились — новая версия."""
    current = await get_form(conn, form_id)
    schema = json.loads(body.schema_.model_dump_json(exclude_none=True))
    async with conn.transaction():
        version = current["current_version"]
        if schema != current["schema"]:
            version += 1
            await conn.execute(
                "INSERT INTO tgid_form_version (form_id, version, schema, created_by) VALUES ($1, $2, $3::jsonb, $4)",
                form_id, version, json.dumps(schema, ensure_ascii=False), user)
        await conn.execute(
            """UPDATE tgid_form SET name = $2, description = $3, scope = $4, current_version = $5, updated_by = $6,
                      updated_at = now() WHERE id = $1""",
            form_id, body.name, body.description.strip(), body.scope, version, user)
    return await get_form(conn, form_id)


async def set_active(conn, form_id: int, active: bool, user: str) -> dict[str, Any]:
    await get_form(conn, form_id)
    await conn.execute("UPDATE tgid_form SET active = $2, updated_by = $3, updated_at = now() WHERE id = $1",
                       form_id, active, user)
    return await get_form(conn, form_id)


async def check_form(conn, form_id: Optional[int]) -> None:
    """Форма для наряда: есть и действует."""
    if form_id and not await conn.fetchval("SELECT active FROM tgid_form WHERE id = $1", form_id):
        raise FormError(422, f"Формы {form_id} нет или она выключена")


# --- заполнения --------------------------------------------------------------------------------

class SubmissionIn(BaseModel):
    form_id: int = Field(ge=1)
    version: int = Field(ge=1)
    work_order_id: int = Field(ge=1)
    node_id: Optional[int] = Field(None, ge=1, description="Узел маршрута обхода — у формы «на каждом узле»")
    answers: dict[str, Any]
    client_id: UUID
    lon: Optional[float] = Field(None, ge=-180, le=180)
    lat: Optional[float] = Field(None, ge=-90, le=90)
    filled_at: Optional[datetime] = None


_SUBMISSION_COLUMNS = ("s.id, s.form_id, s.version, s.work_order_id, s.node_id, s.answers, s.lon, s.lat, "
                       "s.client_id, s.filled_at, s.submitted_by, s.created_at")


def _submission(row) -> dict[str, Any]:
    out = {k: _iso(row[k]) for k in row.keys()}
    out["answers"] = _json(row["answers"])
    out["client_id"] = str(row["client_id"]) if row["client_id"] else None
    if "schema" in row.keys():
        out["schema"] = _json(row["schema"])
    return out


async def submit(conn, body: SubmissionIn, user: str) -> tuple[dict[str, Any], bool]:
    """Записать заполнение; (заполнение, новое ли). Повтор с тем же client_id — прежнее, без записи."""
    found = await conn.fetchrow(
        f"SELECT {_SUBMISSION_COLUMNS} FROM tgid_form_submission s WHERE s.client_id = $1", body.client_id)
    if found is not None:
        return _submission(found), False
    form = await get_form(conn, body.form_id, body.version)
    order = await conn.fetchrow("SELECT id, form_id, inspection_id FROM tgid_work_order WHERE id = $1",
                                body.work_order_id)
    if order is None:
        raise FormError(404, f"Наряд {body.work_order_id} не найден")
    if form["scope"] == "point" and body.node_id is None:
        raise FormError(422, "Форма заполняется на узле обхода: нужен node_id")
    schema = FormSchema.model_validate(form["schema"])
    clean, errors = validate_answers(schema, body.answers)
    if errors:
        labels = {f.name: f.label for f in schema.fields}
        text = "; ".join(f"{labels.get(k, k)}: {v}" for k, v in errors.items())
        raise FormError(422, f"Форма заполнена с ошибками — {text}")
    row = await conn.fetchrow(
        f"""INSERT INTO tgid_form_submission (form_id, version, work_order_id, node_id, answers, lon, lat, client_id,
                                              filled_at, submitted_by)
            VALUES ($1, $2, $3, $4, $5::jsonb, $6, $7, $8, $9, $10)
            ON CONFLICT (client_id) DO NOTHING
            RETURNING id""",
        body.form_id, body.version, body.work_order_id, body.node_id, json.dumps(clean, ensure_ascii=False),
        body.lon, body.lat, body.client_id, body.filled_at, user)
    if row is None:  # параллельный повтор успел раньше
        found = await conn.fetchrow(
            f"SELECT {_SUBMISSION_COLUMNS} FROM tgid_form_submission s WHERE s.client_id = $1", body.client_id)
        return _submission(found), False
    saved = await conn.fetchrow(f"SELECT {_SUBMISSION_COLUMNS} FROM tgid_form_submission s WHERE s.id = $1", row["id"])
    return _submission(saved), True


async def order_submissions(conn, order_id: int) -> list[dict[str, Any]]:
    """Заполнения наряда с полями своей версии и фото (по client_id фото наряда)."""
    rows = await conn.fetch(
        f"""SELECT {_SUBMISSION_COLUMNS}, v.schema, f.name AS form_name,
                   coalesce(nullif(n.nodename, ''), n.externalnodename) AS node_name
              FROM tgid_form_submission s
              JOIN tgid_form f ON f.id = s.form_id
              JOIN tgid_form_version v ON v.form_id = s.form_id AND v.version = s.version
              LEFT JOIN nodes n ON n.id = s.node_id
             WHERE s.work_order_id = $1 ORDER BY s.created_at, s.id""", order_id)
    out = [_submission(r) | {"form_name": r["form_name"], "node_name": r["node_name"]} for r in rows]
    photo_ids = [pid for s in out for f in s["schema"]["fields"] if f["type"] == "image"
                 for pid in s["answers"].get(f["name"], [])]
    photos = {}
    if photo_ids:
        photos = {str(r["client_id"]): r["id"] for r in await conn.fetch(
            "SELECT id, client_id FROM tgid_work_order_photo WHERE client_id = ANY($1::uuid[])", photo_ids)}
    for s in out:
        # фото поля: id фото наряда (ещё не дошедшее с телефона — null)
        s["photos"] = {pid: photos.get(pid) for f in s["schema"]["fields"] if f["type"] == "image"
                       for pid in s["answers"].get(f["name"], [])}
    return out


def _local(dt: Optional[datetime]) -> Optional[datetime]:
    """Время для Excel — местное время сервера города, без пояса (Excel поясов не знает)."""
    return dt.astimezone().replace(tzinfo=None) if dt else None


def _display(field: dict[str, Any], value: Any) -> Any:
    if value is None:
        return ""
    labels = {c["value"]: c["label"] for c in field.get("choices") or ()}
    if field["type"] == "select_one":
        return labels.get(value, value)
    if field["type"] == "select_multiple":
        return ", ".join(labels.get(v, v) for v in value)
    if field["type"] == "acknowledge":
        return "да" if value else ""
    if field["type"] == "image":
        return f"фото: {len(value)}"
    return value


async def export_xlsx(conn, form_id: int) -> tuple[str, bytes]:
    """Заполнения формы в Excel: строка — заполнение, колонки — поля всех версий по порядку появления."""
    from openpyxl import Workbook
    from openpyxl.utils import get_column_letter

    form = await get_form(conn, form_id)
    rows = await conn.fetch(
        """SELECT s.*, v.schema, o.number AS order_number,
                  coalesce(nullif(n.nodename, ''), n.externalnodename) AS node_name
             FROM tgid_form_submission s
             JOIN tgid_form_version v ON v.form_id = s.form_id AND v.version = s.version
             LEFT JOIN tgid_work_order o ON o.id = s.work_order_id
             LEFT JOIN nodes n ON n.id = s.node_id
            WHERE s.form_id = $1 ORDER BY s.created_at, s.id""", form_id)
    columns: dict[str, dict[str, Any]] = {}
    for f in form["schema"]["fields"]:
        if f["type"] != "note":
            columns[f["name"]] = f
    for r in rows:
        for f in _json(r["schema"])["fields"]:
            if f["type"] != "note":
                columns.setdefault(f["name"], f)
    wb = Workbook()
    ws = wb.active
    ws.title = "Заполнения"
    head = ["Наряд", "Узел", "Кто", "Заполнено", "Принято", "Долгота", "Широта", "Версия формы"]
    ws.append(head + [c["label"] for c in columns.values()])
    for r in rows:
        answers = _json(r["answers"])
        fields = {f["name"]: f for f in _json(r["schema"])["fields"]}
        base = [r["order_number"] or "", r["node_name"] or (r["node_id"] or ""), r["submitted_by"],
                _local(r["filled_at"]), _local(r["created_at"]),
                r["lon"], r["lat"], r["version"]]
        ws.append(base + [_display(fields.get(name, col), answers.get(name)) for name, col in columns.items()])
    for i in range(1, len(head) + len(columns) + 1):
        ws.column_dimensions[get_column_letter(i)].width = 18
    ws.freeze_panes = "A2"
    buf = io.BytesIO()
    wb.save(buf)
    safe = re.sub(r"[^\w\-]+", "_", form["name"])[:60] or f"form_{form_id}"
    return f"{safe}.xlsx", buf.getvalue()
