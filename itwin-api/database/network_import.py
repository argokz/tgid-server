"""Импорт сети из файлов (этап 10): SHP, Excel/CSV и координаты узлов.

Эталон десктопа: в gid8 импорта SHP в сеть нет — `python/convertor/shape_pg/shp_ms.py`
грузит SHP (fiona + pyproj) в отдельную «сырую» таблицу MS SQL с перепроецированием,
а `dialog/geof.cpp` (выбор геобаз *.shp/*.mdb) закрыт `#if 0`; импорта координат узлов
из Excel нет ни в gid8, ни в gid6 (`OnCoord` — только параметры проекции). Поэтому
импорт спроектирован по правкам топологии этапа 8 и переиспользует их:

  * nodes  — точки SHP или строки таблицы с X/Y → новые узлы фрагмента (как create_node:
             фрагмент/код/признак от ближайшего узла фрагмента), поля: код
             (externalnodename), наименование, рег. номер, отметка земли;
  * lines  — линии SHP → участки: концы привязываются к узлам фрагмента в допуске
             (иначе создаются новые узлы; новые узлы этого же импорта тоже находятся),
             паспорт трубы — от участка-образца (как create_line), Ду и длина из полей;
  * coords — таблица «узел (id или код) → X, Y»: перенос узлов как move_node (B4):
             концы участков следуют за узлом, длина паспорта пересчитывается по геометрии.

Всё — одной транзакцией через `database.topology._run`: dry-run выполняет импорт и
откатывает (точный отчёт «что будет создано/обновлено», ошибки по строкам), применение
пишет audit_log и журнал отмены топологии (одна операция IMPORT_* — отменяется целиком).
Строка с ошибкой в своей точке сохранения не ломает остальные; при ошибках применение
отказывает целиком, если не задан skip_errors.

Координаты: WGS84 (долгота/широта), местная система (SRID 9998, метры) или координаты
десктопа (nodes.x/y: сантиметры местной системы, Y с обратным знаком). SHP с .prj
перепроецируется в WGS84 сам; без .prj система указывается явно.
"""

from __future__ import annotations

import csv
import io
import math
import os
import re
import tempfile
import zipfile
from dataclasses import dataclass, field
from datetime import date, datetime
from typing import Any, Optional

from database.outage_simulation import invalidate_outage_cache
from database.topology import (
    _capture_incident_lines,
    _create_line_passport,
    _new_node_reference,
    _run,
)

MAX_UPLOAD_BYTES = 20 * 1024 * 1024
MAX_ROWS = 5000
MAX_ACTIONS_IN_REPORT = 500
MAX_SNAP_TOLERANCE_M = 50.0
LOCAL_SRID = 9998

MODES = ("nodes", "lines", "coords")
SOURCE_CRS = ("auto", "wgs84", "local", "desktop")

# Поля назначения по режимам: ключ → (подпись, обязательное)
TARGET_FIELDS: dict[str, dict[str, tuple[str, bool]]] = {
    "nodes": {
        "x": ("X / долгота (для таблицы)", True),
        "y": ("Y / широта (для таблицы)", True),
        "code": ("Код узла (externalnodename)", False),
        "name": ("Наименование", False),
        "regnum": ("Рег. номер", False),
        "elevation": ("Отметка земли, м", False),
    },
    "lines": {
        "regnum": ("Рег. номер участка", False),
        "diameter": ("Внутренний диаметр, мм", False),
        "length": ("Длина, м (иначе по геометрии)", False),
    },
    "coords": {
        "key": ("Узел: id или код", True),
        "x": ("X / долгота (для таблицы)", True),
        "y": ("Y / широта (для таблицы)", True),
    },
}

_SUGGEST: dict[str, tuple[str, ...]] = {
    "x": ("x", "lon", "lng", "long", "longitude", "долгота", "east", "easting", "х", "коорд_x", "coord_x"),
    "y": ("y", "lat", "latitude", "широта", "north", "northing", "у", "коорд_y", "coord_y"),
    "code": ("code", "kod", "код", "код узла", "externalnodename", "uzel", "узел", "номер узла", "node"),
    "name": ("name", "nodename", "наименование", "название", "naim"),
    "regnum": ("regnum", "registnum", "registnumber", "рег", "рег. номер", "рег.номер", "инв", "inventnumber"),
    "elevation": ("z", "elev", "elevation", "отметка", "geomarknodearea", "h", "высота"),
    "diameter": ("d", "dn", "du", "ду", "диаметр", "diameter", "diam", "d_mm"),
    "length": ("l", "len", "length", "длина", "l_m"),
    "key": ("id", "node_id", "nodeid", "sys", "code", "kod", "код", "код узла", "externalnodename", "узел"),
}


class ImportRowErrors(Exception):
    """Применение отклонено: есть ошибки по строкам (без skip_errors). report — отчёт dry-run."""

    def __init__(self, report: dict):
        super().__init__(f"Ошибок в строках: {len(report.get('errors') or [])}")
        self.report = report


class _RowError(Exception):
    pass


# ---------------------------------------------------------------------------
# Разбор файла
# ---------------------------------------------------------------------------

@dataclass
class ParsedSource:
    kind: str  # "shp" | "table"
    columns: list[str]
    rows: list[dict[str, Any]]
    # номер строки для отчёта: строка Excel/CSV (с заголовком) или номер объекта SHP с 1
    row_numbers: list[int]
    geometries: list[Any] = field(default_factory=list)  # shapely-геометрии (SHP)
    geometry_type: Optional[str] = None
    crs: Optional[str] = None  # описание исходной системы SHP (None — .prj нет)
    geometry_srid: Optional[int] = None  # 4326 — SHP уже перепроецирован; None — задать вручную
    sheet: Optional[str] = None
    sheets: list[str] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)


def _json_safe(value: Any) -> Any:
    if value is None:
        return None
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if hasattr(value, "item") and not isinstance(value, (str, bytes)):
        try:
            value = value.item()  # numpy-скаляры
        except (ValueError, AttributeError):
            pass
    if isinstance(value, float) and not math.isfinite(value):
        return None
    if isinstance(value, (int, float, str, bool)):
        return value.strip() if isinstance(value, str) else value
    return str(value)


def _unique_columns(names: list[Any]) -> list[str]:
    out: list[str] = []
    seen: dict[str, int] = {}
    for i, raw in enumerate(names):
        name = str(raw).strip() if raw not in (None, "") else f"column{i + 1}"
        if name in seen:
            seen[name] += 1
            name = f"{name}_{seen[name]}"
        else:
            seen[name] = 1
        out.append(name)
    return out


def _parse_xlsx(content: bytes, sheet: Optional[str]) -> ParsedSource:
    from openpyxl import load_workbook

    try:
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as e:  # noqa: BLE001 — битый/не тот файл
        raise ValueError(f"Не удалось прочитать Excel: {e}") from e
    try:
        sheets = list(wb.sheetnames)
        if sheet and sheet not in sheets:
            raise ValueError(f"Лист «{sheet}» не найден; есть: {', '.join(sheets)}")
        ws = wb[sheet] if sheet else wb.worksheets[0]
        header: Optional[list[str]] = None
        header_row = 0
        rows: list[dict] = []
        numbers: list[int] = []
        for idx, values in enumerate(ws.iter_rows(values_only=True), start=1):
            if header is None:
                if values and any(v not in (None, "") for v in values):
                    header = _unique_columns(list(values))
                    header_row = idx
                continue
            if not values or all(v in (None, "") for v in values):
                continue
            if len(rows) >= MAX_ROWS:
                raise ValueError(f"В файле больше {MAX_ROWS} строк — разбейте импорт на части")
            rows.append({header[i]: _json_safe(v) for i, v in enumerate(values) if i < len(header)})
            numbers.append(idx)
        if header is None:
            raise ValueError("Лист пуст: нет строки заголовков")
        warnings = [] if header_row == 1 else [f"Заголовки взяты из строки {header_row}"]
        return ParsedSource("table", header, rows, numbers, sheet=ws.title, sheets=sheets, warnings=warnings)
    finally:
        wb.close()


def _parse_csv(content: bytes, encoding: Optional[str]) -> ParsedSource:
    text = None
    for enc in ([encoding] if encoding else ["utf-8-sig", "cp1251"]):
        try:
            text = content.decode(enc)
            break
        except (UnicodeDecodeError, LookupError):
            continue
    if text is None:
        raise ValueError("Не удалось определить кодировку CSV (ожидается UTF-8 или Windows-1251)")
    sample = text[:4096]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=";,\t")
        delimiter = dialect.delimiter
    except csv.Error:
        delimiter = ";"
    reader = csv.reader(io.StringIO(text), delimiter=delimiter)
    header: Optional[list[str]] = None
    rows: list[dict] = []
    numbers: list[int] = []
    for idx, values in enumerate(reader, start=1):
        if header is None:
            if any(v.strip() for v in values):
                header = _unique_columns(values)
            continue
        if not any(v.strip() for v in values):
            continue
        if len(rows) >= MAX_ROWS:
            raise ValueError(f"В файле больше {MAX_ROWS} строк — разбейте импорт на части")
        rows.append({header[i]: (v.strip() or None) for i, v in enumerate(values) if i < len(header)})
        numbers.append(idx)
    if header is None:
        raise ValueError("CSV пуст: нет строки заголовков")
    return ParsedSource("table", header, rows, numbers)


def _parse_shp(files: list[tuple[str, bytes]], encoding: Optional[str]) -> ParsedSource:
    import geopandas as gpd

    with tempfile.TemporaryDirectory(prefix="tgid_import_") as tmp:
        for name, content in files:
            base = os.path.basename(name.replace("\\", "/"))
            if base.lower().endswith(".zip"):
                try:
                    with zipfile.ZipFile(io.BytesIO(content)) as zf:
                        for member in zf.infolist():
                            mbase = os.path.basename(member.filename.replace("\\", "/"))
                            if member.is_dir() or not mbase or mbase.startswith("."):
                                continue
                            if member.file_size > MAX_UPLOAD_BYTES:
                                raise ValueError(f"Файл {mbase} в архиве больше 20 МБ")
                            with open(os.path.join(tmp, mbase), "wb") as f:
                                f.write(zf.read(member))
                except zipfile.BadZipFile as e:
                    raise ValueError("Архив повреждён или это не zip") from e
            else:
                with open(os.path.join(tmp, base), "wb") as f:
                    f.write(content)
        shps = sorted(n for n in os.listdir(tmp) if n.lower().endswith(".shp"))
        if not shps:
            raise ValueError("В загрузке нет .shp (нужен zip или файлы .shp + .shx + .dbf [+ .prj, .cpg])")
        warnings = []
        if len(shps) > 1:
            warnings.append(f"В архиве несколько слоёв, взят первый: {shps[0]} (есть: {', '.join(shps)})")
        stem = shps[0][:-4]
        present = {n.lower() for n in os.listdir(tmp)}
        for ext in (".shx", ".dbf"):
            if f"{stem.lower()}{ext}" not in present:
                raise ValueError(f"Нет файла {stem}{ext} — shapefile неполный")
        kwargs: dict[str, Any] = {}
        if encoding:
            kwargs["encoding"] = encoding
        elif f"{stem.lower()}.cpg" not in present:
            # Без .cpg dbf десктопных выгрузок — в Windows-1251
            kwargs["encoding"] = "cp1251"
            warnings.append("Нет .cpg — атрибуты прочитаны в кодировке Windows-1251")
        try:
            gdf = gpd.read_file(os.path.join(tmp, shps[0]), **kwargs)
        except Exception as e:  # noqa: BLE001
            raise ValueError(f"Не удалось прочитать shapefile: {e}") from e
    if len(gdf) > MAX_ROWS:
        raise ValueError(f"В слое больше {MAX_ROWS} объектов — разбейте импорт на части")
    crs_text = None
    srid = None
    if gdf.crs is not None:
        crs_text = gdf.crs.to_string()
        if len(gdf):
            gdf = gdf.to_crs(4326)
        srid = 4326
    else:
        warnings.append("Нет .prj — укажите систему координат вручную")
    columns = [str(c) for c in gdf.columns if c != gdf.geometry.name]
    rows = [{c: _json_safe(v) for c, v in zip(columns, rec)} for rec in gdf[columns].itertuples(index=False, name=None)]
    geoms = list(gdf.geometry)
    types = sorted({g.geom_type for g in geoms if g is not None and not g.is_empty})
    return ParsedSource(
        "shp", columns, rows, list(range(1, len(rows) + 1)),
        geometries=geoms, geometry_type="/".join(types) or None, crs=crs_text, geometry_srid=srid,
        warnings=warnings,
    )


def parse_upload(
    files: list[tuple[str, bytes]], *, encoding: Optional[str] = None, sheet: Optional[str] = None,
) -> ParsedSource:
    """Файлы загрузки → таблица атрибутов (+ геометрия для SHP)."""
    if not files:
        raise ValueError("Файл не передан")
    total = sum(len(c) for _, c in files)
    if total > MAX_UPLOAD_BYTES:
        raise ValueError("Загрузка больше 20 МБ")
    exts = {os.path.splitext(n.lower())[1] for n, _ in files}
    if exts & {".zip", ".shp"}:
        return _parse_shp(files, encoding)
    if len(files) != 1:
        raise ValueError("Excel/CSV загружается одним файлом")
    name, content = files[0]
    ext = os.path.splitext(name.lower())[1]
    if ext in (".xlsx", ".xlsm"):
        return _parse_xlsx(content, sheet)
    if ext in (".csv", ".txt"):
        return _parse_csv(content, encoding)
    if ext == ".xls":
        raise ValueError("Формат .xls (Excel 97–2003) не поддерживается — сохраните файл как .xlsx или CSV")
    raise ValueError(f"Неподдерживаемый файл {name}: нужен .zip/.shp, .xlsx или .csv")


def _norm(name: str) -> str:
    return re.sub(r"[\s_\-.]+", " ", name.strip().lower())


def suggest_mapping(mode: str, columns: list[str], has_geometry: bool) -> dict[str, str]:
    """Сопоставление «поле назначения → колонка файла» по типичным именам."""
    result: dict[str, str] = {}
    used: set[str] = set()
    normalized = {c: _norm(c) for c in columns}
    for target in TARGET_FIELDS.get(mode, {}):
        if has_geometry and target in ("x", "y"):
            continue
        variants = {_norm(v) for v in _SUGGEST.get(target, ())}
        for col in columns:
            if col not in used and normalized[col] in variants:
                result[target] = col
                used.add(col)
                break
    return result


def describe_source(src: ParsedSource, mode: str) -> dict:
    has_geometry = src.kind == "shp"
    return {
        "kind": src.kind,
        "columns": src.columns,
        "row_count": len(src.rows),
        "sample": src.rows[:20],
        "geometry_type": src.geometry_type,
        "crs": src.crs,
        "crs_known": src.geometry_srid is not None,
        "sheet": src.sheet,
        "sheets": src.sheets,
        "warnings": src.warnings,
        "targets": [
            {"key": k, "label": label, "required": req and not (has_geometry and k in ("x", "y"))}
            for k, (label, req) in TARGET_FIELDS.get(mode, {}).items()
            if not (has_geometry and k in ("x", "y"))
        ],
        "suggested_mapping": suggest_mapping(mode, src.columns, has_geometry),
    }


# ---------------------------------------------------------------------------
# Подготовка строк
# ---------------------------------------------------------------------------

def to_float(value: Any) -> Optional[float]:
    """Число из ячейки: 12.5, «12,5», «1 234,5»; пусто → None; не число → ValueError."""
    if value is None:
        return None
    if isinstance(value, bool):
        raise ValueError("не число")
    if isinstance(value, (int, float)):
        f = float(value)
    else:
        s = str(value).strip().replace(" ", "").replace(" ", "")
        if not s:
            return None
        if s.count(",") == 1 and "." not in s:
            s = s.replace(",", ".")
        try:
            f = float(s)
        except ValueError as e:
            raise ValueError(f"«{value}» — не число") from e
    if not math.isfinite(f):
        raise ValueError("не число")
    return f


def _text(value: Any) -> Optional[str]:
    if value is None:
        return None
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    s = str(value).strip()
    return s or None


def point_to_source_xy(x: float, y: float, source_crs: str) -> tuple[float, float, int]:
    """X/Y таблицы → (x, y, srid) для ST_SetSRID: WGS84, местная (м) или десктоп (см, −Y)."""
    if source_crs == "wgs84":
        if not (-180 <= x <= 180 and -90 <= y <= 90):
            raise ValueError("долгота/широта вне диапазона WGS84")
        return x, y, 4326
    if source_crs == "local":
        return x, y, LOCAL_SRID
    if source_crs == "desktop":
        return x / 100.0, -y / 100.0, LOCAL_SRID
    raise ValueError("Укажите систему координат (WGS84, местная или десктоп)")


@dataclass
class _Item:
    row: int
    attrs: dict[str, Any]
    x: Optional[float] = None
    y: Optional[float] = None
    srid: Optional[int] = None
    wkt: Optional[str] = None
    error: Optional[str] = None


def prepare_items(src: ParsedSource, params: dict) -> list[_Item]:
    """Строки файла → элементы импорта: поля по сопоставлению, координаты, ошибки строки."""
    mode = params["mode"]
    mapping: dict[str, str] = params.get("mapping") or {}
    unknown = [c for c in mapping.values() if c and c not in src.columns]
    if unknown:
        raise ValueError(f"Нет колонок в файле: {', '.join(unknown)}")
    fields = TARGET_FIELDS[mode]
    use_geometry = src.kind == "shp"
    for key, (label, required) in fields.items():
        if use_geometry and key in ("x", "y"):
            continue
        if required and not mapping.get(key):
            raise ValueError(f"Не сопоставлено обязательное поле «{label}»")
    if mode == "lines" and not use_geometry:
        raise ValueError("Участки импортируются только из SHP с линиями")
    source_crs = params.get("source_crs") or "auto"
    if use_geometry and src.geometry_srid is None and source_crs in ("auto", "desktop"):
        raise ValueError("У SHP нет .prj: укажите систему координат (WGS84 или местная)")
    if not use_geometry and source_crs == "auto":
        raise ValueError("Укажите систему координат колонок X/Y")

    items: list[_Item] = []
    for idx, (row, number) in enumerate(zip(src.rows, src.row_numbers)):
        item = _Item(row=number, attrs={k: row.get(col) for k, col in mapping.items() if col})
        try:
            if use_geometry:
                geom = src.geometries[idx]
                if geom is None or geom.is_empty:
                    raise ValueError("нет геометрии")
                if src.geometry_srid is not None:
                    item.srid = src.geometry_srid
                else:
                    item.srid = 4326 if source_crs == "wgs84" else LOCAL_SRID
                if mode == "lines":
                    if geom.geom_type == "MultiLineString":
                        if len(geom.geoms) != 1:
                            raise ValueError(f"мультилиния из {len(geom.geoms)} частей — разбейте на участки")
                        geom = geom.geoms[0]
                    if geom.geom_type != "LineString":
                        raise ValueError(f"ожидалась линия, в файле {geom.geom_type}")
                    item.wkt = geom.wkt
                else:
                    if geom.geom_type == "MultiPoint" and len(geom.geoms) == 1:
                        geom = geom.geoms[0]
                    if geom.geom_type != "Point":
                        raise ValueError(f"ожидалась точка, в файле {geom.geom_type}")
                    item.x, item.y = float(geom.x), float(geom.y)
                    if item.srid == 4326:
                        point_to_source_xy(item.x, item.y, "wgs84")
            else:
                x = to_float(item.attrs.get("x"))
                y = to_float(item.attrs.get("y"))
                if x is None or y is None:
                    raise ValueError("нет координат X/Y")
                item.x, item.y, item.srid = point_to_source_xy(x, y, source_crs)
            for key in ("elevation", "diameter", "length"):
                if key in item.attrs:
                    item.attrs[key] = to_float(item.attrs[key])
            if item.attrs.get("diameter") is not None and not (0 < item.attrs["diameter"] <= 2000):
                raise ValueError("диаметр вне 0…2000 мм")
            if item.attrs.get("length") is not None and not (0 < item.attrs["length"] <= 100000):
                raise ValueError("длина вне 0…100 000 м")
            for key in ("code", "name", "regnum", "key"):
                if key in item.attrs:
                    item.attrs[key] = _text(item.attrs[key])
            if mode == "coords" and not item.attrs.get("key"):
                raise ValueError("пустой id/код узла")
        except ValueError as e:
            item.error = str(e)
        items.append(item)
    return items


# ---------------------------------------------------------------------------
# Импорт в транзакции
# ---------------------------------------------------------------------------

async def _to_wgs84(conn, items: list[_Item]) -> None:
    """Точки местной системы → WGS84 одним запросом (x, y, srid заменяются на lon/lat 4326)."""
    local = [it for it in items if it.error is None and it.x is not None and it.srid == LOCAL_SRID]
    if not local:
        return
    rows = await conn.fetch(
        """
        SELECT t.i, ST_X(g) AS lng, ST_Y(g) AS lat
        FROM unnest($1::float8[], $2::float8[]) WITH ORDINALITY AS t(x, y, i),
             LATERAL (SELECT ST_Transform(ST_SetSRID(ST_MakePoint(t.x, t.y), $3), 4326) AS g) p
        ORDER BY t.i
        """,
        [it.x for it in local], [it.y for it in local], LOCAL_SRID,
    )
    for it, r in zip(local, rows):
        it.x, it.y, it.srid = float(r["lng"]), float(r["lat"]), 4326
        if not (-180 <= it.x <= 180 and -90 <= it.y <= 90):
            it.error = "координаты вне области проекции"


async def _insert_node(conn, lng: float, lat: float, fileid: int, attrs: dict) -> int:
    ref = await _new_node_reference(conn, lng, lat, fileid, None, None)
    if ref["externalcodeid"] is None:
        raise _RowError(f"у фрагмента {fileid} нет кода (externalcodes) — узел не попадёт в расчёт")
    return await conn.fetchval(
        """
        INSERT INTO nodes (shape, x, y, removed, archivechangedate, nodetypeid,
                           fileid, externalcodeid, externalsignid, internalnodeid,
                           externalnodename, nodename, registnumber, geomarknodearea)
        VALUES (
          ST_Transform(ST_SetSRID(ST_MakePoint($1, $2), 4326), 9998),
          ST_X(ST_Transform(ST_SetSRID(ST_MakePoint($1, $2), 4326), 9998)) * 100.0,
          -ST_Y(ST_Transform(ST_SetSRID(ST_MakePoint($1, $2), 4326), 9998)) * 100.0,
          0, $3, 1,
          $4, $5, $6, NULL,
          $7, $8, $9, $10
        ) RETURNING id
        """,
        lng, lat, datetime.now(),
        ref["fileid"], ref["externalcodeid"], ref["externalsignid"],
        attrs.get("code"), attrs.get("name"), attrs.get("regnum"), attrs.get("elevation"),
    )


async def _fragment_exists(conn, fileid: Optional[int]) -> None:
    if fileid is None:
        raise ValueError("Укажите фрагмент (fileid)")
    if not await conn.fetchval("SELECT EXISTS (SELECT 1 FROM fragments WHERE id = $1)", fileid):
        raise ValueError(f"Фрагмент {fileid} не найден")


async def _import_nodes(conn, op, items: list[_Item], params: dict, report: dict) -> None:
    fileid = params.get("fileid")
    await _fragment_exists(conn, fileid)
    await _to_wgs84(conn, items)
    codes = [it.attrs.get("code") for it in items if it.error is None and it.attrs.get("code")]
    existing = {
        r["externalnodename"]: r["id"]
        for r in await conn.fetch(
            """
            SELECT DISTINCT ON (externalnodename) externalnodename, id FROM nodes
            WHERE fileid = $1 AND externalnodename = ANY($2::text[])
              AND COALESCE(removed, 0) = 0 AND internalnodeid IS NULL
            ORDER BY externalnodename, id
            """,
            fileid, codes,
        )
    } if codes else {}
    seen: dict[str, int] = {}
    for it in items:
        if it.error:
            continue
        code = it.attrs.get("code")
        if code and code in existing:
            it.error = f"узел с кодом «{code}» уже есть во фрагменте (id {existing[code]})"
            continue
        if code and code in seen:
            it.error = f"код «{code}» повторяется (строка {seen[code]})"
            continue
        try:
            async with conn.transaction():
                node_id = await _insert_node(conn, it.x, it.y, fileid, it.attrs)
        except (_RowError, ValueError) as e:
            it.error = str(e)
            continue
        if code:
            seen[code] = it.row
        op.journal.created("nodes", node_id)
        await op.audit("INSERT", "nodes", node_id, {"import": True, "row": it.row, "lng": it.x, "lat": it.y, **it.attrs})
        report["created_nodes"] += 1
        report["actions"].append({"row": it.row, "action": "create_node", "id": node_id, "code": code})


async def _snap_or_create_node(conn, op, lng: float, lat: float, fileid: int, tolerance: float) -> tuple[int, bool]:
    near = await conn.fetchrow(
        """
        WITH p AS (SELECT ST_Transform(ST_SetSRID(ST_MakePoint($1, $2), 4326), 9998) AS g)
        SELECT n.id, n.externalcodeid FROM nodes n, p
        WHERE COALESCE(n.removed, 0) = 0 AND n.shape IS NOT NULL AND n.fileid = $3
          AND n.internalnodeid IS NULL AND ST_DWithin(n.shape, p.g, $4)
        ORDER BY n.shape <-> p.g, n.id
        LIMIT 1
        """,
        lng, lat, fileid, tolerance,
    )
    if near is not None:
        if near["externalcodeid"] is None:
            raise _RowError(f"узел {near['id']} у конца участка без кода (externalcodeid)")
        return near["id"], False
    return await _insert_node(conn, lng, lat, fileid, {}), True


async def _import_lines(conn, op, items: list[_Item], params: dict, report: dict) -> None:
    fileid = params.get("fileid")
    await _fragment_exists(conn, fileid)
    tolerance = float(params.get("snap_tolerance_m", 1.0))
    for it in items:
        if it.error:
            continue
        try:
            async with conn.transaction():
                g = await conn.fetchrow(
                    """
                    WITH src AS (SELECT ST_Force2D(ST_SetSRID(ST_GeomFromText($1), $2)) AS g),
                         w AS (SELECT ST_Transform(g, 4326) AS g FROM src)
                    SELECT ST_X(ST_StartPoint(w.g)) AS x1, ST_Y(ST_StartPoint(w.g)) AS y1,
                           ST_X(ST_EndPoint(w.g)) AS x2, ST_Y(ST_EndPoint(w.g)) AS y2,
                           ST_Length(ST_Transform(w.g, 9998)) AS len
                    FROM w
                    """,
                    it.wkt, it.srid,
                )
                if g is None or g["len"] is None or g["len"] < 0.01:
                    raise _RowError("участок нулевой длины")
                n1, new1 = await _snap_or_create_node(conn, op, g["x1"], g["y1"], fileid, tolerance)
                n2, new2 = await _snap_or_create_node(conn, op, g["x2"], g["y2"], fileid, tolerance)
                if n1 == n2:
                    raise _RowError(f"оба конца привязаны к узлу {n1} (участок короче допуска привязки?)")
                dup = await conn.fetchval(
                    """
                    SELECT id FROM linesobj
                    WHERE COALESCE(removed, 0) = 0
                      AND ((nodeid1 = $1 AND nodeid2 = $2) OR (nodeid1 = $2 AND nodeid2 = $1))
                    LIMIT 1
                    """,
                    n1, n2,
                )
                if dup:
                    raise _RowError(f"между узлами {n1} и {n2} уже есть участок {dup}")
                line_id = await conn.fetchval(
                    """
                    WITH g AS (SELECT ST_Transform(ST_Force2D(ST_SetSRID(ST_GeomFromText($3), $4)), 9998) AS g)
                    INSERT INTO linesobj (nodeid1, nodeid2, shape, removed, archivechangedate, fileid,
                                          internalnodeid, registnum)
                    SELECT $1, $2,
                           ST_SetPoint(
                             ST_SetPoint(g.g, 0, (SELECT shape FROM nodes WHERE id = $1)),
                             ST_NumPoints(g.g) - 1, (SELECT shape FROM nodes WHERE id = $2)),
                           0, $5, $6, NULL, $7
                    FROM g
                    RETURNING id
                    """,
                    n1, n2, it.wkt, it.srid, datetime.now(), fileid, it.attrs.get("regnum"),
                )
                passport = await _create_line_passport(conn, line_id, n1, n2)
                if it.attrs.get("diameter") is not None or it.attrs.get("length") is not None:
                    await conn.execute(
                        """
                        UPDATE heatpipesections
                        SET diameterinternal = COALESCE($2, diameterinternal),
                            pipesectlength = COALESCE($3, pipesectlength)
                        WHERE lineid = $1
                        """,
                        line_id, it.attrs.get("diameter"), it.attrs.get("length"),
                    )
                passport_ids = [r["id"] for r in await conn.fetch("SELECT id FROM heatpipesections WHERE lineid = $1", line_id)]
        except (_RowError, ValueError) as e:
            it.error = str(e)
            continue
        for node_id, is_new in ((n1, new1), (n2, new2)):
            if is_new:
                op.journal.created("nodes", node_id)
                await op.audit("INSERT", "nodes", node_id, {"import": True, "row": it.row, "line_end": True})
                report["created_nodes"] += 1
        op.journal.created("linesobj", line_id)
        for pid in passport_ids:
            op.journal.created("heatpipesections", pid)
        await op.audit(
            "INSERT", "linesobj", line_id,
            {"import": True, "row": it.row, "nodeid1": n1, "nodeid2": n2, "passport": passport, **it.attrs},
        )
        report["created_lines"] += 1
        report["actions"].append({
            "row": it.row, "action": "create_line", "id": line_id, "nodeid1": n1, "nodeid2": n2,
            "new_nodes": [n for n, new in ((n1, new1), (n2, new2)) if new],
            "length": round(float(g["len"]), 2), "passport": passport.get("source"),
        })


async def _resolve_coord_targets(conn, items: list[_Item], params: dict) -> None:
    match_by = params.get("match_by", "code")
    fileid = params.get("fileid")
    live = [it for it in items if it.error is None]
    if match_by == "id":
        ids: dict[int, _Item] = {}
        for it in live:
            try:
                nid = int(float(it.attrs["key"]))
                if not 0 < nid < 2**31:
                    raise ValueError
            except (ValueError, OverflowError):
                it.error = f"«{it.attrs['key']}» — не номер узла"
                continue
            ids[nid] = it
            it.attrs["node_id"] = nid
        found = {
            r["id"]: r for r in await conn.fetch(
                "SELECT id, fileid FROM nodes WHERE id = ANY($1::int[]) AND COALESCE(removed, 0) = 0",
                list(ids),
            )
        }
        for it in live:
            nid = it.attrs.get("node_id")
            if it.error or nid is None:
                continue
            if nid not in found:
                it.error = f"узел {nid} не найден или удалён"
            elif fileid is not None and found[nid]["fileid"] != fileid:
                it.error = f"узел {nid} из другого фрагмента ({found[nid]['fileid']})"
        return
    if fileid is None:
        raise ValueError("Для сопоставления по коду укажите фрагмент (коды повторяются между фрагментами)")
    await _fragment_exists(conn, fileid)
    rows = await conn.fetch(
        """
        SELECT externalnodename AS code, array_agg(id ORDER BY id) AS ids FROM nodes
        WHERE fileid = $1 AND externalnodename = ANY($2::text[])
          AND COALESCE(removed, 0) = 0 AND internalnodeid IS NULL
        GROUP BY externalnodename
        """,
        fileid, [it.attrs["key"] for it in live],
    )
    by_code = {r["code"]: list(r["ids"]) for r in rows}
    for it in live:
        ids_ = by_code.get(it.attrs["key"])
        if not ids_:
            it.error = f"узел с кодом «{it.attrs['key']}» не найден во фрагменте {fileid}"
        elif len(ids_) > 1:
            it.error = f"код «{it.attrs['key']}» неоднозначен: узлы {', '.join(map(str, ids_[:5]))}"
        else:
            it.attrs["node_id"] = ids_[0]


async def _import_coords(conn, op, items: list[_Item], params: dict, report: dict) -> None:
    await _resolve_coord_targets(conn, items, params)
    await _to_wgs84(conn, items)
    recalc = bool(params.get("recalc_lengths", True))
    build_missing = bool(params.get("build_missing_lines", False))
    seen: dict[int, int] = {}
    for it in items:
        if it.error:
            continue
        node_id = it.attrs["node_id"]
        if node_id in seen:
            it.error = f"узел {node_id} уже задан в строке {seen[node_id]}"
            continue
        seen[node_id] = it.row
        await op.journal.capture_ids(conn, "nodes", [node_id])
        await _capture_incident_lines(conn, op, [node_id])
        old = await conn.fetchrow(
            "SELECT ST_Distance(shape, ST_Transform(ST_SetSRID(ST_MakePoint($2, $3), 4326), 9998)) AS shift, "
            "shape IS NULL AS had_no_shape FROM nodes WHERE id = $1",
            node_id, it.x, it.y,
        )
        now = datetime.now()
        await conn.execute(
            """
            UPDATE nodes SET
              shape = ST_Transform(ST_SetSRID(ST_MakePoint($1, $2), 4326), 9998),
              x = ST_X(ST_Transform(ST_SetSRID(ST_MakePoint($1, $2), 4326), 9998)) * 100.0,
              y = -ST_Y(ST_Transform(ST_SetSRID(ST_MakePoint($1, $2), 4326), 9998)) * 100.0,
              archivechangedate = $4
            WHERE id = $3
            """,
            it.x, it.y, node_id, now,
        )
        await conn.execute(
            """
            UPDATE linesobj SET shape = ST_SetPoint(shape, 0, (SELECT shape FROM nodes WHERE id = $1)),
                   archivechangedate = $2
            WHERE nodeid1 = $1 AND shape IS NOT NULL AND COALESCE(removed, 0) = 0
            """,
            node_id, now,
        )
        await conn.execute(
            """
            UPDATE linesobj SET shape = ST_SetPoint(shape, ST_NumPoints(shape) - 1, (SELECT shape FROM nodes WHERE id = $1)),
                   archivechangedate = $2
            WHERE nodeid2 = $1 AND shape IS NOT NULL AND COALESCE(removed, 0) = 0
            """,
            node_id, now,
        )
        built = 0
        if build_missing:
            built = len(await conn.fetch(
                """
                UPDATE linesobj l
                SET shape = ST_MakeLine(n1.shape, n2.shape), archivechangedate = $2
                FROM nodes n1, nodes n2
                WHERE (l.nodeid1 = $1 OR l.nodeid2 = $1) AND l.shape IS NULL AND COALESCE(l.removed, 0) = 0
                  AND n1.id = l.nodeid1 AND n2.id = l.nodeid2
                  AND n1.shape IS NOT NULL AND n2.shape IS NOT NULL
                RETURNING l.id
                """,
                node_id, now,
            ))
        recalculated = 0
        if recalc:
            recalculated = len(await conn.fetch(
                """
                UPDATE heatpipesections h SET pipesectlength = round(ST_Length(l.shape)::numeric, 2)
                FROM linesobj l
                WHERE h.lineid = l.id AND (l.nodeid1 = $1 OR l.nodeid2 = $1)
                  AND l.shape IS NOT NULL AND COALESCE(l.removed, 0) = 0
                RETURNING h.lineid
                """,
                node_id,
            ))
        shift = None if old is None or old["shift"] is None else round(float(old["shift"]), 2)
        await op.audit("MOVE", "nodes", node_id, {
            "import": True, "row": it.row, "lng": it.x, "lat": it.y, "shift_m": shift,
            "recalculated_lines": recalculated, "built_lines": built,
        })
        report["updated_nodes"] += 1
        report["recalculated_lines"] += recalculated
        report["built_lines"] += built
        report["actions"].append({
            "row": it.row, "action": "move_node", "id": node_id, "shift_m": shift,
            "had_no_shape": bool(old and old["had_no_shape"]), "recalculated_lines": recalculated,
            "built_lines": built,
        })


_OPERATIONS = {"nodes": "IMPORT_NODES", "lines": "IMPORT_LINES", "coords": "IMPORT_COORDS"}
_HANDLERS = {"nodes": _import_nodes, "lines": _import_lines, "coords": _import_coords}


def validate_params(params: dict) -> dict:
    mode = params.get("mode")
    if mode not in MODES:
        raise ValueError(f"mode — одно из {', '.join(MODES)}")
    if (params.get("source_crs") or "auto") not in SOURCE_CRS:
        raise ValueError(f"source_crs — одно из {', '.join(SOURCE_CRS)}")
    if params.get("match_by", "code") not in ("id", "code"):
        raise ValueError("match_by — id или code")
    tol = float(params.get("snap_tolerance_m", 1.0))
    if not (0 <= tol <= MAX_SNAP_TOLERANCE_M):
        raise ValueError(f"Допуск привязки — 0…{MAX_SNAP_TOLERANCE_M:g} м")
    return params


async def run_import(src: ParsedSource, params: dict, *, dry_run: bool, actor: Optional[str]) -> dict:
    """Импорт (dry-run — выполнить и откатить). Ошибки строк без skip_errors → ImportRowErrors."""
    validate_params(params)
    mode = params["mode"]
    items = prepare_items(src, params)
    skip_errors = bool(params.get("skip_errors"))

    def new_report() -> dict:
        return {
            "mode": mode, "dry_run": dry_run, "total_rows": len(items), "fileid": params.get("fileid"),
            "created_nodes": 0, "created_lines": 0, "updated_nodes": 0, "recalculated_lines": 0,
            "built_lines": 0, "errors": [], "actions": [], "warnings": list(src.warnings),
        }

    async def body(conn, op):
        report = new_report()
        await _HANDLERS[mode](conn, op, items, params, report)
        report["errors"] = [{"row": it.row, "message": it.error} for it in items if it.error]
        report["ok_rows"] = report["total_rows"] - len(report["errors"])
        report["actions_truncated"] = len(report["actions"]) > MAX_ACTIONS_IN_REPORT
        report["actions"] = report["actions"][:MAX_ACTIONS_IN_REPORT]
        if not dry_run:
            if report["errors"] and not skip_errors:
                raise ImportRowErrors(report)
            if report["ok_rows"] == 0:
                raise ValueError("Нечего импортировать: все строки с ошибками")
        return report

    result = await _run(dry_run, body, None if dry_run else _OPERATIONS[mode], actor)
    if not dry_run:
        invalidate_outage_cache()
    return result
