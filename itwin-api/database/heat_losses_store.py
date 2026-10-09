"""Нормативные теплопотери: запуск, хранение и просмотр результатов (web к database/heat_losses_norm.py).

Расчёт — строка `calculation` (fileid = NULL, чтобы «последний расчёт фрагмента» для гидравлики
его не выбирал; фрагмент и параметры — в calc_params, module = "heat_losses_norm"):
- `ut_teplo_out` — удельные потери участков по месяцам (подача и обратка отдельными строками);
- `heatlosses_report_out` — листы по источникам (sql/migrations/20260928_heat_losses_report_out.sql).
Обе таблицы — *_out с calculationid: DELETE /api/v1/calculations/{id} удаляет их строки.

«Условия работы» (`prepare_work_conditions`) — перенос MainController.set_cond_env_temperatures:
строки heatLosesSourceMonths источника из климата сезона и развёрнутого температурного графика,
плюс строка heatLosesSource со значениями по умолчанию, если её нет.
"""

from __future__ import annotations

import io
import json
from datetime import date, datetime
from typing import Any, Optional

import asyncpg
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.utils import get_column_letter

from database import heat_losses_norm as hl

REPORT_TABLE = "heatlosses_report_out"
SHEETS = ("material_characteristics", "month_temperatures", "winter_norms", "summer_norms",
          "avg_month_loses", "avg_year_loses", "loads", "capacities", "heat_tests", "repair_heat_tests",
          "net_water_loses", "net_water_year_loses", "overalls", "tank_batteries", "tank_batteries_loses",
          "fillings", "pressings", "flushings_hs", "flushings", "sarz_flows", "sarz_rets")
LOSES_TYPE_LABELS = {"norm": "нормативные", "fact": "фактические"}
UT_TEPLO_COLUMNS = ("calculationid", "lineid", "externalsignlineid", "truba", "diametr", "tol", "diametr_usl",
                    "dlina", "name_typ", "kti", "kolwork", "kod_owner", "year", "q",
                    *(f"q{m:02d}" for m in range(1, 13)), "kod_ist")
PARAMS_LIKE = '%"module": "heat_losses_norm"%'


class HeatLossStoreError(RuntimeError):
    """Нет таблицы листов (миграция не применена) и т.п."""


def _date(v: Any) -> Optional[date]:
    if v is None:
        return None
    return v.date() if isinstance(v, datetime) else v


def season_label(season: dict[str, Any]) -> str:
    d1, d2 = _date(season.get("d1")), _date(season.get("d2"))
    fmt = lambda d: d.strftime("%d.%m.%Y") if d else "?"
    return f"{fmt(d1)} – {fmt(d2)}"


async def report_table_exists(conn: asyncpg.Connection) -> bool:
    return bool(await conn.fetchval("SELECT to_regclass($1) IS NOT NULL", REPORT_TABLE))


async def save_run(conn: asyncpg.Connection, inputs: hl.HeatLossInputs, result: dict[str, Any], *,
                   user: str) -> int:
    """Пишет расчёт в одной транзакции; id расчёта."""
    if not await report_table_exists(conn):
        raise HeatLossStoreError(
            "Нет таблицы heatlosses_report_out: примените sql/migrations/20260928_heat_losses_report_out.sql")
    totals = hl.summary_totals(result)
    params = {
        "module": hl.MODULE, "loses_type": inputs.loses_type,
        "season_id": inputs.season.get("id"), "season": season_label(inputs.season),
        "city": inputs.season.get("city"),
        "heat_source_ids": inputs.heat_source_ids, "ready_source_ids": result["ready_source_ids"],
        "fragment_id": inputs.fragment_id, "line_count": len(inputs.line_ids) if inputs.line_ids else None,
        "section_rows": len(result["sections"]), "totals": totals,
    }
    name = f"Теплопотери {LOSES_TYPE_LABELS.get(inputs.loses_type, inputs.loses_type)}, " \
           f"сезон {season_label(inputs.season)}"
    if inputs.fragment_id is not None:
        name += f", фрагмент {inputs.fragment_id}"
    async with conn.transaction():
        calc_id = await conn.fetchval(
            "INSERT INTO calculation (fileid, tn, date1, name, user_gid, calc_params) "
            "VALUES (NULL, NULL, now(), $1, $2, $3) RETURNING id",
            name, user, json.dumps(hl.json_safe(params), ensure_ascii=False))
        records = []
        for s in result["sections"]:
            records.append(tuple(
                calc_id if c == "calculationid" else s.get(c) for c in UT_TEPLO_COLUMNS))
        if records:
            await conn.copy_records_to_table("ut_teplo_out", records=records, columns=list(UT_TEPLO_COLUMNS))
        rows: list[tuple] = []
        for sheet in SHEETS:
            per: dict[Any, list] = {}
            for r in result[sheet]:
                per.setdefault(r.get("heatsourceid"), []).append(r)
            for hs, items in per.items():
                rows.append((calc_id, hs, sheet, json.dumps(hl.json_safe(items), ensure_ascii=False)))
        for hs, tot in (result.get("year_totals") or {}).items():
            rows.append((calc_id, int(hs), "year_totals", json.dumps(hl.json_safe(tot), ensure_ascii=False)))
        rows.append((calc_id, None, "sources", json.dumps(hl.json_safe(result["sources"]), ensure_ascii=False)))
        rows.append((calc_id, None, "totals", json.dumps(hl.json_safe(totals), ensure_ascii=False)))
        betas = sorted({(s["diametr_usl"], s["name_typ"], s["beta"]) for s in result["sections"]},
                       key=lambda x: (str(x[0]), str(x[1])))
        rows.append((calc_id, None, "betas", json.dumps(hl.json_safe(
            [{"diametr_usl": d, "name_typ": t, "beta": b} for d, t, b in betas]), ensure_ascii=False)))
        await conn.executemany(
            f"INSERT INTO {REPORT_TABLE} (calculationid, heatsourceid, sheet, rows) VALUES ($1, $2, $3, $4::jsonb)",
            rows)
    return calc_id


async def run(conn: asyncpg.Connection, *, season_id: int, heat_source_ids: Optional[list[int]] = None,
              fragment_id: Optional[int] = None, line_ids: Optional[list[int]] = None, user: str,
              save: bool = True, loses_type: str = "norm") -> dict[str, Any]:
    inputs = await hl.load_inputs(conn, season_id=season_id, heat_source_ids=heat_source_ids,
                                  fragment_id=fragment_id, line_ids=line_ids, loses_type=loses_type)
    if not inputs.heat_source_ids:
        raise hl.HeatLossInputError("Нет участков с источником теплоснабжения для расчёта")
    result = hl.compute(inputs)
    missing = [s for s in result["sources"] if not s["has_months"]]
    summary = {
        "sources": result["sources"], "ready_source_ids": result["ready_source_ids"],
        "missing_conditions": [s["id"] for s in missing], "section_rows": len(result["sections"]),
        "totals": hl.summary_totals(result),
    }
    if not result["ready_source_ids"]:
        raise hl.HeatLossInputError(
            f"Ни у одного источника не заданы «Условия работы» ({hl.LOSES_TABLES[loses_type]['months']}): "
            + ", ".join(str(s["id"]) for s in missing))
    if save:
        summary["calculation_id"] = await save_run(conn, inputs, result, user=user)
    return hl.json_safe(summary)


def _params(raw: Optional[str]) -> dict[str, Any]:
    try:
        value = json.loads(raw or "")
    except (TypeError, ValueError):
        return {}
    return value if isinstance(value, dict) else {}


async def list_runs(conn: asyncpg.Connection, *, limit: int = 50, offset: int = 0,
                    fragment_id: Optional[int] = None) -> dict[str, Any]:
    rows = await conn.fetch(
        "SELECT id, name, user_gid, date1, calc_params FROM calculation WHERE calc_params LIKE $1 "
        "ORDER BY id DESC", PARAMS_LIKE)
    items = []
    for r in rows:
        p = _params(r["calc_params"])
        if p.get("module") != hl.MODULE:
            continue
        if fragment_id is not None and p.get("fragment_id") != fragment_id:
            continue
        items.append({"id": r["id"], "name": r["name"], "user_gid": r["user_gid"],
                      "calculated_at": r["date1"].isoformat() if r["date1"] else None, "params": p})
    return {"total": len(items), "items": items[offset:offset + limit]}


async def get_run(conn: asyncpg.Connection, calculation_id: int) -> Optional[dict[str, Any]]:
    calc = await conn.fetchrow(
        "SELECT id, name, user_gid, date1, calc_params FROM calculation WHERE id = $1", calculation_id)
    if calc is None:
        return None
    params = _params(calc["calc_params"])
    if params.get("module") != hl.MODULE:
        return None
    if not await report_table_exists(conn):
        raise HeatLossStoreError("Нет таблицы heatlosses_report_out")
    sheets: dict[str, Any] = {s: [] for s in SHEETS}
    per_source_totals: dict[str, Any] = {}
    extra: dict[str, Any] = {}
    for r in await conn.fetch(
            f"SELECT heatsourceid, sheet, rows FROM {REPORT_TABLE} WHERE calculationid = $1 ORDER BY id",
            calculation_id):
        data = json.loads(r["rows"]) if isinstance(r["rows"], str) else r["rows"]
        if r["sheet"] in sheets:
            sheets[r["sheet"]].extend(data)
        elif r["sheet"] == "year_totals":
            per_source_totals[str(r["heatsourceid"])] = data
        else:
            extra[r["sheet"]] = data
    counts = await conn.fetchrow(
        "SELECT count(*)::int AS rows, count(DISTINCT lineid)::int AS lines, sum(dlina) AS length "
        "FROM ut_teplo_out WHERE calculationid = $1", calculation_id)
    return {
        "id": calc["id"], "name": calc["name"], "user_gid": calc["user_gid"],
        "calculated_at": calc["date1"].isoformat() if calc["date1"] else None, "params": params,
        "sources": extra.get("sources") or [], "totals": extra.get("totals") or params.get("totals"),
        "source_totals": per_source_totals, "betas": extra.get("betas") or [],
        "section_counts": dict(counts) if counts else {}, "sheets": sheets,
    }


def _month_days(month_temps: list[dict[str, Any]]) -> dict[str, dict[int, float]]:
    """Суток работы по месяцам источника (строки «Условий работы», без сводных)."""
    days: dict[str, dict[int, float]] = {}
    for r in month_temps:
        if r.get("m") is None or r["m"] > 12:
            continue
        per = days.setdefault(str(r["heatsourceid"]), {})
        per[r["m"]] = per.get(r["m"], 0) + (r.get("workcount") or 0)
    return days


def section_losses(row: dict[str, Any], beta: Optional[float], days: dict[int, float]) -> dict[str, Any]:
    """Потери участка: Гкал/ч в среднем за год и Гкал за год (q × длина × β × K)."""
    base = None
    if None not in (row.get("dlina"), beta, row.get("kti")):
        base = row["dlina"] * beta * row["kti"] / 1e6
    hourly = row["q"] * base if base is not None and row.get("q") is not None else None
    year = None
    if base is not None:
        vals = [row.get(f"q{m:02d}") * base * 24 * days.get(m, 0)
                for m in range(1, 13) if row.get(f"q{m:02d}") is not None]
        year = sum(vals) if vals else None
    return {"beta": beta, "loss_gcal_h": hourly, "loss_gcal_year": year}


async def get_sections(conn: asyncpg.Connection, calculation_id: int, *, heat_source_id: Optional[int] = None,
                       line_id: Optional[int] = None, page: int = 1, page_size: int = 100,
                       run: Optional[dict[str, Any]] = None) -> dict[str, Any]:
    run = run or await get_run(conn, calculation_id)
    if run is None:
        return {"total": 0, "items": []}
    betas = {(b["diametr_usl"], b["name_typ"]): b["beta"] for b in run.get("betas") or []}
    days = _month_days(run["sheets"]["month_temperatures"])
    clauses, args = ["u.calculationid = $1"], [calculation_id]
    if heat_source_id is not None:
        args.append(str(heat_source_id))
        clauses.append(f"u.kod_ist = ${len(args)}")
    if line_id is not None:
        args.append(line_id)
        clauses.append(f"u.lineid = ${len(args)}")
    where = " AND ".join(clauses)
    total = await conn.fetchval(f"SELECT count(*) FROM ut_teplo_out u WHERE {where}", *args)
    limit_sql = ""
    if page_size > 0:
        args += [page_size, (page - 1) * page_size]
        limit_sql = f" LIMIT ${len(args) - 1} OFFSET ${len(args)}"
    rows = await conn.fetch(
        f"""SELECT u.*, l.fileid, l.nodeid1, l.nodeid2,
                   -- участок внутренней схемы узла: точка узла-владельца (схема нарисована в условных координатах)
                   CASE WHEN l.shape IS NULL THEN NULL
                        ELSE ST_X(ST_Transform(COALESCE((SELECT owner.shape FROM nodes owner WHERE owner.id = l.internalnodeid), ST_PointOnSurface(l.shape)), 4326)) END AS longitude,
                   CASE WHEN l.shape IS NULL THEN NULL
                        ELSE ST_Y(ST_Transform(COALESCE((SELECT owner.shape FROM nodes owner WHERE owner.id = l.internalnodeid), ST_PointOnSurface(l.shape)), 4326)) END AS latitude
              FROM ut_teplo_out u LEFT JOIN linesobj l ON l.id = u.lineid
             WHERE {where} ORDER BY u.kod_ist, u.lineid, u.truba, u.id{limit_sql}""", *args)
    items = []
    for r in rows:
        d = dict(r)
        d.update(section_losses(d, betas.get((d.get("diametr_usl"), d.get("name_typ"))),
                                days.get(d.get("kod_ist") or "", {})))
        items.append(hl.json_safe(d))
    return {"total": total, "page": page, "page_size": page_size, "items": items}


# ---------------------------------------------------------------- Excel


MAT_HEADERS = [("typnet1", "Тип сети"), ("diameterexternal", "Наружный диаметр, мм"),
               ("lenp_n", "Длина надз., подающий, м"), ("leno_n", "Длина надз., обратный, м"),
               ("lenpodzp", "Длина подз., подающий, м"), ("lenpodzo", "Длина подз., обратный, м"),
               ("lenall", "Суммарная длина трубопроводов, м"), ("len_tr", "Протяжённость по трассе, м"),
               ("mn_p", "М надз., подающий, м²"), ("mn_o", "М надз., обратный, м²"),
               ("mp_p", "М подз., подающий, м²"), ("mp_o", "М подз., обратный, м²"), ("m", "М всего, м²"),
               ("vv", "Ёмкость, м³")]
TEMP_HEADERS = [("m1", "Месяц"), ("sezon1", "Период работы"), ("tn", "Наружного воздуха, °С"),
                ("tpod", "Воздуха в техн. подвале, °С"), ("tgr", "Грунта, °С"), ("tgp", "Сетевой воды, подающий, °С"),
                ("tgo", "Сетевой воды, обратный, °С"), ("tx", "Подпитки, °С"), ("workcount", "Продолжительность, дней")]
NORM_HEADERS = [("typnet1", "Тип сети"), ("a5000", "Класс / режим"), ("diametercondit", "Ду, мм"),
                ("lennp", "Надз.: длина под., м"), ("lenno", "Надз.: длина обр., м"),
                ("qnp", "Надз.: норма под., ккал/(м·ч)"), ("qno", "Надз.: норма обр., ккал/(м·ч)"),
                ("potnp", "Надз.: ТП под., ккал/ч"), ("potno", "Надз.: ТП обр., ккал/ч"),
                ("lenkp", "Канал.: длина под., м"), ("lenko", "Канал.: длина обр., м"),
                ("qk", "Канал.: норма 2-трубн., ккал/(м·ч)"), ("lenbp", "Бесканал.: длина под., м"),
                ("lenbo", "Бесканал.: длина обр., м"), ("qb", "Бесканал.: норма 2-трубн., ккал/(м·ч)"),
                ("potp", "Подземная: ТП 2-трубн., ккал/ч")]
LOSS_HEADERS = [("monthname", "Месяц"), ("potnp", "Надземная, подающий"), ("potno", "Надземная, обратный"),
                ("potpodz", "Подземная"), ("potall", "Сумма через изоляцию"), ("v1", "С утечкой"),
                ("vall", "Суммарные")]
_MONTH = ("monthname", "Месяц")
WATER_SHEETS = [
    # ключ листа, имя листа Excel (как у десктопа), заголовок, колонки, формат чисел, строка «ВСЕГО»
    ("loads", "Нагрузка", "РАСЧЁТНЫЕ НАГРУЗКИ ПОТРЕБИТЕЛЕЙ, Гкал/ч",
     [("heatsourcename", "Источник"), ("got_pr", "Отопление"), ("gvent_pr", "Вентиляция"), ("ggvs_pr", "ГВС ср.")],
     "0.000", True),
    ("capacities", "Емкость", "ОБЪЁМЫ ТЕПЛОВЫХ СЕТЕЙ И СИСТЕМ ТЕПЛОПОТРЕБЛЕНИЯ, м³",
     [("name", "Источник"), ("v1", "Магистральные"), ("v1leto", "Магистр., лето"), ("v2", "Распределительные"),
      ("v2leto", "Распред., лето"), ("vpodv", "В подвалах"), ("vpodvleto", "Подвалы, лето"), ("vob", "Обвязка"),
      ("vobleto", "Обвязка, лето"), ("vot", "Системы отопления"), ("votleto", "Отопл., лето"),
      ("vvent", "Системы вентиляции"), ("vventleto", "Вент., лето"), ("vgvs", "Системы ГВС"),
      ("vgvsleto", "ГВС, лето"), ("vall", "Всего"), ("vallleto", "Всего, лето"),
      ("podp", "Норм. утечка, м³/ч"), ("podpleto", "Норм. утечка, лето")], "0.00", True),
    ("heat_tests", "К(исп)", "КОЭФФИЦИЕНТЫ ПО РЕЗУЛЬТАТАМ ИСПЫТАНИЙ",
     [("name", "Источник"), *((c, c.replace("coeff", "").replace("norms", " N")) for c in hl.HEAT_TEST_COLUMNS)],
     "0.00", False),
    ("repair_heat_tests", "К(исп) Ремонт", "КОЭФФИЦИЕНТЫ ПО РЕЗУЛЬТАТАМ ИСПЫТАНИЙ (ПОСЛЕ РЕМОНТА)",
     [("name", "Источник"), *((c, c.replace("coeff", "").replace("norms", " N")) for c in hl.HEAT_TEST_COLUMNS)],
     "0.00", False),
    ("net_water_loses", "ТехнПСВ", "ТЕХНОЛОГИЧЕСКИЕ ПОТЕРИ СЕТЕВОЙ ВОДЫ, м³",
     [_MONTH, ("fillingg", "Заполнение"), ("avggpressingg", "Опрессовка"), ("avggflushingg", "Промывка"),
      ("avggsarzg", "САРЗ"), ("normg", "Нормативная утечка"), ("gall", "Всего")], "0.00", True),
    ("net_water_year_loses", "ТехнТП", "ТЕХНОЛОГИЧЕСКИЕ ПОТЕРИ ТЕПЛА С СЕТЕВОЙ ВОДОЙ, Гкал",
     [_MONTH, ("fillingq", "Заполнение"), ("avggpressingq", "Опрессовка"), ("avggflushingq", "Промывка"),
      ("avggsarzq", "САРЗ"), ("normq", "Нормативная утечка"), ("qall", "Всего")], "0.00", True),
    ("overalls", "ИТОГО", "ИТОГО: ПОТЕРИ ТЕПЛА (Гкал) И СЕТЕВОЙ ВОДЫ (м³)",
     [_MONTH, ("isolq", "Через изоляцию, Гкал"), ("qtb", "Баки-аккумуляторы, Гкал"),
      ("normq", "С норм. утечкой, Гкал"), ("reglq", "Регламентные, Гкал"), ("normg", "Норм. утечка, м³"),
      ("reglg", "Регламентные, м³"), ("gall", "Вода всего, м³"), ("allq", "Тепло всего, Гкал")], "0.00", True),
    ("tank_batteries", "БакАк", "БАКИ-АККУМУЛЯТОРЫ",
     [("mesto", "Место установки"), ("designcapacity", "Объём, м³"), ("quantity", "Количество"),
      ("height", "Высота, мм"), ("diameter", "Диаметр, мм")], "0.00", False),
    ("tank_batteries_loses", "БакАкТП", "ПОТЕРИ ТЕПЛА БАКАМИ-АККУМУЛЯТОРАМИ, Гкал/ч",
     [_MONTH, ("tn", "tн, °С"), ("tgo", "t2, °С"), ("monthloses", "Потери, Гкал/ч"), ("workcount", "Суток"),
      ("yearloses", "Потери за период, Гкал/ч")], "0.00", False),
    ("fillings", "Заполнение", "ПОТЕРИ С ЗАПОЛНЕНИЕМ ТРУБОПРОВОДОВ И СИСТЕМ",
     [_MONTH, ("magistralshare", "Доля магистр., %"), ("distsiteshare", "Доля распред., %"),
      ("heatingsystemshare", "Доля систем, %"), ("gmag", "G магистр., м³"), ("grs", "G распред., м³"),
      ("gtep", "G систем, м³"), ("nettemperature", "t сетевой воды"), ("tx", "t подпитки"),
      ("qms", "Q магистр., Гкал"), ("qrs", "Q распред., Гкал"), ("qtep", "Q систем, Гкал")], "0.00", True),
    ("pressings", "Опрессовка", "ПОТЕРИ ПРИ ОПРЕССОВКЕ",
     [_MONTH, ("opr", "Опрессовка"), ("tset", "t сетевой воды"), ("tn", "tн"), ("percent1", "Объём, %"),
      ("v1", "V магистр."), ("v2", "V распред."), ("vpodv", "V подвалы"), ("vobm", "V обвязка магистр."),
      ("vobr", "V обвязка прочая"), ("vall", "V всего"), ("avgqpressing", "Q, Гкал")], "0.00", True),
    ("flushings_hs", "ПромывкаСО", "ПОТЕРИ ПРИ ПРОМЫВКЕ СИСТЕМ ОТОПЛЕНИЯ",
     [_MONTH, ("flushinghs_temp1", "t воды"), ("tn", "tн"), ("flushinghs", "Кратность"), ("vot1", "V жилых"),
      ("vot2", "V общественных"), ("vall", "V всего"), ("q", "Q, Гкал")], "0.00", True),
    ("flushings", "ПромывкаТС", "ПОТЕРИ ПРИ ПРОМЫВКЕ ТЕПЛОВЫХ СЕТЕЙ",
     [_MONTH, ("kolv", "Трубопроводов"), ("flushing_temp1", "t воды"), ("tn", "tн"), ("flushing", "Кратность"),
      ("v1", "V магистр."), ("v2", "V распред."), ("vpodv", "V подвалы"), ("vobm", "V обвязка магистр."),
      ("vobr", "V обвязка прочая"), ("vall", "V всего"), ("q", "Q, Гкал")], "0.00", True),
]
for _key, _title, _sfx in (("sarz_flows", "САРЗ_Под", "ПОДАЮЩИЙ"), ("sarz_rets", "САРЗ_Обр", "ОБРАТНЫЙ")):
    WATER_SHEETS.append((_key, _title, f"СЛИВЫ САРЗ, {_sfx} ТРУБОПРОВОД",
                         [_MONTH, ("netwaterexp", "Расход слива, м³/ч"), ("workcount", "Суток"),
                          ("regcount", "Регуляторов"), ("avggsarzg", "G регуляторов, м³"),
                          ("regcountnode", "Регуляторов в узлах"), ("avggsarznodeg", "G в узлах, м³"),
                          ("avggsarzgall", "G всего, м³"), ("tgp", "t1"), ("tn", "tн"), ("qsarz", "Q, Гкал")],
                         "0.00", True))
SECTION_HEADERS = [("kod_ist", "Источник"), ("lineid", "Участок (lineid)"), ("truba", "Труба 1-под./2-обр."),
                   ("name_typ", "Прокладка"), ("diametr", "Dвн, мм"), ("diametr_usl", "Ду, мм"),
                   ("dlina", "Длина, м"), ("year", "Класс (год прокладки)"), ("kolwork", "Раб. 5000 ч"),
                   ("kti", "K"), ("beta", "β"), ("q", "q ср.год, ккал/(м·ч)"),
                   *((f"q{m:02d}", f"q {hl.MONTH_NAMES[m]}") for m in range(1, 13)),
                   ("loss_gcal_h", "Потери ср.год, Гкал/ч"), ("loss_gcal_year", "Потери за год, Гкал")]

_THIN = Side(style="thin", color="000000")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)


def _write_table(ws, title: str, subtitle: str, headers, blocks: list[tuple[str, list[dict]]],
                 totals: Optional[list[tuple[str, dict]]] = None, number_format: str = "0.000") -> None:
    """Таблица листа: заголовок, шапка, блоки по источникам, итоги (номер строки ведётся вручную:
    ws.max_row в openpyxl пересчитывается по всем ячейкам и делает запись квадратичной)."""
    ncol = len(headers)
    ws.cell(row=1, column=1, value=title).font = Font(bold=True, size=12)
    ws.cell(row=2, column=1, value=subtitle)
    head_row = 4
    for c, (_, label) in enumerate(headers, start=1):
        cell = ws.cell(row=head_row, column=c, value=label)
        cell.font = Font(bold=True)
        cell.alignment = Alignment(wrap_text=True, horizontal="center", vertical="center")
        cell.border = _BORDER
        ws.column_dimensions[get_column_letter(c)].width = 16
    row = head_row + 1

    def put(values: list, bold: bool = False) -> None:
        for c, value in enumerate(values, start=1):
            cell = ws.cell(row=row, column=c, value=value)
            cell.border = _BORDER
            if bold:
                cell.font = Font(bold=True)
            if isinstance(value, float):
                cell.number_format = number_format

    for block_title, rows in blocks:
        if block_title:
            ws.cell(row=row, column=1, value=block_title).font = Font(bold=True)
            row += 1
        for r in rows:
            put([r.get(k) for k, _ in headers])
            row += 1
    for label, values in totals or []:
        put([label] + [values.get(k) for k, _ in headers[1:]], bold=True)
        row += 1
    ws.freeze_panes = ws.cell(row=head_row + 1, column=1)
    _ = ncol


def build_excel(run: dict[str, Any], sections: list[dict[str, Any]]) -> bytes:
    """Листы десктопа (МатХарМаг … САРЗ_Обр, кроме ПТ/УТ) + участки и итоги."""
    wb = Workbook()
    sub = f"{run['name']} · расчёт {run['id']}"
    names = {str(s["id"]): f"{s.get('name') or s.get('sourcename') or 'Источник'} (№{s['id']})"
             for s in run.get("sources") or []}

    def blocks(sheet: str, sort=None):
        per: dict[str, list] = {}
        for r in run["sheets"][sheet]:
            per.setdefault(str(r.get("heatsourceid")), []).append(r)
        return [(names.get(hs, f"Источник №{hs}"), rows) for hs, rows in per.items()]

    ws = wb.active
    ws.title = "Итоги"
    total_headers = [("period", "Период"), ("potnp", "Надземная, подающий"), ("potno", "Надземная, обратный"),
                     ("potpodz", "Подземная"), ("potall", "Через изоляцию"), ("v1", "С утечкой"),
                     ("vall", "Всего")]
    rows_tot = []
    for hs, tot in (run.get("source_totals") or {}).items():
        for p, label in (("heating", "отопительный"), ("summer", "летний"), ("year", "год")):
            rows_tot.append({"period": f"{names.get(hs, hs)}: {label}", **tot[p]})
    grand = run.get("totals") or {}
    _write_table(ws, "НОРМИРУЕМЫЕ ГОДОВЫЕ ПОТЕРИ ТЕПЛА, Гкал", sub, total_headers,
                 [("По источникам", rows_tot)],
                 totals=[(f"ВСЕГО: {lbl}", grand.get(p) or {}) for p, lbl in
                         (("heating", "отопительный"), ("summer", "летний"), ("year", "год"))], number_format="0.00")
    _write_table(wb.create_sheet("МатХарМаг"), "МАТЕРИАЛЬНАЯ ХАРАКТЕРИСТИКА ТЕПЛОВОЙ СЕТИ", sub, MAT_HEADERS,
                 blocks("material_characteristics"), number_format="0.00")
    _write_table(wb.create_sheet("МесТемп"),
                 "СРЕДНЕМЕСЯЧНЫЕ И СРЕДНЕГОДОВЫЕ ТЕМПЕРАТУРЫ ОКРУЖАЮЩЕЙ СРЕДЫ И СЕТЕВОЙ ВОДЫ", sub, TEMP_HEADERS,
                 blocks("month_temperatures"), number_format="0.0")
    _write_table(wb.create_sheet("НормыЗима"),
                 "УДЕЛЬНЫЕ ТЕПЛОВЫЕ ПОТЕРИ ТРУБОПРОВОДАМИ ПРИ СРЕДНЕГОДОВЫХ УСЛОВИЯХ · период: отопительный", sub,
                 NORM_HEADERS, blocks("winter_norms"), number_format="0.00")
    _write_table(wb.create_sheet("НормыЛето"),
                 "УДЕЛЬНЫЕ ТЕПЛОВЫЕ ПОТЕРИ ТРУБОПРОВОДАМИ ПРИ СРЕДНЕГОДОВЫХ УСЛОВИЯХ · период: летний", sub,
                 NORM_HEADERS, blocks("summer_norms"), number_format="0.00")
    _write_table(wb.create_sheet("МесПотери"), "НОРМИРУЕМЫЕ СРЕДНЕМЕСЯЧНЫЕ ПОТЕРИ ТЕПЛА, Гкал/ч", sub,
                 LOSS_HEADERS, blocks("avg_month_loses"), number_format="0.000")
    year_blocks = []
    for title, rows in blocks("avg_year_loses"):
        year_blocks.append((title, rows[:14]))
    _write_table(wb.create_sheet("ГодПотери"), "НОРМИРУЕМЫЕ ГОДОВЫЕ ПОТЕРИ ТЕПЛА, Гкал", sub, LOSS_HEADERS,
                 year_blocks, number_format="0.00")
    for key, title, heading, headers, fmt, with_total in WATER_SHEETS:
        rows = run["sheets"].get(key) or []
        total = None
        if with_total and rows:
            vals = {k: sum(r[k] for r in rows if isinstance(r.get(k), (int, float)) and not isinstance(r.get(k), bool))
                    for k, _ in headers[1:]}
            total = [("ВСЕГО", vals)]
        _write_table(wb.create_sheet(title), heading, sub, headers, blocks(key), totals=total, number_format=fmt)
    _write_table(wb.create_sheet("Участки"), "УДЕЛЬНЫЕ НОРМАТИВНЫЕ ПОТЕРИ ПО УЧАСТКАМ (ut_teplo_out)", sub,
                 SECTION_HEADERS, [("", sections)], number_format="0.000")
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------- «Условия работы»


async def season_climate(conn: asyncpg.Connection, season: dict[str, Any]) -> list[dict[str, Any]]:
    """Климат сезона: строки heatLoses сезона, иначе города (как get_heat_loss_season)."""
    rows = await conn.fetch(
        """
        WITH ranked AS (
            SELECT hl.m, hl.tn, hl.tpod, hl.tgr,
                   row_number() OVER (PARTITION BY hl.m ORDER BY (hl.heatlosesmainid = $1) DESC NULLS LAST,
                                      hl.id DESC) AS priority
              FROM heatloses hl LEFT JOIN cities c ON c.id = hl.cityid
             WHERE hl.heatlosesmainid = $1 OR (hl.heatlosesmainid IS NULL AND lower(c.name) = lower($2))
        )
        SELECT m, tn, tpod, tgr FROM ranked WHERE priority = 1 ORDER BY m
        """, season["id"], season.get("city") or "")
    return [dict(r) for r in rows]


async def prepare_work_conditions(conn: asyncpg.Connection, *, heat_source_id: int, season_id: int,
                                  tx: Optional[dict[int, float]] = None, t_percent: Optional[float] = None,
                                  dry_run: bool = False) -> dict[str, Any]:
    """Строки heatLosesSourceMonths источника (14 шт.) и heatLosesSource по умолчанию.

    tx — температура подпитки по месяцам (у десктопа вводится вручную в «Условиях работы»);
    без неё — прежнее значение месяца источника, иначе DEFAULT колонки (0).
    """
    season = await conn.fetchrow("SELECT * FROM heatlosesmain WHERE id = $1", season_id)
    if season is None:
        raise hl.HeatLossInputError(f"Сезон {season_id} не найден")
    source = await conn.fetchrow("SELECT id, name FROM heatsources WHERE id = $1", heat_source_id)
    if source is None:
        raise hl.HeatLossInputError(f"Источник {heat_source_id} не найден")
    season = dict(season)
    climate = await season_climate(conn, season)
    graph = [dict(r) for r in await conn.fetch(
        "SELECT hsourceid, tn, t1, t2 FROM deployedtempgraphs WHERE hsourceid = $1 ORDER BY tn", heat_source_id)]
    map_tg, bounds = hl.temp_graph_maps(graph)
    start, end = _date(season["d1"]), _date(season["d2"])
    if start is None or end is None:
        raise hl.HeatLossInputError("У сезона не заданы даты начала и конца отопительного периода")
    rows = hl.work_condition_months(heat_source_id, climate, start, end, map_tg, bounds)
    old_tx = {r["m"]: r["tx"] for r in await conn.fetch(
        "SELECT m, tx FROM heatlosessourcemonths WHERE heatsourceid = $1 ORDER BY r", heat_source_id)}
    for r in rows:
        value = (tx or {}).get(r["m"])
        if value is None:
            value = old_tx.get(r["m"])
        r["tx"] = value
    has_source_row = bool(await conn.fetchval(
        "SELECT count(*) FROM heatlosessource WHERE heatsourceid = $1", heat_source_id))
    result = {"heat_source_id": heat_source_id, "season_id": season_id, "months": hl.json_safe(rows),
              "creates_source_parameters": not has_source_row, "dry_run": dry_run}
    if dry_run:
        return result
    async with conn.transaction():
        await conn.execute("DELETE FROM heatlosessourcemonths WHERE heatsourceid = $1", heat_source_id)
        for r in rows:
            await conn.execute(
                "INSERT INTO heatlosessourcemonths (heatsourceid, r, m, sezon, tn, tpod, tgr, tgp, tgo, workcount, tx) "
                "VALUES ($1, $2, $3, $4, $5, $6, $7, $8, $9, $10, COALESCE($11::float8, 0))",
                heat_source_id, r["r"], r["m"], r["sezon"], r["tn"], r["tpod"], r["tgr"], r["tgp"], r["tgo"],
                r["workcount"], r["tx"])
        if not has_source_row:
            await conn.execute("INSERT INTO heatlosessource (heatsourceid) VALUES ($1)", heat_source_id)
        if t_percent is not None:
            await conn.execute("UPDATE heatlosessource SET t_percent = $2 WHERE heatsourceid = $1",
                               heat_source_id, float(t_percent))
    return result
