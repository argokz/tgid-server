"""Телеметрия (фаза 3, «Цифровой двойник»): датчики сети, показания, сравнение «факт против модели».

Датчик стоит в узле (давление, температура) или на участке (расход) и измеряет подачу или обратку. Показания
приходят пачками от шлюза SCADA (POST /api/v2/telemetry/readings) или файлом выгрузки ОДПУ (CSV, Excel). Сравнение
с моделью — последний расчёт фрагмента, где стоит датчик:
- давление — напор над отметкой в узле, м (us_out.pih); показания манометра переводятся в метры водяного столба;
- температура — температура сетевой воды в узле, °C (us_out.t);
- расход — расход на участке, т/ч (ut_out.a13 по модулю: знак — направление относительно линии).

Хранение — tgid_sensor_reading (датчик, время) → значение в единицах датчика; с TimescaleDB — гипертаблица.
"""
from __future__ import annotations

import csv
import io
import math
import re
from dataclasses import dataclass
from datetime import datetime, timedelta, timezone
from typing import Any, Literal, Optional

from pydantic import BaseModel, Field, field_validator, model_validator

from database.ut_out_columns import (
    US_PIEZO_HEAD_M,
    US_SIGN_RETURN,
    US_SIGN_SUPPLY,
    US_TEMPERATURE_C,
    UT_FLOW_TH,
    UT_SIGN_RETURN,
    UT_SIGN_SUPPLY,
)

Kind = Literal["pressure", "temperature", "flow"]
Pipe = Literal["supply", "return"]


@dataclass(frozen=True)
class Unit:
    kind: str
    to_base: float  # множитель к единице сравнения: м вод. ст., °C, т/ч
    label: str


# Расход в м³/ч → т/ч: плотность сетевой воды ~975 кг/м³ (70–90 °C); погрешность — меньше порога «норма»
UNITS: dict[str, Unit] = {
    "kgf_cm2": Unit("pressure", 10.0, "кгс/см²"),
    "bar": Unit("pressure", 10.197, "бар"),
    "mpa": Unit("pressure", 101.97, "МПа"),
    "kpa": Unit("pressure", 0.10197, "кПа"),
    "m": Unit("pressure", 1.0, "м вод. ст."),
    "c": Unit("temperature", 1.0, "°C"),
    "t_h": Unit("flow", 1.0, "т/ч"),
    "m3_h": Unit("flow", 0.975, "м³/ч"),
}
BASE_UNITS = {"pressure": "м вод. ст.", "temperature": "°C", "flow": "т/ч"}
KIND_LABELS = {"pressure": "Давление", "temperature": "Температура", "flow": "Расход"}
PIPE_LABELS = {"supply": "подача", "return": "обратка"}

# Пороги расхождения с моделью: давление и расход — в процентах, температура — в градусах
OK_PCT, WARN_PCT = 5.0, 10.0
OK_DEG, WARN_DEG = 2.0, 5.0
# Показание старше — датчик «молчит»
STALE_MIN = 60
MAX_BATCH = 10_000
MAX_FILE_ROWS = 200_000
FUTURE_SLACK = timedelta(hours=24)
SERIES_POINTS = 500
_BUCKETS_MIN = (1, 5, 15, 30, 60, 180, 360, 720, 1440)


class TelemetryError(Exception):
    def __init__(self, status: int, message: str):
        super().__init__(message)
        self.status = status
        self.message = message


# --- датчики -----------------------------------------------------------------------------------

class SensorIn(BaseModel):
    code: str = Field(min_length=1, max_length=100, description="Код датчика в SCADA / ОДПУ")
    name: str = Field(min_length=1, max_length=200)
    kind: Kind
    pipe: Pipe = "supply"
    unit: str
    node_id: Optional[int] = Field(None, ge=1)
    line_id: Optional[int] = Field(None, ge=1)
    active: bool = True
    note: str = Field("", max_length=2000)

    @field_validator("code", "name")
    @classmethod
    def _strip(cls, v: str) -> str:
        v = v.strip()
        if not v:
            raise ValueError("пустое значение")
        return v

    @model_validator(mode="after")
    def _place(self) -> "SensorIn":
        unit = UNITS.get(self.unit)
        if unit is None:
            raise ValueError(f"единица — одна из: {', '.join(UNITS)}")
        if unit.kind != self.kind:
            raise ValueError(f"единица «{unit.label}» не для {KIND_LABELS[self.kind].lower()}")
        if self.kind == "flow":
            if not self.line_id:
                raise ValueError("расходомер ставится на участок: нужен line_id")
            self.node_id = None
        else:
            if not self.node_id:
                raise ValueError("датчик давления и температуры ставится в узел: нужен node_id")
            self.line_id = None
        return self


_SENSOR_SELECT = """
    SELECT s.*, coalesce(n.fileid, l.fileid) AS fileid,
           coalesce(nullif(n.nodename, ''), n.externalnodename) AS node_name,
           CASE WHEN s.node_id IS NOT NULL THEN ST_X(ST_Transform(n.shape, 4326))
                ELSE ST_X(ST_Transform(ST_PointOnSurface(l.shape), 4326)) END AS lon,
           CASE WHEN s.node_id IS NOT NULL THEN ST_Y(ST_Transform(n.shape, 4326))
                ELSE ST_Y(ST_Transform(ST_PointOnSurface(l.shape), 4326)) END AS lat
      FROM tgid_sensor s
      LEFT JOIN nodes n ON n.id = s.node_id
      LEFT JOIN linesobj l ON l.id = s.line_id
"""


def _iso(v: Any) -> Any:
    return v.isoformat() if isinstance(v, datetime) else v


def _sensor(row) -> dict[str, Any]:
    out = {k: _iso(row[k]) for k in row.keys()}
    unit = UNITS.get(row["unit"])
    out["unit_label"] = unit.label if unit else row["unit"]
    out["kind_label"] = KIND_LABELS.get(row["kind"], row["kind"])
    out["pipe_label"] = PIPE_LABELS.get(row["pipe"], row["pipe"])
    return out


async def _check_place(conn, body: SensorIn) -> None:
    if body.node_id and not await conn.fetchval("SELECT EXISTS (SELECT 1 FROM nodes WHERE id = $1)", body.node_id):
        raise TelemetryError(422, f"Узла {body.node_id} нет")
    if body.line_id and not await conn.fetchval("SELECT EXISTS (SELECT 1 FROM linesobj WHERE id = $1)", body.line_id):
        raise TelemetryError(422, f"Участка {body.line_id} нет")


async def get_sensor(conn, sensor_id: int) -> dict[str, Any]:
    row = await conn.fetchrow(f"{_SENSOR_SELECT} WHERE s.id = $1", sensor_id)
    if row is None:
        raise TelemetryError(404, f"Датчик {sensor_id} не найден")
    return _sensor(row)


async def create_sensor(conn, body: SensorIn, user: str) -> dict[str, Any]:
    await _check_place(conn, body)
    if await conn.fetchval("SELECT EXISTS (SELECT 1 FROM tgid_sensor WHERE code = $1)", body.code):
        raise TelemetryError(409, f"Датчик с кодом «{body.code}» уже есть")
    sensor_id = await conn.fetchval(
        """INSERT INTO tgid_sensor (code, name, kind, pipe, unit, node_id, line_id, active, note, created_by)
           VALUES ($1, $2, $3, $4, $5, $6, $7, $8, $9, $10) RETURNING id""",
        body.code, body.name, body.kind, body.pipe, body.unit, body.node_id, body.line_id, body.active, body.note, user)
    return await get_sensor(conn, sensor_id)


async def update_sensor(conn, sensor_id: int, body: SensorIn) -> dict[str, Any]:
    await get_sensor(conn, sensor_id)
    await _check_place(conn, body)
    if await conn.fetchval("SELECT EXISTS (SELECT 1 FROM tgid_sensor WHERE code = $1 AND id <> $2)", body.code, sensor_id):
        raise TelemetryError(409, f"Датчик с кодом «{body.code}» уже есть")
    await conn.execute(
        """UPDATE tgid_sensor SET code = $2, name = $3, kind = $4, pipe = $5, unit = $6, node_id = $7, line_id = $8,
                  active = $9, note = $10, updated_at = now() WHERE id = $1""",
        sensor_id, body.code, body.name, body.kind, body.pipe, body.unit, body.node_id, body.line_id, body.active,
        body.note)
    return await get_sensor(conn, sensor_id)


async def delete_sensor(conn, sensor_id: int) -> dict[str, Any]:
    sensor = await get_sensor(conn, sensor_id)
    count = await conn.fetchval("SELECT count(*) FROM tgid_sensor_reading WHERE sensor_id = $1", sensor_id)
    await conn.execute("DELETE FROM tgid_sensor WHERE id = $1", sensor_id)
    return {"id": sensor_id, "code": sensor["code"], "readings_deleted": count}


# --- показания ---------------------------------------------------------------------------------

class ReadingIn(BaseModel):
    sensor: str = Field(min_length=1, max_length=100, description="Код датчика")
    ts: datetime
    value: float

    @field_validator("value")
    @classmethod
    def _finite(cls, v: float) -> float:
        if not math.isfinite(v):
            raise ValueError("значение — конечное число")
        return v


class ReadingsIn(BaseModel):
    readings: list[ReadingIn] = Field(min_length=1, max_length=MAX_BATCH)


def _aware(ts: datetime) -> datetime:
    """Время без пояса — местное время сервера города (так пишут выгрузки ОДПУ)."""
    return ts if ts.tzinfo else ts.astimezone()


async def ingest(conn, readings: list[tuple[str, datetime, float]]) -> dict[str, Any]:
    """Записать показания (код, время, значение); повтор того же времени — замена значения.

    Неизвестные и выключенные коды, время из будущего — пропускаются и перечисляются в ответе.
    """
    codes = {code for code, _, _ in readings}
    ids = {r["code"]: r["id"] for r in await conn.fetch(
        "SELECT id, code FROM tgid_sensor WHERE code = ANY($1::text[]) AND active", list(codes))}
    limit = datetime.now(timezone.utc) + FUTURE_SLACK
    # одно время датчика дважды в пачке — побеждает последнее (как при повторной отправке)
    latest: dict[tuple[int, datetime], float] = {}
    future = 0
    for code, ts, value in readings:
        sid = ids.get(code)
        if sid is None:
            continue
        ts = _aware(ts)
        if ts > limit:
            future += 1
            continue
        latest[(sid, ts)] = float(value)
    if latest:
        await conn.execute(
            """INSERT INTO tgid_sensor_reading (sensor_id, ts, value)
               SELECT * FROM unnest($1::int[], $2::timestamptz[], $3::float8[])
               ON CONFLICT (sensor_id, ts) DO UPDATE SET value = EXCLUDED.value""",
            [k[0] for k in latest], [k[1] for k in latest], list(latest.values()))
    unknown = sorted(codes - set(ids))
    return {"accepted": len(latest), "skipped": len(readings) - len(latest), "unknown_codes": unknown[:50],
            "future": future}


# --- файл ОДПУ ---------------------------------------------------------------------------------

_TIME_HEAD = re.compile(r"^(time|ts|timestamp|datetime|date|время|дата|дата и время|дата/время)$")
_CODE_HEAD = re.compile(r"^(code|sensor|tag|код|датчик|тег)$")
_VALUE_HEAD = re.compile(r"^(value|val|значение|показание)$")
_TIME_FORMATS = ("%d.%m.%Y %H:%M:%S", "%d.%m.%Y %H:%M", "%d.%m.%Y", "%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M",
                 "%Y-%m-%d", "%d/%m/%Y %H:%M:%S", "%d/%m/%Y %H:%M")


def parse_time(v: Any) -> Optional[datetime]:
    if isinstance(v, datetime):
        return v
    text = str(v or "").strip()
    if not text:
        return None
    try:
        return datetime.fromisoformat(text.replace("Z", "+00:00"))
    except ValueError:
        pass
    for fmt in _TIME_FORMATS:
        try:
            return datetime.strptime(text, fmt)
        except ValueError:
            continue
    return None


def parse_number(v: Any) -> Optional[float]:
    if isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return float(v) if math.isfinite(v) else None
    text = str(v or "").strip().replace(" ", "").replace(" ", "").replace(",", ".")
    if not text:
        return None
    try:
        num = float(text)
    except ValueError:
        return None
    return num if math.isfinite(num) else None


def _table(data: bytes, filename: str) -> list[list[Any]]:
    if filename.lower().endswith((".xlsx", ".xlsm")) or data[:2] == b"PK":
        from openpyxl import load_workbook

        try:
            wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        except Exception as exc:  # noqa: BLE001 — битый файл
            raise TelemetryError(422, "Файл Excel не открывается") from exc
        return [list(r) for r in wb.worksheets[0].iter_rows(values_only=True)]
    for enc in ("utf-8-sig", "cp1251"):
        try:
            text = data.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    else:
        raise TelemetryError(422, "Кодировка файла не распознана (нужна UTF-8 или Windows-1251)")
    try:
        dialect: Any = csv.Sniffer().sniff(text[:4096], delimiters=";,\t")
    except csv.Error:
        dialect = _Semicolon  # не распознан — «;», как в выгрузках Excel с русской локалью
    return [row for row in csv.reader(io.StringIO(text), dialect)]


class _Semicolon(csv.excel):
    delimiter = ";"


def parse_file(data: bytes, filename: str) -> tuple[list[tuple[str, datetime, float]], list[str]]:
    """Выгрузка ОДПУ → показания. Два вида:
    - «длинный»: колонки время, код датчика, значение (по строке на показание);
    - «широкий»: первая колонка — время, остальные колонки — коды датчиков (как выгружают вычислители).
    """
    rows = [r for r in _table(data, filename) if any(c not in (None, "") for c in r)]
    if len(rows) < 2:
        raise TelemetryError(422, "В файле нет показаний")
    if len(rows) > MAX_FILE_ROWS:
        raise TelemetryError(413, f"Строк больше {MAX_FILE_ROWS} — разбейте файл")
    head = [str(c or "").strip() for c in rows[0]]
    low = [h.lower() for h in head]
    notes: list[str] = []
    out: list[tuple[str, datetime, float]] = []
    t_col = next((i for i, h in enumerate(low) if _TIME_HEAD.match(h)), None)
    c_col = next((i for i, h in enumerate(low) if _CODE_HEAD.match(h)), None)
    v_col = next((i for i, h in enumerate(low) if _VALUE_HEAD.match(h)), None)
    bad = 0
    if t_col is not None and c_col is not None and v_col is not None:
        for r in rows[1:]:
            get = lambda i: r[i] if i < len(r) else None  # noqa: E731
            ts, value, code = parse_time(get(t_col)), parse_number(get(v_col)), str(get(c_col) or "").strip()
            if ts is None or value is None or not code:
                bad += 1
                continue
            out.append((code, ts, value))
    else:
        t_col = 0 if t_col is None else t_col
        codes = [(i, h) for i, h in enumerate(head) if i != t_col and h]
        if not codes:
            raise TelemetryError(422, "Не найдены колонки: «время, код, значение» или «время, <код датчика>…»")
        for r in rows[1:]:
            ts = parse_time(r[t_col] if t_col < len(r) else None)
            if ts is None:
                bad += 1
                continue
            for i, code in codes:
                value = parse_number(r[i] if i < len(r) else None)
                if value is not None:
                    out.append((code, ts, value))
    if bad:
        notes.append(f"строк без времени или значения: {bad}")
    if not out:
        raise TelemetryError(422, "В файле не найдено ни одного показания")
    return out, notes


# --- чтение ------------------------------------------------------------------------------------

def bucket_minutes(start: datetime, end: datetime, points: int = SERIES_POINTS) -> int:
    span = max((end - start).total_seconds() / 60, 1)
    for b in _BUCKETS_MIN:
        if span / b <= points:
            return b
    return _BUCKETS_MIN[-1]


async def series(conn, sensor_id: int, start: datetime, end: datetime) -> dict[str, Any]:
    """Показания за период: среднее, минимум и максимум по интервалам (до SERIES_POINTS точек)."""
    sensor = await get_sensor(conn, sensor_id)
    start, end = _aware(start), _aware(end)
    if end <= start:
        raise TelemetryError(422, "Конец периода раньше начала")
    step = bucket_minutes(start, end)
    rows = await conn.fetch(
        f"""SELECT date_bin(interval '{step} minutes', ts, timestamptz '2000-01-01') AS t,
                   avg(value) AS avg, min(value) AS min, max(value) AS max, count(*) AS n
              FROM tgid_sensor_reading WHERE sensor_id = $1 AND ts >= $2 AND ts < $3
             GROUP BY 1 ORDER BY 1""", sensor_id, start, end)
    return {"sensor": sensor, "step_min": step,
            "points": [{"t": _iso(r["t"]), "avg": r["avg"], "min": r["min"], "max": r["max"], "n": r["n"]}
                       for r in rows]}


def deviation(kind: str, fact: float, model: float) -> dict[str, Any]:
    """Расхождение «факт − модель» и оценка: ok, warn, bad (давление и расход — в %, температура — в градусах)."""
    diff = fact - model
    if kind == "temperature":
        size = abs(diff)
        status = "ok" if size <= OK_DEG else "warn" if size <= WARN_DEG else "bad"
        return {"diff": round(diff, 3), "pct": None, "status": status}
    if abs(model) < 1e-6:
        return {"diff": round(diff, 3), "pct": None, "status": "ok" if abs(diff) < 1e-6 else "bad"}
    pct = diff / abs(model) * 100
    status = "ok" if abs(pct) <= OK_PCT else "warn" if abs(pct) <= WARN_PCT else "bad"
    return {"diff": round(diff, 3), "pct": round(pct, 1), "status": status}


async def _latest_calcs(conn, fileids: list[int]) -> dict[int, dict[str, Any]]:
    rows = await conn.fetch(
        "SELECT DISTINCT ON (fileid) fileid, id, date1 FROM calculation WHERE fileid = ANY($1::int[]) "
        "ORDER BY fileid, id DESC", fileids)
    return {r["fileid"]: {"id": r["id"], "at": _iso(r["date1"])} for r in rows}


async def overview(conn, *, only_active: bool = False) -> list[dict[str, Any]]:
    """Датчики с последним показанием, значением модели (последний расчёт фрагмента) и расхождением."""
    where = "WHERE s.active" if only_active else ""
    sensors = [_sensor(r) for r in await conn.fetch(f"{_SENSOR_SELECT} {where} ORDER BY s.name, s.id")]
    if not sensors:
        return []
    ids = [s["id"] for s in sensors]
    last = {r["sensor_id"]: r for r in await conn.fetch(
        """SELECT DISTINCT ON (sensor_id) sensor_id, ts, value FROM tgid_sensor_reading
            WHERE sensor_id = ANY($1::int[]) ORDER BY sensor_id, ts DESC""", ids)}
    calcs = await _latest_calcs(conn, sorted({s["fileid"] for s in sensors if s["fileid"]}))
    node_rows = [(s["node_id"], calcs[s["fileid"]]["id"]) for s in sensors
                 if s["node_id"] and s["fileid"] in calcs]
    line_rows = [(s["line_id"], calcs[s["fileid"]]["id"]) for s in sensors
                 if s["line_id"] and s["fileid"] in calcs]
    node_model: dict[tuple[int, int, int], Any] = {}
    if node_rows:
        for r in await conn.fetch(
                f"""SELECT u.nodeid, u.calculationid, u.externalsign, u.{US_PIEZO_HEAD_M}::float AS head,
                           u.{US_TEMPERATURE_C}::float AS t
                      FROM us_out u JOIN unnest($1::int[], $2::int[]) AS q(node, calc)
                        ON u.nodeid = q.node AND u.calculationid = q.calc""",
                [n for n, _ in node_rows], [c for _, c in node_rows]):
            node_model[(r["nodeid"], r["calculationid"], r["externalsign"])] = r
    line_model: dict[tuple[int, int, int], float] = {}
    if line_rows:
        for r in await conn.fetch(
                f"""SELECT u.lineid, u.calculationid, u.externalsignlineid AS sign, abs(u.{UT_FLOW_TH})::float AS flow
                      FROM ut_out u JOIN unnest($1::int[], $2::int[]) AS q(line, calc)
                        ON u.lineid = q.line AND u.calculationid = q.calc""",
                [n for n, _ in line_rows], [c for _, c in line_rows]):
            line_model[(r["lineid"], r["calculationid"], r["sign"])] = r["flow"]
    stale_before = datetime.now(timezone.utc) - timedelta(minutes=STALE_MIN)
    for s in sensors:
        unit = UNITS[s["unit"]]
        reading = last.get(s["id"])
        s["last"] = {"ts": _iso(reading["ts"]), "value": reading["value"]} if reading else None
        s["stale"] = reading is None or reading["ts"] < stale_before
        calc = calcs.get(s["fileid"])
        s["calculation"] = calc
        model_base: Optional[float] = None
        if calc and s["node_id"]:
            sign = US_SIGN_SUPPLY if s["pipe"] == "supply" else US_SIGN_RETURN
            r = node_model.get((s["node_id"], calc["id"], sign))
            if r is not None:
                model_base = r["head"] if s["kind"] == "pressure" else r["t"]
        elif calc and s["line_id"]:
            sign = UT_SIGN_SUPPLY if s["pipe"] == "supply" else UT_SIGN_RETURN
            model_base = line_model.get((s["line_id"], calc["id"], sign))
        # модель — в единицах датчика (как его видит оператор) и в единицах сравнения
        s["model"] = None if model_base is None else round(model_base / unit.to_base, 4)
        s["base_unit"] = BASE_UNITS[s["kind"]]
        s["deviation"] = None
        if reading is not None and model_base is not None:
            s["deviation"] = deviation(s["kind"], reading["value"] * unit.to_base, model_base)
            s["deviation"]["base_unit"] = BASE_UNITS[s["kind"]]
    return sensors


def geojson(sensors: list[dict[str, Any]]) -> dict[str, Any]:
    features = []
    for s in sensors:
        if s.get("lon") is None:
            continue
        dev = s.get("deviation") or {}
        features.append({
            "type": "Feature", "id": s["id"],
            "geometry": {"type": "Point", "coordinates": [s["lon"], s["lat"]]},
            "properties": {"id": s["id"], "name": s["name"], "kind": s["kind"], "pipe": s["pipe"],
                           "value": (s.get("last") or {}).get("value"), "unit": s["unit_label"],
                           "status": "stale" if s.get("stale") else dev.get("status", "no_model"),
                           "pct": dev.get("pct"), "diff": dev.get("diff")},
        })
    return {"type": "FeatureCollection", "features": features}


# --- тревоги -----------------------------------------------------------------------------------

ALERT_LABELS = {"bad": "расхождение с моделью", "silent": "датчик молчит", "normal": "вернулся в норму"}
MAX_ALERTS = 500


def alert_level(sensor: dict[str, Any]) -> Optional[str]:
    """Состояние для тревог: silent — нет свежих показаний, bad — расхождение, normal — норма или «внимание».

    None — оценки нет (нет расчёта): состояние не меняется, тревоги нет.
    """
    if sensor.get("stale"):
        return "silent"
    dev = sensor.get("deviation")
    if not dev:
        return None
    return "bad" if dev["status"] == "bad" else "normal"


async def check_alerts(conn) -> list[int]:
    """Сверить состояние действующих датчиков с прошлым и записать тревоги о переходах; вернуть id новых.

    Первая оценка датчика тревоги не даёт (иначе включение тревог разослало бы всё разом) — только запоминается.
    Вызывать под замком (database/webhooks.py): два процесса не должны записать одну тревогу дважды.
    """
    sensors = await overview(conn, only_active=True)
    prev = {r["id"]: r["alert_status"] for r in await conn.fetch(
        "SELECT id, alert_status FROM tgid_sensor WHERE active")}
    new_ids: list[int] = []
    for s in sensors:
        level = alert_level(s)
        before = prev.get(s["id"])
        if level is None or level == before:
            continue
        await conn.execute("UPDATE tgid_sensor SET alert_status = $2, alert_since = now() WHERE id = $1", s["id"], level)
        if before is None:
            continue
        dev = s.get("deviation") or {}
        last = s.get("last") or {}
        new_ids.append(await conn.fetchval(
            """INSERT INTO tgid_sensor_alert (sensor_id, status, prev_status, value, model, diff, pct)
               VALUES ($1, $2, $3, $4, $5, $6, $7) RETURNING id""",
            s["id"], level, before, last.get("value"), s.get("model"), dev.get("diff"), dev.get("pct")))
    return new_ids


async def list_alerts(conn, *, limit: int = 100, open_only: bool = False) -> list[dict[str, Any]]:
    where = "WHERE a.acknowledged_at IS NULL AND a.status <> 'normal'" if open_only else ""
    rows = await conn.fetch(
        f"""SELECT a.*, s.code, s.name, s.kind, s.pipe, s.unit, s.node_id, s.line_id
              FROM tgid_sensor_alert a JOIN tgid_sensor s ON s.id = a.sensor_id
              {where} ORDER BY a.at DESC, a.id DESC LIMIT $1""", max(1, min(MAX_ALERTS, limit)))
    out = []
    for r in rows:
        item = {k: _iso(r[k]) for k in r.keys()}
        item["label"] = ALERT_LABELS.get(r["status"], r["status"])
        unit = UNITS.get(r["unit"])
        item["unit_label"] = unit.label if unit else r["unit"]
        out.append(item)
    return out


async def acknowledge(conn, alert_id: int, user: str) -> dict[str, Any]:
    row = await conn.fetchrow(
        """UPDATE tgid_sensor_alert SET acknowledged_by = $2, acknowledged_at = now()
            WHERE id = $1 AND acknowledged_at IS NULL RETURNING id, acknowledged_at""", alert_id, user)
    if row is None:
        raise TelemetryError(404, f"Тревога {alert_id} не найдена или уже принята")
    return {"id": row["id"], "acknowledged_by": user, "acknowledged_at": _iso(row["acknowledged_at"])}


# --- пьезометр ---------------------------------------------------------------------------------

async def facts_at_nodes(conn, node_ids: list[int]) -> dict[int, dict[str, Any]]:
    """Последние показания датчиков давления и температуры в узлах: давление — в метрах водяного столба
    (как напор над отметкой у модели), температура — °C. Ключи: p_supply, p_return, t_supply, t_return, at, stale.
    """
    if not node_ids or not await conn.fetchval("SELECT to_regclass('public.tgid_sensor_reading') IS NOT NULL"):
        return {}
    rows = await conn.fetch(
        """SELECT s.node_id, s.kind, s.pipe, s.unit, r.ts, r.value
             FROM tgid_sensor s
             JOIN LATERAL (SELECT ts, value FROM tgid_sensor_reading WHERE sensor_id = s.id
                            ORDER BY ts DESC LIMIT 1) r ON true
            WHERE s.active AND s.node_id = ANY($1::int[]) AND s.kind IN ('pressure', 'temperature')""", node_ids)
    stale_before = datetime.now(timezone.utc) - timedelta(minutes=STALE_MIN)
    out: dict[int, dict[str, Any]] = {}
    for r in rows:
        item = out.setdefault(r["node_id"], {})
        key = f"{'p' if r['kind'] == 'pressure' else 't'}_{r['pipe']}"
        unit = UNITS.get(r["unit"])
        if unit is None:
            continue
        item[key] = round(r["value"] * unit.to_base, 3)
        # время и «молчит» — по самому старому из показаний узла
        if "at" not in item or r["ts"] < item["_ts"]:
            item["_ts"] = r["ts"]
            item["at"] = _iso(r["ts"])
            item["stale"] = r["ts"] < stale_before
    for item in out.values():
        item.pop("_ts", None)
    return out
