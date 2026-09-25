"""Калькулятор шайб и элеваторов против десктопных эталонов.

- элеватор: sety/dross/drsh2.py вызывается напрямую (тот же код, что в расчётном движке);
- шайбы и расходы: формулы ячеек бланка gid8/python/dross/dross.py (G20..G43),
  перенесённые сюда дословно, потому что dross.py — скрипт, строящий Excel при импорте.
"""

import io
import math

import openpyxl
import pytest

from database.throttling_calc import (
    calculate_elevator_parameters,
    calculate_orifice_diameter,
    calculate_orifice_plate_full,
    calculate_throttling_sheet,
    generate_throttling_excel,
)
import sety.dross.drsh2 as drsh2_module
from sety.dross.drsh2 import drsh2


@pytest.fixture(autouse=True)
def _quiet_desktop_print(monkeypatch):
    # w_print десктопа требует CLI-конфигурацию sety; в тестах сообщения не нужны
    monkeypatch.setattr(drsh2_module, "w_print", lambda *a, **k: None)


def _excel_round(x, digits):
    """ROUND из Excel: половина — от нуля (Python round — банковское)."""
    q = 10 ** digits
    return math.floor(abs(x) * q + 0.5) / q * (1 if x >= 0 else -1)


def _dross_orifice(g, h, coeff=10.0):
    """=IF(ROUND(k*POWER(G*G/H,1/4),1)<3,3,ROUND(k*POWER(G*G/H,1/4),1))"""
    d = _excel_round(coeff * math.pow(g * g / h, 0.25), 1)
    return 3 if d < 3 else d


def _drsh2_number(nal):
    # drsh2: нулевой номер (dg <= 10) — стандартного элеватора №0 нет, у нас это №1
    return max(1, nal)


@pytest.mark.parametrize("g", [0.5, 2.0, 7.3, 15.0, 40.0])
@pytest.mark.parametrize("h", [3.0, 8.0, 20.0, 45.0])
def test_orifice_matches_dross_formula(g, h):
    assert calculate_orifice_diameter(flow_g=g, delta_h=h) == pytest.approx(_dross_orifice(g, h), abs=0.05)
    assert calculate_orifice_diameter(flow_g=g, delta_h=h, coeff=9.6) == pytest.approx(
        _dross_orifice(g, h, 9.6), abs=0.05)


def test_orifice_not_applicable_without_head():
    assert calculate_orifice_diameter(flow_g=5, delta_h=0) is None
    res = calculate_orifice_plate_full(flow_g=5, delta_h=4, scheme="bezelevator")
    assert res["diameter_orifice_mm"] is None and "Недостаточный" in res["warning"]


def test_orifice_requires_inputs():
    with pytest.raises(ValueError, match="расход G или тепловую нагрузку"):
        calculate_orifice_plate_full(delta_h=10)
    with pytest.raises(ValueError, match="располагаемый напор"):
        calculate_orifice_plate_full(flow_g=5)


@pytest.mark.parametrize("g, u, hc, h", [
    (10.0, 1.4, 1.5, 20.0),     # пример из аудита: десктоп — горловина 37.6 мм, №6
    (2.0, 1.2, 1.0, 12.0),
    (0.4, 2.2, 0.8, 30.0),
    (25.0, 1.8, 2.0, 40.0),
])
def test_elevator_matches_drsh2(g, u, hc, h):
    t1, t2 = 150.0, 70.0
    t3 = (t1 + u * t2) / (1 + u)          # u = (t1 - t3) / (t3 - t2)
    res = calculate_elevator_parameters(flow_g=g, p1=6.0 + h / 10.0, p2=6.0, t1=t1, t2=t2, t3=t3,
                                        delta_h_system=hc)
    ferr, hrc, dc, nal = drsh2(g, 3.0, u, hc, h, "тест")
    assert res["mixing_ratio_u"] == pytest.approx(u, abs=0.01)
    assert res["nozzle_diameter_mm"] == pytest.approx(dc, abs=0.06)
    neck = 8.5 * math.sqrt(g * (1 + u) / math.sqrt(hc))   # dg из drsh2 (функция его не возвращает)
    assert res["mixing_chamber_diameter_mm"] == pytest.approx(neck, abs=0.06)
    assert res["elevator_number"] == _drsh2_number(nal)


def test_elevator_worked_example():
    t3 = (150.0 + 1.4 * 70.0) / 2.4
    res = calculate_elevator_parameters(flow_g=10, p1=8.0, p2=6.0, t1=150, t2=70, t3=t3, delta_h_system=1.5)
    assert res["mixing_chamber_diameter_mm"] == pytest.approx(37.6, abs=0.1)
    assert res["elevator_number"] == 6


def test_elevator_head_warnings():
    low = calculate_elevator_parameters(flow_g=5, p1=6.1, p2=6.0, t1=150, t2=70, t3=95, delta_h_system=1.5)
    assert any("меньше требуемого" in w for w in low["warnings"])
    high = calculate_elevator_parameters(flow_g=5, p1=12.0, p2=6.0, t1=150, t2=70, t3=95, delta_h_system=1.5)
    assert any("повышенном напоре" in w for w in high["warnings"])


def test_sheet_flows_match_dross():
    data = {"p1": 7.0, "p2": 4.5, "q_heating_gcal": 0.8, "q_vent_gcal": 0.12, "q_gvs_gcal": 0.3,
            "t1": 132, "t2": 70}
    calc = calculate_throttling_sheet(data)
    q_kcal = lambda gcal: gcal * 1_000_000
    assert calc["g_ot"] == pytest.approx(q_kcal(0.8) / ((132 - 70) * 1000))          # G20
    assert calc["g_v"] == pytest.approx(q_kcal(0.12) / ((132 - 70) * 1000))         # G21
    assert calc["g_gvs"] == pytest.approx(q_kcal(0.3) / 2.4 / 60000)                # G17, G22
    assert calc["h_ras"] == pytest.approx(25.0)                                      # G23
    assert calc["d_heater"] == pytest.approx(_dross_orifice(calc["g_ot"], 20.0), abs=0.05)  # G41 — расход отопления
    assert calc["d_gvs_circulation"] == pytest.approx(_dross_orifice(calc["g_gvs"], 20.0), abs=0.05)  # G43
    assert calc["d_pre_nozzle"] == pytest.approx(_dross_orifice(calc["g_ot"], 25.0), abs=0.05)  # G32


def test_generate_throttling_excel_without_placeholders():
    data = {"district": "Район 1", "consumer_name": "Дом", "p1": 6.0, "p2": 4.0,
            "q_heating_gcal": 1.2, "q_vent_gcal": 0.2, "q_gvs_gcal": 0.4, "t1": 150.0, "t2": 70.0,
            "signers": [{"position": "Инженер", "name": "Иванов И.И."}]}
    wb = openpyxl.load_workbook(io.BytesIO(generate_throttling_excel(data)))
    ws = wb.active
    text = " ".join(str(c.value) for row in ws.iter_rows() for c in row if c.value is not None)
    assert "Иванов И.И." in text
    assert "Чупин" not in text and "Бегимбетов" not in text
    assert ws.protection.sheet is True
