"""Быстрая запись строк ведомостей (reports_generator._inject_rows) даёт ту же книгу, что openpyxl.

Строки данных листов со значениями str/int/float/Decimal/bool/None сериализуются в XML напрямую,
минуя openpyxl (pt всей сети: 50–60 с → около 6 с). Значения ячеек, листы, шапка, закрепление
и автофильтр должны совпадать с прежним путём, где все строки писал openpyxl.
"""

import datetime as dt
import io
import zipfile
from decimal import Decimal

import openpyxl
from lxml import etree

import reports_generator as R

TRICKY_ROWS = [
    ["a & b <c> \"d\" 'e'", 1, 2.5, Decimal("3.25"), True, None],
    ["  пробелы по краям  ", -7, 1e-9, Decimal("0"), False, ""],
    ["упр\x01авл\x1fяющие\x0b", 0, float("nan"), None, None, "строка\nвторая"],
    ["возврат\rкаретки", 10**15, float("inf"), Decimal("-1.5"), True, "x" * 40000],
    [None, None, None, None, None, None],
    ["короткая строка"],
]
HEADERS = ["Текст", "Целое", "Дробное", "Decimal", "Флаг", "Прочее"]


def _reference(data: R.SheetData, notes):
    """Прежний путь: все строки пишет openpyxl write_only."""
    wb = openpyxl.Workbook(write_only=True)
    sheets = [((data.title or "Ведомость")[:31], data)] + [(t[:31], s) for t, s in data.extra_sheets]
    for title, sheet in sheets:
        R._write_sheet(wb.create_sheet(title), sheet.headers, sheet.rows, with_data=True)
    if notes:
        ws = wb.create_sheet("Примечание")
        ws.column_dimensions["A"].width = 120
        for note in notes:
            ws.append([note])
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def _read(content: bytes):
    wb = openpyxl.load_workbook(io.BytesIO(content))
    result = {}
    for ws in wb.worksheets:
        values = [tuple(r) for r in ws.iter_rows(values_only=True)]
        result[ws.title] = (values, ws.freeze_panes, ws.auto_filter.ref)
    return result


def test_fast_rows_match_openpyxl_values_and_layout():
    extra = R.SheetData(headers=["№", "Имя"], rows=[[i, f"имя {i}"] for i in range(1, 2501)], total=2500)
    data = R.SheetData(headers=HEADERS, rows=TRICKY_ROWS, total=len(TRICKY_ROWS),
                       extra_sheets=[("Второй лист", extra)])
    notes = ["Примечание к ведомости"]

    fast = R._render_workbook("Ведомость", data, notes)
    ref = _reference(data, notes)

    assert _read(fast) == _read(ref)
    # все XML-части книги корректны (Excel не откроет книгу с битым XML)
    with zipfile.ZipFile(io.BytesIO(fast)) as zf:
        for name in zf.namelist():
            if name.endswith(".xml") or name.endswith(".rels"):
                etree.fromstring(zf.read(name))
            # перезаписанные листы сжаты, как и остальные части
            assert zf.getinfo(name).compress_type == zipfile.ZIP_DEFLATED


def test_control_characters_removed_and_long_strings_cut():
    data = R.SheetData(headers=HEADERS, rows=TRICKY_ROWS, total=len(TRICKY_ROWS))
    values = _read(R._render_workbook("Ведомость", data, []))["Ведомость"][0]
    assert values[3][0] == "управляющие"
    assert values[4][5] == "x" * R._EXCEL_MAX_STR


def test_sheets_with_dates_fall_back_to_openpyxl():
    rows = [["ТУ-1", dt.date(2025, 3, 1)], ["ТУ-2", dt.datetime(2025, 3, 2, 10, 30)]]
    assert R._fast_width(rows) is None
    data = R.SheetData(headers=["Номер", "Дата"], rows=rows, total=2)
    values = _read(R._render_workbook("ТУ", data, []))["ТУ"][0]
    assert values[1] == ("ТУ-1", dt.datetime(2025, 3, 1))
    assert values[2] == ("ТУ-2", dt.datetime(2025, 3, 2, 10, 30))


def test_sheet_paths_follow_workbook_order_not_creation_ids():
    sheets = [(f"Лист {i}", R.SheetData(headers=["a"], rows=[[i]], total=1)) for i in range(2, 5)]
    data = R.SheetData(headers=["a"], rows=[[1]], total=1, extra_sheets=sheets)
    result = _read(R._render_workbook("Лист 1", data, ["примечание"]))
    assert list(result) == ["Лист 1", "Лист 2", "Лист 3", "Лист 4", "Примечание"]
    for i in range(1, 5):
        assert result[f"Лист {i}"][0] == [("a",), (i,)]
