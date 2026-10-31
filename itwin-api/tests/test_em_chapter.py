"""Глава «Электронная модель» (database/em_chapter.py): связность, классы диаметра, копии фрагментов, Word."""
from io import BytesIO

from database import em_chapter as em


def test_graph_stats_rings_components_and_feeds():
    # кольцо 1-2-3-1 с источником 1, тупик 3-4; отдельная часть 5-6 без питания; 7-8 через закрытый участок
    lines = {1: (1, 2), 2: (2, 3), 3: (3, 1), 4: (3, 4), 5: (5, 6), 6: (4, 7), 7: (7, 8)}
    g = em.graph_stats(lines, sources={1, 8}, closed={6})
    assert g["nodes"] == 8 and g["edges"] == 6
    assert g["components"] == 3 and g["rings"] == 1
    assert g["fed_components"] == 2 and g["unfed_nodes"] == 2 and g["multi_source_components"] == 0
    g2 = em.graph_stats(lines, sources={1, 8}, closed=set())
    assert g2["multi_source_components"] == 1 and g2["components"] == 2


def test_diameter_class_labels():
    assert em.diameter_class(80) == "до 100" and em.diameter_class(100) == "до 100"
    assert em.diameter_class(150) == "101–200" and em.diameter_class(1000) == "свыше 700"
    assert em.diameter_class(None) == "нет данных" and em.diameter_class(0) == "нет данных"


def test_top_rows_folds_tail():
    rows = [[f"r{i}", i, float(i)] for i in range(em.TOP_ROWS + 5)]
    out = em.top_rows(rows, 2)
    assert len(out) == em.TOP_ROWS + 1 and out[0][0] == f"r{em.TOP_ROWS + 4}"
    assert out[-1][0] == "прочие (5)" and out[-1][1] == sum(range(5))


def test_copy_fragments_by_share():
    lines = [
        {"id": 1, "fileid": 10, "ends": "a"}, {"id": 2, "fileid": 10, "ends": "b"},
        {"id": 3, "fileid": 20, "ends": "a"}, {"id": 4, "fileid": 20, "ends": "b"}, {"id": 5, "fileid": 20, "ends": "c"},
        {"id": 6, "fileid": 5, "ends": "c"}, {"id": 7, "fileid": 5, "ends": "d"}, {"id": 8, "fileid": 5, "ends": "e"},
    ]
    # фрагмент 10 целиком повторяет более новый 20; у 5 совпадает 1 из 3 — не копия
    assert em.copy_fragments(lines) == {10: 1.0}


def test_docx_has_all_items():
    from docx import Document

    data = {"generated_at": "2026-10-31T18:00+05:00", "basis": em.BASIS, "fragments": [{"id": 1, "name": "тм-1"}],
            "copies": [{"id": 2, "name": "тм-1 копия", "share": 0.9}], "lines": 10, "km": 1.5,
            "sections": [em.section(k, em.OK, [f"текст {k}"],
                                    [{"title": "Т", "header": ["a", "b"], "rows": [["x", 1]]}] if k == "б" else [])
                         for k in em.ITEMS]}
    doc = Document(BytesIO(em.to_docx(data, "Город")))
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "Глава 3. Электронная модель системы теплоснабжения — Город" in text
    assert all(f"(подпункт «{k}»)" in text for k in em.ITEMS) and "тм-1 копия (90 %)" in text
    assert len(doc.tables) == 2 and len(doc.tables[0].rows) == len(em.ITEMS) + 1
