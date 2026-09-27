"""Excel-отчёты десктопа (database/excel_reports.py): каталог, SQL и шаблоны без базы.

Выполнение каждого запроса на живой базе — scripts/golden/reports_smoke.py.
"""

import datetime as dt
import io
import re

import openpyxl
import pytest

from database import excel_reports as X

SQL_FILES = sorted(p.stem for p in X.SQL_DIR.glob("*.sql"))
REPORT_SQL = sorted(p for p in SQL_FILES if not p.startswith("_"))
# справочники без фрагмента, как в десктопе
GLOBAL_SQL = {"vh_organizacii", "vh_teplosnabzhayushaya_sistema"}

# Приметы диалекта MS SQL / dBase-схемы десктопа, которых не должно остаться
MSSQL_PATTERNS = {
    "TOP n": r"\bSELECT\s+TOP\b",
    "IIF()": r"\bIIF\s*\(",
    "ISNULL()": r"\bISNULL\s*\(",
    "[идентификатор]": r"\[[^\]\n]+\]",
    "$fileID$/$calculationID$": r"\$[A-Za-z]+\$",
    "NVARCHAR": r"\bN?VARCHAR\s*\(\s*MAX\s*\)|\bNVARCHAR\b",
    "GETDATE()": r"\bGETDATE\s*\(",
    "YEAR()": r"\bYEAR\s*\(",
    "LEN()": r"\bLEN\s*\(",
    "CHARINDEX()": r"\bCHARINDEX\s*\(",
    "DATEDIFF()": r"\bDATEDIFF\s*\(",
    "CONVERT()": r"\bCONVERT\s*\(",
    "APPLY": r"\b(OUTER|CROSS)\s+APPLY\b",
    "dbo.": r"\bdbo\.",
    "строка + строка": r"'\s*\+|\+\s*'",
    "#include": r"#include",
}


def _code(sql: str) -> str:
    """Текст без комментариев и строковых литералов — чтобы проверять только код."""
    sql = "\n".join(line for line in sql.splitlines() if not line.lstrip().startswith("--"))
    return re.sub(r"'(?:[^']|'')*'", "''", sql)


def test_sql_directory_holds_ported_desktop_queries():
    # 44 SQL gid6 excel2/sql2: 39 перенесены, 5 — нет (причины в NOT_PORTED)
    assert len(REPORT_SQL) == 39
    assert len(X.NOT_PORTED) == 5


@pytest.mark.parametrize("name", REPORT_SQL)
def test_report_sql_is_postgresql(name):
    sql = X.load_sql(name)
    code = _code(sql)
    for label, pattern in MSSQL_PATTERNS.items():
        assert not re.search(pattern, code, re.IGNORECASE), f"{name}: {label}"
    assert "{{" not in sql, f"{name}: неразрешённое включение"


@pytest.mark.parametrize("name", REPORT_SQL)
def test_report_sql_uses_only_positional_params(name):
    sql = X.load_sql(name)
    params = {int(n) for n in re.findall(r"\$(\d+)", sql)}
    if name not in GLOBAL_SQL:
        assert 1 in params, f"{name}: нет параметра фрагмента"
    assert params <= {1, 2}, name
    if 2 in params:
        assert 1 in params, name
    # значения не подставляются в текст: нет format-плейсхолдеров
    assert not re.search(r"\{[a-z_]*\}|%s|%\(", _code(sql)), name


@pytest.mark.parametrize("name", REPORT_SQL)
def test_report_sql_brackets_balanced_and_single_statement(name):
    code = _code(X.load_sql(name))
    assert code.count("(") == code.count(")"), name
    assert ";" not in code, name
    assert re.match(r"\s*(SELECT|WITH)\b", code, re.IGNORECASE), name


def test_include_expands_shared_subqueries():
    sql = X.load_sql("out_pt_teplo")
    assert "FROM realconsumers rc" in sql and "FROM generalizedconsumers gc" in sql
    assert X.sql_params(sql) == 2
    assert X.sql_params(X.load_sql("out_pt_raschetnye_nagruzki")) == 1


def test_unknown_or_escaping_include_is_rejected():
    with pytest.raises(ValueError):
        X._read("../../main")
    with pytest.raises(ValueError):
        X._read("no_such_report")


def test_catalog_sheets_reference_existing_sql_and_template_sheets():
    ids = [r.id for r in X.REPORTS]
    assert len(ids) == len(set(ids))
    used = set()
    for report in X.REPORTS:
        assert report.sheets, report.id
        for sheet in report.sheets:
            assert sheet.sql in REPORT_SQL, (report.id, sheet.sql)
            used.add(sheet.sql)
        if report.template is None:
            assert report.headers, report.id
            continue
        wb = openpyxl.load_workbook(X.TEMPLATE_DIR / f"{report.template}.xlsx")
        for sheet in report.sheets:
            ws = wb.worksheets[sheet.sheet - 1]
            assert ws.title == sheet.title, (report.id, sheet.sheet, ws.title)
            # строка шапки — строка номеров колонок «1, 2, 3 ...» десктопного шаблона
            assert ws.cell(sheet.header_row, 1).value in (1, 1.0, "1"), (report.id, sheet.sheet)
            assert ws.cell(sheet.header_row + 1, 1).value is None, (report.id, sheet.sheet)
    # каждый перенесённый запрос выводится хотя бы в одном отчёте
    assert set(REPORT_SQL) - used == set()


def test_catalog_marks_reports_that_need_a_calculation():
    items = {i["id"]: i for i in X.catalog()}
    assert items["out_ut"]["uses_calculation"] is True
    assert items["out_ut"]["params"]["calculation_id"] == "optional"
    assert items["gut"]["uses_calculation"] is False
    assert items["gut"]["params"]["calculation_id"] is None
    assert all(i["params"]["fragment_id"] == "required" for i in items.values())


def test_dop_values_follow_desktop_rules():
    calc = {"name": "Расчет фактического режима", "date1": dt.datetime(2026, 9, 26, 21, 53), "tn": -25.0}
    assert X.dop_value("$out_rezhim", calc) == "Фактический"
    assert X.dop_value("$out_rezhim", {**calc, "name": "Расчет планового режима"}) == "Плановый"
    assert X.dop_value("$out_date", calc) == "26.09.2026"
    assert X.dop_value("$out_time", calc) == "21:53"
    assert X.dop_value("$TnZ", calc) == "Tнаруж -25°С"
    assert X.dop_value("$out_rezhim", None) == ""


def test_render_writes_rows_under_template_header_and_skips_service_columns():
    report = X.get_report("out_zd")
    results = []
    for sheet in report.sheets:
        cols = ["kod_p", "uzel_p", "_ni_id"] if sheet.sql.startswith("out_") else ["id", "sost"]
        rows = [("6-8", "ТК1", 42)] if sheet.sql.startswith("out_") else [(1, "открыто")]
        results.append((sheet, cols, rows))
    calc = {"id": 3, "name": "x", "date1": dt.datetime(2026, 9, 26), "tn": -25.0}
    wb = openpyxl.load_workbook(io.BytesIO(X.render_report(report, calc, results)))
    assert wb.sheetnames == [s.title for s in report.sheets]
    ws = wb["Все"]
    assert [ws.cell(16, c).value for c in (1, 2, 3)] == ["6-8", "ТК1", None]
    assert wb["Вх.Задвижки"].cell(13, 2).value == "открыто"


def test_render_labels_extra_columns_and_formats_dates():
    report = X.get_report("gus")
    sheet = report.sheets[0]
    cols = ["id", "kod", "uzel", "pr", "g1", "g2", "name_typ", "extra", "changed"]
    rows = [(1, "6-8", "У1", "П", 1.5, 2.5, "тип", "x", dt.date(2024, 5, 3))]
    ws = openpyxl.load_workbook(io.BytesIO(X.render_report(report, None, [(sheet, cols, rows)]))).worksheets[0]
    assert ws.cell(sheet.header_row + 1, 2).value == "6-8"
    assert ws.cell(sheet.header_row, 9).value == "changed"
    assert ws.cell(sheet.header_row + 1, 9).number_format == "DD.MM.YYYY"


def test_catalog_routes_registered():
    import main

    paths = main.app.openapi()["paths"]
    assert "/api/reports/catalog" in paths
    assert "/api/reports/catalog/{report_id}/excel" in paths
    params = {p["name"]: p for p in paths["/api/reports/catalog/{report_id}/excel"]["get"]["parameters"]}
    assert params["fragment_id"]["required"] is True
    assert params["calculation_id"]["required"] is False
