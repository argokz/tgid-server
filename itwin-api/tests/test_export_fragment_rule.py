"""QA F13/F14: единое правило «участок во фрагменте» (linesobj.fileid) для SHP/DXF/GeoJSON и
ведомости Excel по фрагменту без молчаливой обрезки (без БД — подменённое соединение)."""

from __future__ import annotations

import asyncio
import io
import os
from contextlib import asynccontextmanager

import openpyxl
import pytest
from fastapi import HTTPException
from fastapi.testclient import TestClient

os.environ.setdefault("JWT_SECRET", "test-secret")

import main  # noqa: E402
import reports_generator as rg  # noqa: E402
from database import export_shp, file_jobs  # noqa: E402
from database.fragment_filter import LINE_IN_FRAGMENTS_SQL, parse_fragment_ids  # noqa: E402
from routers import exports_p4  # noqa: E402


class FakeConn:
    def __init__(self, *, total=0, rows=None):
        self.calls: list[tuple[str, tuple]] = []
        self.total = total
        self.rows = rows or []

    async def fetch(self, sql, *args):
        self.calls.append((sql, args))
        return self.rows[: args[-1]] if args and isinstance(args[-1], int) else self.rows

    async def fetchval(self, sql, *args):
        self.calls.append((sql, args))
        return self.total


def _acquire(conn):
    @asynccontextmanager
    async def _cm():
        yield conn

    return _cm


# --- F13: SHP / DXF / GeoJSON ------------------------------------------------------------

@pytest.mark.parametrize("sql", [
    export_shp._LINES_SQL,
    exports_p4._DXF_LINES_SQL,
    exports_p4._LINES_WITH_PASSPORT_SQL,
])
def test_line_exports_select_fragment_by_line_fileid(sql):
    assert LINE_IN_FRAGMENTS_SQL in sql
    assert "n1.fileid" not in sql.lower()
    assert "join nodes" not in sql.lower()
    assert "COALESCE(lo.removed, 0) = 0" in sql
    assert "lo.shape IS NOT NULL" in sql


def test_shp_export_passes_fragments_to_line_and_node_queries(monkeypatch):
    conn = FakeConn()
    monkeypatch.setattr(export_shp, "acquire_conn", _acquire(conn))
    asyncio.run(export_shp.export_network_to_shp([74, 74], limit=123))
    assert [c[1] for c in conn.calls] == [([74], 123), ([74], 123)]
    nodes_sql = conn.calls[1][0]
    # узлы — свои и концы участков фрагмента (часть узлов фрагмента 74 числится во фрагменте 99)
    assert "lo.fileid = ANY($1::int[])" in nodes_sql and "lo.nodeid1" in nodes_sql


def test_shp_export_whole_network_passes_null_fragments(monkeypatch):
    conn = FakeConn()
    monkeypatch.setattr(export_shp, "acquire_conn", _acquire(conn))
    asyncio.run(export_shp.export_network_to_shp(None))
    assert conn.calls[0][1][0] is None


def test_parse_fragment_ids_merges_dedupes_and_rejects_garbage():
    assert parse_fragment_ids(74, "99, 74,,5") == [5, 74, 99]
    assert parse_fragment_ids(None, None) is None
    for bad in ("1,x", "0", "-3"):
        with pytest.raises(HTTPException) as exc:
            parse_fragment_ids(None, bad)
        assert exc.value.status_code == 400


def test_shp_endpoint_rejects_bad_fragments_instead_of_exporting_whole_network():
    with TestClient(main.app) as client:
        assert client.get("/api/export/shp", params={"fragments": "74,abc"}).status_code == 400


def test_dxf_endpoint_accepts_fragment_list():
    params = {p["name"] for p in main.app.openapi()["paths"]["/api/export/dxf"]["get"]["parameters"]}
    assert {"fragment_id", "fragments", "limit"} <= params


# --- F14: ведомости Excel ----------------------------------------------------------------

def _sheet(content: bytes):
    wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True)
    data = list(wb.worksheets[0].iter_rows(values_only=True))
    notes = [r[0] for r in wb["Примечание"].iter_rows(values_only=True)] if "Примечание" in wb.sheetnames else []
    return data, notes


def _line_rows(n):
    return [{"id": i, "nodeid1": i, "nodeid2": i + 1, "pipesectlength": 10.0, "diameterinternal": 100,
             "diametercondit": 100, "diameterexternal": 108, "wallthickness": 4} for i in range(1, n + 1)]


def test_ut_by_fragment_filters_by_line_fileid_and_is_not_capped(monkeypatch):
    conn = FakeConn(total=4018, rows=_line_rows(4018))
    monkeypatch.setattr(rg, "acquire_conn", _acquire(conn))
    report = asyncio.run(rg.build_excel_report("ut", fragment_ids=[74]))
    count_sql, count_args = conn.calls[0]
    rows_sql, rows_args = conn.calls[1]
    assert LINE_IN_FRAGMENTS_SQL in count_sql and LINE_IN_FRAGMENTS_SQL in rows_sql
    assert "JOIN LATERAL" in rows_sql  # одна строка на участок
    assert count_args == ([74],) and rows_args == ([74], rg.EXCEL_MAX_DATA_ROWS)
    assert (report.rows, report.total, report.truncated) == (4018, 4018, False)
    data, notes = _sheet(report.content)
    assert len(data) == 4018 + 1
    assert notes == ["Отбор по фрагментам: 74."]
    assert report.headers()["X-Report-Fragments"] == "74"


def test_whole_network_over_limit_is_reported_not_silently_cut(monkeypatch):
    monkeypatch.setattr(rg, "MAX_NETWORK_REPORT_ROWS", 3)
    conn = FakeConn(total=10, rows=_line_rows(10))
    monkeypatch.setattr(rg, "acquire_conn", _acquire(conn))
    report = asyncio.run(rg.build_excel_report("ut"))
    assert conn.calls[1][1] == (None, 3)
    assert (report.rows, report.total, report.truncated) == (3, 10, True)
    assert report.headers()["X-Report-Truncated"] == "1"
    data, notes = _sheet(report.content)
    assert len(data) == 3 + 1
    assert len(notes) == 1 and "3 строк из 10" in notes[0] and "Выберите фрагмент" in notes[0]


def test_whole_network_default_limit_is_far_above_old_5000():
    assert rg.report_limit(None) == rg.MAX_NETWORK_REPORT_ROWS >= 200_000
    assert rg.report_limit([74]) == rg.EXCEL_MAX_DATA_ROWS


def test_paged_registers_query_each_selected_fragment(monkeypatch):
    seen = []

    async def fake_armatures(conn, *, page, page_size, fragment_id=None):
        seen.append((fragment_id, page_size))
        items = [{"id": fragment_id * 100 + i, "line_id": i} for i in range(2)]
        return {"items": items, "total": 2}

    monkeypatch.setattr(rg, "get_network_armatures", fake_armatures)
    monkeypatch.setattr(rg, "acquire_conn", _acquire(FakeConn()))
    report = asyncio.run(rg.build_excel_report("zd", fragment_ids=[99, 74]))
    assert [s[0] for s in seen] == [74, 99]
    assert report.rows == report.total == 4
    data, _ = _sheet(report.content)
    assert [r[0] for r in data[1:]] == [7400, 7401, 9900, 9901]


def test_register_without_fragment_binding_says_so(monkeypatch):
    async def fake_tu(conn, *, page, page_size):
        return {"items": [{"id": 1}], "total": 1}

    import database.technical_conditions as tc

    monkeypatch.setattr(tc, "get_technical_conditions", fake_tu)
    monkeypatch.setattr(rg, "acquire_conn", _acquire(FakeConn()))
    report = asyncio.run(rg.build_excel_report("tu", fragment_ids=[74]))
    assert report.fragment_filter_applied is False
    assert report.headers()["X-Report-Fragments"] == ""
    _, notes = _sheet(report.content)
    assert "не привязана к фрагментам" in notes[0]


def test_report_filename_marks_fragments():
    assert rg.report_filename("ut", fragment_ids=[74]) == "report_ut_f74.xlsx"
    assert rg.report_filename("ut", fragment_ids=[1, 2]) == "report_ut_frag.xlsx"
    assert rg.report_filename("tu-balance", year=2025) == "report_tu-balance_2025.xlsx"


def test_excel_endpoint_passes_fragments_and_exposes_completeness_headers(monkeypatch):
    captured = {}

    async def fake_build(doc_type, *, year=None, fragment_ids=None):
        captured.update(doc_type=doc_type, fragment_ids=fragment_ids)
        return rg.ExcelReport(b"xlsx", rows=5, total=9, fragment_ids=fragment_ids, fragment_filter_applied=True)

    monkeypatch.setattr("routers.reports.build_excel_report", fake_build)
    with TestClient(main.app) as client:
        r = client.get("/api/reports/excel/ut", params={"fragments": "99,74"})
    assert r.status_code == 200
    assert captured == {"doc_type": "ut", "fragment_ids": [74, 99]}
    assert r.headers["x-report-truncated"] == "1"
    assert r.headers["x-report-total"] == "9"
    assert "X-Report-Truncated" in r.headers["access-control-expose-headers"]
    assert "report_ut_frag.xlsx" in r.headers["content-disposition"]


def test_file_job_report_excel_accepts_fragments():
    assert file_jobs.validate_params("report_excel", {"doc_type": "ut", "fragments": [74]})["fragments"] == [74]
    with pytest.raises(ValueError):
        file_jobs.validate_params("report_excel", {"doc_type": "ut", "fragments": [0]})


def test_file_job_report_excel_forwards_fragments_and_headers(monkeypatch):
    async def fake_build(doc_type, *, year=None, fragment_ids=None):
        return rg.ExcelReport(b"xlsx", rows=2, total=2, fragment_ids=fragment_ids, fragment_filter_applied=True)

    monkeypatch.setattr(rg, "build_excel_report", fake_build)
    result = asyncio.run(file_jobs.build("report_excel", {"doc_type": "ut", "fragments": [74]}))
    assert result.filename == "report_ut_f74.xlsx"
    assert result.headers["X-Report-Fragments"] == "74"
    assert result.headers["X-Report-Truncated"] == "0"
