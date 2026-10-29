"""Отключения (database/outages.py, routers/outages.py): снимок изоляции, пересечения, статусы, уведомления."""

import asyncio
import io
import json
from contextlib import asynccontextmanager
from datetime import datetime, timezone

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from database import outages as og

ISOLATION = {
    "summary": {"total_with_downstream_load_gcal_h": 3.5},
    "valves_to_close": [
        {"id": 2, "display_name": "З-2", "lineid": 20, "required": False, "nominal_diameter": 300,
         "state_name": "открыта", "lng": 76.9, "lat": 43.2},
        {"id": 1, "display_name": "З-1", "lineid": 10, "required": True, "nominal_diameter": 500,
         "state_name": "открыта", "lng": 76.91, "lat": 43.21},
    ],
    "affected_consumers": [{"id": 5, "consumer_type": "real", "node_id": 50, "name": "Школа №1", "total_load": 1.5,
                            "longitude": 76.9, "latitude": 43.2}],
    "downstream_consumers": [{"id": 6, "consumer_type": "generalized", "node_id": 60, "name": "Квартал 3",
                              "total_load": 2.0, "longitude": 76.92, "latitude": 43.22}],
    "geojson": {"isolated_pipes": {"features": [{"properties": {"id": 10}}]},
                "downstream_pipes": {"features": [{"properties": {"id": 11}}]}},
}


class _Tx:
    async def __aenter__(self):
        return self

    async def __aexit__(self, *exc):
        return False


class FakeConn:
    def __init__(self, *, revision_ok=True):
        self.revision_ok = revision_ok
        self.executed: list[tuple[str, tuple]] = []
        self.inserted: tuple = ()
        self.others: list[dict] = []

    def transaction(self):
        return _Tx()

    async def execute(self, sql, *args):
        self.executed.append((sql, args))

    async def fetchval(self, sql, *args):
        if "ST_AsGeoJSON" in sql:
            assert args[0] == [10, 11]  # зона — по участкам изоляции и ниже по сети
            return json.dumps({"type": "Polygon", "coordinates": [[[76.9, 43.2], [76.91, 43.2], [76.9, 43.21],
                                                                   [76.9, 43.2]]]})
        if sql.lstrip().startswith("INSERT INTO tgid_outage"):
            self.inserted = args
            return 3
        if "UPDATE tgid_outage" in sql:
            self.executed.append((sql, args))
            return 3 if self.revision_ok else None
        return None

    async def fetchrow(self, sql, *args):
        if "FROM linesobj" in sql:
            return {"fileid": 74, "lon": 76.905, "lat": 43.205}
        return None

    async def fetch(self, sql, *args):
        if "FROM tgid_outage" in sql:
            return self.others
        return []


def _outage(**kw):
    base = {"id": 3, "number": "О-2026-00003", "title": "Замена задвижки", "kind": "planned", "kind_label": "Плановое",
            "status": "planned", "status_label": "Запланировано", "line_id": 10, "node_id": None,
            "planned_start": "2026-11-02T04:00:00+00:00", "planned_end": "2026-11-02T13:00:00+00:00",
            "reason": "замена задвижки", "consumers_count": 2, "load_gcal_h": 3.5, "revision": 2,
            "valves": [{"id": 1, "name": "З-1", "line_id": 10, "required": True, "diameter": 500, "state": "открыта",
                        "lng": 76.9, "lat": 43.2}],
            "consumers": [{"id": 5, "type": "real", "node_id": 50, "name": "Школа №1", "load": 1.5, "lng": 76.9,
                           "lat": 43.2, "downstream": False},
                          {"id": 6, "type": "generalized", "node_id": 60, "name": "Квартал 3", "load": 2.0,
                           "lng": 76.92, "lat": 43.22, "downstream": True}]}
    base.update(kw)
    return base


@pytest.fixture
def isolation(monkeypatch):
    calls = []

    async def fake(conn, line_id=None, node_id=None):
        calls.append((line_id, node_id))
        return ISOLATION

    monkeypatch.setattr(og, "simulate_outage_isolation", fake)
    return calls


@pytest.fixture
def stub_get(monkeypatch):
    state = {"outage": _outage()}

    async def fake_get(conn, outage_id):
        return dict(state["outage"])

    monkeypatch.setattr(og, "get_outage", fake_get)
    return state


def test_create_takes_snapshot(isolation, stub_get):
    conn = FakeConn()
    start = datetime(2026, 11, 2, 9, tzinfo=timezone.utc)
    asyncio.run(og.create(conn, og.OutageIn(title="Замена", line_id=10, planned_start=start,
                                            planned_end=start.replace(hour=18), public=True), "ivan"))
    args = conn.inserted
    assert isolation == [(10, None)]
    # title, kind, line, node, work_order, fileid, lon, lat, start, end, reason, public, note, count, load, valves…
    assert args[2:6] == (10, None, None, 74) and args[13:15] == (2, 3.5)
    valves, consumers = json.loads(args[15]), json.loads(args[16])
    assert [v["id"] for v in valves] == [1, 2]  # обязательные — первыми
    assert [c["downstream"] for c in consumers] == [False, True] and consumers[1]["type"] == "generalized"
    assert json.loads(args[17])["type"] == "Polygon"
    assert any("потребителей 2, задвижек 2" in a[4] for s, a in conn.executed if "tgid_outage_event" in s)


def test_create_validation(isolation):
    start = datetime(2026, 11, 2, 9, tzinfo=timezone.utc)
    with pytest.raises(og.OutageError, match="раньше начала"):
        asyncio.run(og.create(FakeConn(), og.OutageIn(title="x", line_id=1, planned_start=start,
                                                      planned_end=start), "u"))
    with pytest.raises(og.OutageError, match="участок или узел"):
        asyncio.run(og.create(FakeConn(), og.OutageIn(title="x", planned_start=start,
                                                      planned_end=start.replace(hour=12)), "u"))


def test_isolation_error_is_422(monkeypatch):
    async def fail(conn, line_id=None, node_id=None):
        raise ValueError("Трубопровод 99 удалён.")

    monkeypatch.setattr(og, "simulate_outage_isolation", fail)
    start = datetime(2026, 11, 2, 9, tzinfo=timezone.utc)
    with pytest.raises(og.OutageError) as e:
        asyncio.run(og.create(FakeConn(), og.OutageIn(title="x", line_id=99, planned_start=start,
                                                      planned_end=start.replace(hour=12)), "u"))
    assert e.value.status == 422 and "удалён" in e.value.message


def test_conflicts_need_shared_consumers():
    conn = FakeConn()
    conn.others = [
        {"id": 4, "number": "О-4", "title": "a", "status": "planned", "planned_start": datetime(2026, 11, 2, 8),
         "planned_end": datetime(2026, 11, 2, 12), "consumers": json.dumps([{"type": "real", "id": 5}])},
        {"id": 5, "number": "О-5", "title": "b", "status": "planned", "planned_start": datetime(2026, 11, 2, 8),
         "planned_end": datetime(2026, 11, 2, 12), "consumers": json.dumps([{"type": "generalized", "id": 5}])},
    ]
    found = asyncio.run(og.conflicts(conn, _outage()))
    assert [(c["number"], c["shared_consumers"]) for c in found] == [("О-4", 1)]  # тот же id, но другой тип — нет


def test_transitions(stub_get):
    with pytest.raises(og.OutageError) as e:
        asyncio.run(og.transition(FakeConn(), 3, og.OutageTransition(to="finished", revision=2), "u"))
    assert e.value.status == 409
    conn = FakeConn()
    asyncio.run(og.transition(conn, 3, og.OutageTransition(to="active", revision=2, note="задвижки закрыты"), "u"))
    sql, args = next((s, a) for s, a in conn.executed if "UPDATE tgid_outage" in s)
    assert "actual_start = COALESCE(actual_start, $4)" in sql and args[2] == "active"
    with pytest.raises(og.OutageError) as e:
        asyncio.run(og.transition(FakeConn(revision_ok=False), 3, og.OutageTransition(to="announced", revision=1), "u"))
    assert e.value.status == 409
    stub_get["outage"] = _outage(status="active", status_label="Идёт")
    with pytest.raises(og.OutageError) as e:
        asyncio.run(og.recompute(FakeConn(), 3, 2, "u"))
    assert e.value.status == 409


def test_announcement_and_workbook():
    text = og.announcement(_outage())
    assert text.startswith("Плановое отключение теплоснабжения О-2026-00003: с 02.11.2026")
    assert "Причина: замена задвижки." in text and "(2): Школа №1; Квартал 3." in text
    many = _outage(consumers=[{"id": i, "type": "real", "name": f"Дом {i}"} for i in range(30)])
    assert "и ещё 10" in og.announcement(many)
    wb = load_workbook(io.BytesIO(og.notification_workbook(_outage())))
    assert wb.sheetnames == ["Уведомление", "Задвижки"]
    rows = [[c.value for c in r] for r in wb["Уведомление"].iter_rows()]
    assert rows[0][0].startswith("Плановое отключение О-2026-00003")
    assert [1, "Школа №1", "реальный", 50, 1.5, None, 43.2, 76.9] in rows
    assert [2, "Квартал 3", "обобщённый", 60, 2.0, "да", 43.22, 76.92] in rows
    assert [c.value for c in wb["Задвижки"][2]][:5] == [1, "З-1", 10, 500, "да"]


def test_public_feed_hides_names(monkeypatch):
    import main
    from auth import invalidate_user_status
    from routers import outages as router

    monkeypatch.setenv("AUTH_DISABLED", "true")
    invalidate_user_status()
    area = json.dumps({"type": "Polygon", "coordinates": [[[0, 0], [1, 0], [0, 1], [0, 0]]]})

    class FeedConn:
        async def fetch(self, sql, *args):
            assert "AND public" in sql and args[0] == list(og.PUBLIC_STATUSES)
            return [{"id": 3, "number": "О-3", "title": "Замена задвижки у школы", "kind": "planned",
                     "status": "announced", "planned_start": datetime(2026, 11, 2, 4),
                     "planned_end": datetime(2026, 11, 2, 13), "consumers_count": 2, "load_gcal_h": 3.5,
                     "public_note": "Плановые работы", "area": area}]

    @asynccontextmanager
    async def acquire():
        yield FeedConn()

    monkeypatch.setattr(router, "acquire_conn", acquire)
    client = TestClient(main.app)
    monkeypatch.setenv("PUBLIC_OUTAGES_ENABLED", "false")
    assert client.get("/api/v1/public/outages").status_code == 404
    monkeypatch.setenv("PUBLIC_OUTAGES_ENABLED", "true")
    props = client.get("/api/v1/public/outages").json()["features"][0]["properties"]
    assert props["note"] == "Плановые работы" and props["consumers_count"] == 2
    assert "title" not in props and "load_gcal_h" not in props
