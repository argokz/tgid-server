"""Формы обходов: описание формы, проверка ответов, импорт XLSForm, выгрузка."""

import io
import uuid

import pytest
from openpyxl import Workbook

from database import forms as fm

CHAMBER = {
    "fields": [
        {"name": "hatch", "type": "select_one", "label": "Люк", "required": True,
         "choices": [{"value": "ok", "label": "Исправен"}, {"value": "broken", "label": "Повреждён"},
                     {"value": "none", "label": "Нет люка"}]},
        {"name": "hatch_note", "type": "text", "label": "Что с люком",
         "relevant": {"field": "hatch", "op": "!=", "value": "ok"}, "required": True},
        {"name": "water", "type": "acknowledge", "label": "Вода в камере"},
        {"name": "water_level", "type": "decimal", "label": "Уровень воды, см", "min": 0, "max": 300,
         "relevant": {"field": "water", "op": "=", "value": "1"}, "required": True},
        {"name": "steam", "type": "select_multiple", "label": "Признаки",
         "choices": [{"value": "steam", "label": "Пар"}, {"value": "leak", "label": "Течь"}]},
        {"name": "leak_where", "type": "text", "label": "Где течь",
         "relevant": {"field": "steam", "op": "selected", "value": "leak"}},
        {"name": "temp", "type": "integer", "label": "Температура воздуха", "min": -50, "max": 120},
        {"name": "checked_on", "type": "date", "label": "Дата"},
        {"name": "info", "type": "note", "label": "Сфотографируйте камеру сверху"},
        {"name": "photos", "type": "image", "label": "Фото камеры"},
    ],
}


def schema(raw=None):
    return fm.FormSchema.model_validate(raw or CHAMBER)


def test_schema_rules():
    assert [f.name for f in schema().fields][:2] == ["hatch", "hatch_note"]
    bad = [
        {"fields": [{"name": "Bad Name", "type": "text", "label": "x"}]},
        {"fields": [{"name": "a", "type": "select_one", "label": "x"}]},  # без вариантов
        {"fields": [{"name": "a", "type": "text", "label": "x"}, {"name": "a", "type": "text", "label": "y"}]},
        {"fields": [{"name": "a", "type": "text", "label": "x", "relevant": {"field": "b", "value": "1"}},
                    {"name": "b", "type": "text", "label": "y"}]},  # условие на поле ниже
        {"fields": [{"name": "n", "type": "note", "label": "только пояснение"}]},
        {"fields": [{"name": "a", "type": "integer", "label": "x", "min": 5, "max": 1}]},
    ]
    for raw in bad:
        with pytest.raises(ValueError):
            fm.FormSchema.model_validate(raw)
    # у не числовых полей границы и у не выборов варианты отбрасываются
    s = fm.FormSchema.model_validate({"fields": [{"name": "a", "type": "text", "label": "x", "min": 1,
                                                  "choices": [{"value": "1", "label": "1"}]}]})
    assert s.fields[0].min is None and s.fields[0].choices is None


def test_validate_answers_relevance_and_types():
    photo = str(uuid.uuid4())
    clean, errors = fm.validate_answers(schema(), {
        "hatch": "ok", "hatch_note": "скрыто условием", "water": True, "water_level": "12,5",
        "steam": ["leak", "steam"], "leak_where": "на подаче", "temp": 18, "checked_on": "2026-10-31T10:00",
        "info": "x", "photos": [photo], "unknown": 1,
    })
    assert errors == {}
    assert clean == {"hatch": "ok", "water": True, "water_level": 12.5, "steam": ["steam", "leak"],
                     "leak_where": "на подаче", "temp": 18, "checked_on": "2026-10-31", "photos": [photo]}


def test_validate_answers_errors():
    _, errors = fm.validate_answers(schema(), {
        "hatch": "broken", "water": True, "water_level": 500, "steam": ["fire"], "temp": 1.5,
        "checked_on": "завтра", "photos": ["не uuid"],
    })
    assert errors == {
        "hatch_note": "обязательное поле", "water_level": "не больше 300", "steam": "нет такого варианта",
        "temp": "нужно целое число", "checked_on": "дата — ГГГГ-ММ-ДД", "photos": "фото — до 5",
    }
    _, errors = fm.validate_answers(schema(), {})
    assert errors == {"hatch": "обязательное поле"}
    _, errors = fm.validate_answers(schema(), {"hatch": "ok", "water": "да"})
    assert errors == {"water": "нужно да или пусто"}


def xlsform(survey, choices=None, settings=None) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "survey"
    for row in survey:
        ws.append(row)
    if choices:
        wc = wb.create_sheet("choices")
        for row in choices:
            wc.append(row)
    if settings:
        st = wb.create_sheet("settings")
        for row in settings:
            st.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_parse_xlsform_subset():
    data = xlsform(
        [
            ["type", "name", "label::English (en)", "label::Русский (ru)", "hint", "required", "relevant", "constraint"],
            ["start", "start", "", "", "", "", "", ""],
            ["begin_group", "g1", "Chamber", "Камера", "", "", "", ""],
            ["select_one yesno", "hatch_ok", "Hatch ok?", "Люк исправен?", "", "yes", "", ""],
            ["text", "hatch_why", "Why", "Что с люком", "", "true()", "${hatch_ok} = 'no'", ""],
            ["integer", "temp", "Temp", "Температура", "", "", "", ". >= -50 and . <= 120"],
            ["decimal", "depth", "Depth", "Глубина", "", "", "", ". > 0 or . = 0"],
            ["geopoint", "where", "Where", "Где", "", "", "", ""],
            ["select_multiple signs", "signs", "Signs", "Признаки", "", "", "", ""],
            ["text", "leak", "Leak", "Где течь", "", "", "selected(${signs}, 'leak')", ""],
            ["select_one missing", "bad", "Bad", "Плохой список", "", "", "", ""],
            ["image", "photo", "Photo", "Фото", "сверху", "", "", ""],
            ["end_group", "", "", "", "", "", "", ""],
        ],
        [["list_name", "name", "label::Русский (ru)"], ["yesno", "yes", "Да"], ["yesno", "no", "Нет"],
         ["signs", "steam", "Пар"], ["signs", "leak", "Течь"]],
        [["form_title", "form_id"], ["Осмотр камеры", "chamber"]],
    )
    form, warnings = fm.parse_xlsform(data)
    assert form["name"] == "Осмотр камеры"
    fields = {f["name"]: f for f in form["schema"]["fields"]}
    assert list(fields) == ["hatch_ok", "hatch_why", "temp", "depth", "signs", "leak", "photo"]
    assert fields["hatch_ok"]["label"] == "Люк исправен?" and fields["hatch_ok"]["required"] is True
    assert fields["hatch_ok"]["choices"] == [{"value": "yes", "label": "Да"}, {"value": "no", "label": "Нет"}]
    assert fields["hatch_why"]["relevant"] == {"field": "hatch_ok", "op": "=", "value": "no"}
    assert (fields["temp"]["min"], fields["temp"]["max"]) == (-50, 120)
    assert "min" not in fields["depth"]
    assert fields["leak"]["relevant"] == {"field": "signs", "op": "selected", "value": "leak"}
    assert fields["photo"]["type"] == "image" and fields["photo"]["hint"] == "сверху"
    text = "\n".join(warnings)
    assert "группа «Камера»" in text and "geopoint" in text and "missing" in text and ". > 0 or . = 0" in text
    fm.FormSchema.model_validate(form["schema"])  # перенесённое — правильная форма


def test_parse_xlsform_rejects_garbage():
    with pytest.raises(fm.FormError):
        fm.parse_xlsform(b"not a zip")
    with pytest.raises(fm.FormError) as exc:
        fm.parse_xlsform(xlsform([["type", "name", "label"], ["geopoint", "p", "Где"]]))
    assert "ни одного" in exc.value.message


def test_display_for_export():
    f = {"type": "select_multiple", "choices": [{"value": "a", "label": "А"}, {"value": "b", "label": "Б"}]}
    assert fm._display(f, ["a", "b"]) == "А, Б"
    assert fm._display({"type": "image"}, ["x", "y"]) == "фото: 2"
    assert fm._display({"type": "acknowledge"}, True) == "да"
