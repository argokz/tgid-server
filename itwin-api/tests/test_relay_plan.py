"""План перекладок (database/relay_plan.py): расценки, жадный отбор, объекты, сравнение «старые первыми»."""

import pytest

from database import relay_plan as rp
from database.wear import flow_factor, priority

PRICES = [{"dn_max_mm": 100, "price_per_m": 50.0}, {"dn_max_mm": 300, "price_per_m": 120.0},
          {"dn_max_mm": 700, "price_per_m": 300.0}]


def _item(lineid, benefit, cost, fileid=1, n1=None, n2=None, age=None, prob=None):
    return {"lineid": lineid, "benefit": benefit, "cost": cost, "fileid": fileid, "node1": n1, "node2": n2,
            "age": age, "prob": prob if prob is not None else benefit, "len_m": cost * 1000}


def test_price_for_classes():
    assert rp.price_for(80, PRICES) == 50 and rp.price_for(100, PRICES) == 50
    assert rp.price_for(250, PRICES) == 120 and rp.price_for(1000, PRICES) == 300
    assert rp.price_for(None, PRICES) == 120 and rp.price_for(80, []) is None


def test_select_by_benefit_per_cost_within_budget():
    items = [_item(1, 0.5, 2.0), _item(2, 0.3, 0.5), _item(3, 0.2, 0.5), _item(4, 0.9, 5.0), _item(5, 0.0, 0.1)]
    chosen = rp.select(items, 3.2)
    # на единицу стоимости: 2 (0,6), 3 (0,4), 1 (0,25), 4 (0,18); 4 не влезает; польза 0 — не берётся
    assert [x["lineid"] for x in chosen] == [2, 3, 1]
    # постоянные затраты объекта 1,0: 2 (0,3/1,5), 1 (0,5/3), 3 (0,2/1,5) — обрезок 3 теперь после 1
    assert [x["lineid"] for x in rp.select(items, 3.2, overhead=1.0)] == [2, 1, 3]


def test_build_runs_between_branches_with_supply_return_corridor():
    def line(i, n1, n2, dn=100, fileid=1, b=0.1):
        return {**_item(i, b, 0.1, fileid=fileid, n1=n1, n2=n2), "dn": dn}

    items = [
        line(1, 10, 11), line(2, 11, 10),                 # подача и обратка 10–11 — один коридор
        line(3, 11, 12), line(4, 12, 13),                 # 11 и 12 — проходные: трасса 10–13
        line(5, 13, 14), line(6, 13, 15),                 # 13 — развилка: 13–14 и 13–15 — отдельные трассы
        line(7, 15, 16, dn=200),                          # смена диаметра в 15 — граница трассы
        line(8, 16, 17, fileid=2),                        # другой фрагмент
        {**_item(9, 0.5, 0.2), "dn": None},              # без узлов
    ]
    runs = {tuple(sorted(x["lineid"] for x in r["lines"])) for r in rp.build_runs(items)}
    assert runs == {(1, 2, 3, 4), (5,), (6,), (7,), (8,), (9,)}
    r = next(r for r in rp.build_runs(items) if len(r["lines"]) == 4)
    assert r["benefit"] == pytest.approx(0.4) and r["cost"] == pytest.approx(0.4) and r["lineid"] == 1


def test_oldest_first_comparison():
    items = [_item(1, 0.1, 1, age=50), _item(2, 0.8, 1, age=5), _item(3, 0.2, 1, age=None)]
    assert [x["lineid"] for x in rp.oldest_first(items, 2)] == [1, 2]


def test_flow_factor_shared_with_wear():
    assert flow_factor(None) == 1 and flow_factor(500) == 1.25 and flow_factor(-3000) == 1.5
    assert priority(50.0, 1.4, 500) == round(0.5 * 1.4 * 1.25, 4)


def test_severity_by_objective():
    x = {"consumers": 12, "load": 0.8, "consequence": 2.0, "flow_t_h": 1000}
    assert rp.severity("consumers", x) == 12 and rp.severity("load", x) == 0.8
    assert rp.severity("wear", x) == pytest.approx(2.0 * 1.5) and rp.severity("failures", x) == 1
    # последствия отключения неизвестны — участок по ним не выбирается
    assert rp.severity("consumers", {**x, "consumers": None}) == 0


def test_expected_shares():
    lines = [{"prob": 0.5, "consumers": 10, "load": 1.0, "social": 1}, {"prob": 0.1, "consumers": None,
                                                                         "load": None, "social": None}]
    totals = {"failures": 2.0, "consumers": 20.0, "load": 0.0, "social": 1.0}
    e = rp.expected(lines, totals)
    assert e["lines"] == 2 and e["expected_failures"] == 0.6 and e["expected_failures_share"] == 0.3
    assert e["expected_consumers"] == 5 and e["expected_consumers_share"] == 0.25
    assert e["expected_load_share"] is None and e["expected_social_share"] == 0.5


def test_plan_workbook_sheets():
    from io import BytesIO

    from openpyxl import load_workbook

    old = {"lines": 1, **{f"expected_{k}": 0.1 for k in ("failures", "consumers", "load", "social")},
           **{f"expected_{k}_share": 0.01 for k in ("failures", "consumers", "load", "social")}}
    p = {"unit": "km", "currency": None, "budget": 5, "objective_label": "отключаемые потребители",
         "horizon_years": 2, "job_overhead_m": 25, "run_id": 1, "oldest_first": old, **old, "lines": 2,
         "objects": [{"n": 1, "fragment": "тм-1", "lines": 2, "length_m": 40.0, "dn_min": 100, "dn_max": 150,
                      "expected_failures": 0.3, "consumers_max": 12, "load_max": 0.5, "social_max": 1, "cost": 0.04}],
         "chosen_lines": [{"object": 1, "lineid": 7, "fragment": "тм-1", "len_m": 20.0, "dn": 100, "age": 31.0,
                           "prob": 0.2, "consumers": 12, "load": 0.5, "social": 1, "cost": 0.02},
                          {"object": 1, "lineid": 8, "fragment": "тм-1", "len_m": 20.0, "dn": 150, "age": None,
                           "prob": 0.1, "consumers": None, "load": None, "social": None, "cost": 0.02}]}
    wb = load_workbook(BytesIO(rp.plan_workbook(p)))
    assert wb.sheetnames == ["Итоги", "Объекты", "Участки"]
    assert wb["Объекты"].max_row == 2 and wb["Участки"].max_row == 3
    assert wb["Участки"]["B2"].value == 7 and wb["Итоги"]["B2"].value == "отключаемые потребители"


def test_merge_copies_keeps_one_pipe_per_place():
    def x(lid, fileid, prob, ends, consumers=None):
        return {"lineid": lid, "fileid": fileid, "prob": prob, "ends": ends, "consumers": consumers, "load": None,
                "social": None}

    e1, e2 = ((0, 0), (10, 0), 10), ((10, 0), (20, 0), 10)
    items = [
        x(1, 1, 0.30, e1), x(2, 1, 0.05, e1),           # фрагмент 1: подача и обратка на одной геометрии
        x(3, 2, 0.10, e1, 40), x(4, 2, 0.20, e1, 10),   # копия во фрагменте 2
        x(5, 1, 0.20, e2), x(6, 2, 0.01, e2),
        x(7, 3, 0.50, None),                            # без геометрии — не склеивается
        x(8, 1, 0.40, ((5, 5), (6, 6), 1)),             # только во фрагменте 1
    ]
    out, dropped = rp.merge_copies(items)
    assert dropped == 3 and sorted(i["lineid"] for i in out) == [1, 2, 5, 7, 8]
    by = {i["lineid"]: i for i in out}
    # остаётся фрагмент 1 (суммарный риск 0,95 против 0,31); вероятности — попарно наибольшие по убыванию
    assert by[1]["prob"] == 0.30 and by[2]["prob"] == 0.10 and by[1]["consumers"] == 40
    assert items[1]["prob"] == 0.05  # исходные участки не меняются
