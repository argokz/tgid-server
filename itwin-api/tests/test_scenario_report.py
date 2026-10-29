"""Excel-отчёт «сценарий против основы» (database/scenario_report.py) и его адрес."""

import io
from contextlib import asynccontextmanager

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from database import scenario_report

RESULT = {
    "scenario_id": 4, "name": "Перекладка", "warnings": ["Разная температура наружного воздуха: -25 и -20 °C"],
    "base_calculation": {"id": 710, "fileid": 74, "tn": -25.0, "name": None, "date": "2026-10-29T21:26:47.5",
                         "emergency": False},
    "scenario_calculation": {"id": 21, "fileid": 115, "tn": -20.0, "name": "Вариант", "date": None, "emergency": True},
    "consumers": {"base": 958, "scenario": 944, "underheated_base": 429, "underheated_scenario": 414,
                  "load_base": 2765.7, "load_scenario": 2404.7, "flow_base": 40227.6, "flow_scenario": 39275.5,
                  "became_underheated": 1, "recovered": 0, "lost": 0, "new": 0, "changed": 0},
    "lines": {"base": 1692, "scenario": 1638, "overloaded_base": 10, "overloaded_scenario": 12,
              "became_overloaded": 1, "became_fast": 0, "relieved": 0, "lost": 0, "new": 0, "changed": 0},
    "thresholds": {"underheat_pct": 100.0, "head_delta_m": 0.5, "spec_loss_pa_m": 80.0, "velocity_ms": 1.5},
    "total": 3, "truncated": True,
    "items": [
        {"object": "consumer", "kind": "became_underheated", "base_id": 5, "copy_id": 905, "name": "ТК-5",
         "head": [12.0, 9.5], "supply": [100.0, 92.0], "flow": [10.0, None], "lat": 43.2, "lng": 76.9},
        {"object": "line", "kind": "became_overloaded", "base_id": 7, "copy_id": 907,
         "flow": [100.0, 140.0], "velocity": [1.0, 1.4], "loss": [60.0, 95.5], "lat": None, "lng": None},
    ],
}


def test_workbook_sheets_and_values():
    wb = load_workbook(io.BytesIO(scenario_report.workbook(RESULT)))
    assert wb.sheetnames == ["Сводка", "Потребители", "Участки"]
    summary = [[c.value for c in row] for row in wb["Сводка"].iter_rows()]
    flat = str(summary)
    assert "№710, фрагмент 74, Тн -25 °C, плановый, 29.10.2026 21:26" in flat
    assert "№21 «Вариант», фрагмент 115, Тн -20 °C, аварийный" in flat
    assert ["С недотопом", 429, 414, -15] in [row[:4] for row in summary]
    assert ["Стал недотоп", 1, 0] in [row[:3] for row in summary]
    assert any("первые 2 изменений из 3" in str(row[0]) for row in summary)
    cons = [c.value for c in wb["Потребители"][2]]
    assert cons[:4] == ["Стал недотоп", "ТК-5", 5, 905] and cons[4:7] == [12.0, 9.5, -2.5]
    assert cons[10:13] == [10.0, None, None]  # нет значения в сценарии — разницы нет
    line = [c.value for c in wb["Участки"][2]]
    assert line[:3] == ["Участок перегружен", 7, 907] and line[9:12] == [60.0, 95.5, 35.5]


def test_endpoint(monkeypatch):
    import main
    from auth import invalidate_user_status
    from routers import scenarios

    monkeypatch.setenv("AUTH_DISABLED", "true")
    invalidate_user_status()
    seen = {}

    async def fake_compare(conn, scenario_id, base_id, scen_id, limit):
        seen["limit"] = limit
        return RESULT if scenario_id == 4 else None

    @asynccontextmanager
    async def acquire():
        yield object()

    monkeypatch.setattr(scenarios, "compare_scenario_calculations", fake_compare)
    monkeypatch.setattr(scenarios, "acquire_conn", acquire)
    client = TestClient(main.app)
    params = {"base_calculation_id": 710, "scenario_calculation_id": 21}
    r = client.get("/api/v1/scenarios/4/compare.xlsx", params=params)
    assert r.status_code == 200 and seen["limit"] == scenario_report.REPORT_ITEMS_LIMIT
    assert "filename*=UTF-8''" in r.headers["content-disposition"]
    assert load_workbook(io.BytesIO(r.content)).sheetnames[0] == "Сводка"
    assert client.get("/api/v1/scenarios/5/compare.xlsx", params=params).status_code == 404
