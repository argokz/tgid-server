"""Пьезометр показывает ровно us_out одного расчёта (sety/out/us_out.py), а не измеренные давления."""

import asyncio
from unittest.mock import AsyncMock

import networkx as nx
import pytest

from routers.piezometer import _assemble_path_data, _route_calculation_id


def _node(node_id, z, pih_pod, pih_obr, meas_pod=None, meas_obr=None):
    return {"id": node_id, "geomarktoptube": z, "meas_pod": meas_pod, "meas_obr": meas_obr,
            "label": str(node_id), "lng": 76.9, "lat": 43.2, "t_pod": 130.0, "t_obr": 70.0,
            "pih_pod": pih_pod, "pih_obr": pih_obr}


def test_heads_come_from_us_out_not_measured_pressure():
    G = nx.Graph()
    G.add_edge(1, 2, weight=30.0)
    conn = AsyncMock()
    # у узла 1 есть измеренное давление 55 м — оно не должно подменять расчётные 62.2 м
    conn.fetch.return_value = [_node(1, 799.2, 62.2, 20.1, meas_pod=55.0), _node(2, 800.0, 61.0, 21.0)]

    path = asyncio.run(_assemble_path_data(conn, G, [1, 2], 9))

    sql, nodes, cid = conn.fetch.call_args.args
    assert nodes == [1, 2] and cid == 9
    # nodes.calcpress* по умолчанию 0 (gid8 baza.sql) — 0 означает «не измерено»
    assert "NULLIF(n.calcpressflow, 0)" in sql
    # pt_out и us_out — того же расчёта, что и весь маршрут
    assert "pt_out WHERE nodeid = n.id AND calculationid = $2" in sql
    assert "us1.calculationid = $2" in sql and "us2.calculationid = $2" in sql
    assert path[0]["h_pod"] == pytest.approx(799.2 + 62.2)
    assert path[0]["h_obr"] == pytest.approx(799.2 + 20.1)
    assert path[0]["h_pod_meas"] == pytest.approx(799.2 + 55.0)
    assert path[1]["h_pod_meas"] is None
    assert path[1]["distance"] == 30.0


def test_route_uses_one_calculation():
    conn = AsyncMock()
    assert asyncio.run(_route_calculation_id(conn, [1, 2], 5)) == 5
    conn.fetchval.assert_not_called()
    conn.fetchval.return_value = 9
    assert asyncio.run(_route_calculation_id(conn, [1, 2], None)) == 9
    assert "max(calculationid) FROM us_out" in conn.fetchval.call_args.args[0]


# --- этап 10: двойной и статический пьезометр, сохранённые направления -----------------

def test_static_head_is_max_mark_plus_building_plus_5():
    """gid8 Pjezo.cpp m_stat: горизонталь h_max + 5, h_max = max(geoMarkTopTube + hz)."""
    from routers.piezometer import STATIC_RESERVE_M, _static_head

    conn = AsyncMock()
    conn.fetchrow.return_value = {"id": 7, "label": "ТК-7", "z": 800.0, "hz": 30.0}
    res = asyncio.run(_static_head(conn, [74]))
    sql, frags = conn.fetchrow.call_args.args
    assert frags == [74]
    assert "rc.buildheight, gc.maxbuildingheight" in sql and "geomarktoptube > 0" in sql
    assert STATIC_RESERVE_M == 5.0
    assert res["value"] == pytest.approx(835.0) and res["node_id"] == 7
    assert asyncio.run(_static_head(conn, [])) is None


def test_second_calculation_heads_on_same_route():
    from routers.piezometer import _attach_second_calc

    conn = AsyncMock()
    conn.fetch.return_value = [
        {"nodeid": 1, "externalsign": 1, "pih": 60.0},
        {"nodeid": 1, "externalsign": 2, "pih": 20.0},
    ]
    path = [{"node_id": 1, "z": 800.0}, {"node_id": 2, "z": 801.0}]
    asyncio.run(_attach_second_calc(conn, path, 3))
    assert conn.fetch.call_args.args[2] == 3
    assert path[0]["h_pod_2"] == pytest.approx(860.0) and path[0]["h_obr_2"] == pytest.approx(820.0)
    assert path[1]["h_pod_2"] is None and path[1]["h_obr_2"] is None


def test_excel_has_static_and_second_calc_columns():
    import io

    import openpyxl

    from database.piezo_excel import generate_piezometer_excel

    path = [
        {"node_id": 1, "distance": 0, "z": 800, "h_pod": 860, "h_obr": 820, "h_pod_2": 858, "h_obr_2": 821, "label": "A"},
        {"node_id": 2, "distance": 50, "z": 801, "h_pod": 859, "h_obr": 821, "h_pod_2": 857, "h_obr_2": 822, "label": "B"},
    ]
    data = generate_piezometer_excel(path, [], static_head=835.0, second_calculation_id=3)
    ws = openpyxl.load_workbook(io.BytesIO(data))["Пьезометрический профиль"]
    head = [c.value for c in ws[1]]
    assert head[4] == "Статический напор (м)" and "расчёт 3" in head[5] and head[-1] == "Узел"
    assert [c.value for c in ws[2]][4:7] == [835.0, 858, 821]
    # без опций — прежний набор колонок
    ws0 = openpyxl.load_workbook(io.BytesIO(generate_piezometer_excel(path, [])))["Пьезометрический профиль"]
    assert len(ws0[1]) == 5


def test_direction_writes_need_editor_and_mutations(monkeypatch):
    import os

    from fastapi.testclient import TestClient

    os.environ.setdefault("JWT_SECRET", "test-secret")
    import main
    from auth import create_access_token

    def bearer(role):
        return {"Authorization": "Bearer " + create_access_token(username=f"t-{role}", role=role)}

    monkeypatch.setenv("AUTH_DISABLED", "false")
    client = TestClient(main.app)
    body = {"name": "x", "nodes": [1, 2]}
    monkeypatch.setenv("MUTATIONS_ENABLED", "false")
    assert client.post("/piezometer/directions", json=body, headers=bearer("editor")).status_code == 503
    assert client.delete("/piezometer/directions/1", headers=bearer("editor")).status_code == 503
    monkeypatch.setenv("MUTATIONS_ENABLED", "true")
    assert client.post("/piezometer/directions", json=body, headers=bearer("viewer")).status_code == 403
    assert client.delete("/piezometer/directions/1", headers=bearer("viewer")).status_code == 403
    assert client.post("/piezometer/directions", json={"name": "x", "nodes": [1]},
                       headers=bearer("editor")).status_code == 422
